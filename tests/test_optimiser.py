"""CV optimiser improvements (user's decisions 2026-10-14): plain wording, keyword limits,
years of experience, length advice, Word read-back warnings, and the optimised mode's red
layer (proposal kinds, compressions, declined proposals, the pass toward the target score)."""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.core.config import Settings
from src.cv.ats import check_docx, layout_warnings
from src.cv.ats_score import grade, score_cv, term_found
from src.cv.claims import jargon, third_person
from src.cv.cover_letter import _paragraph_reason
from src.cv.docx_exporter import export_docx
from src.cv.docx_original import write_like_original
from src.cv.fit_check import WEIGHTS, compare_fit
from src.cv.models import (
    AdditionCandidate,
    Bullet,
    CompressionCandidate,
    CompressionProposal,
    EvidencedReading,
    EvidencedReadings,
    JDAnalysis,
    MasterCV,
    MetRequirement,
    ProjectedFit,
    RatedLine,
    RequirementEvidence,
    ReviewAddition,
    RewrittenBullet,
    TailoredCV,
    TailoredCVEdits,
    TailoringPlan,
    TailoringSettings,
    TermBridge,
)
from src.cv.proposals import (
    AdditionContext,
    apply_compression,
    check_additions,
    check_compressions,
    enriched,
    merged_heading,
    rank_additions,
    rewordings,
    with_proposals,
)
from src.cv.tailor import (
    MAX_ROLE_SKILLS,
    PLAN_SYSTEM,
    apply_plan,
    career_years,
    cv_to_text,
    length_notes,
    literal_terms,
    printed_text,
    reader_feedback,
    screen,
    tailor,
)
from src.jobs.models import JobPosting
from src.services import declined, tailored_documents
from src.services.workspace import Workspace
from tests.conftest import FakeLLM

JD = JDAnalysis(
    job_title="Senior ML Engineer",
    hard_skills=["Python", "PyTorch", "Kubernetes", "MLOps"],
    must_have=["Python"],
)
JOB_TEXT = (
    "Senior ML Engineer. Build recommendation services in Python and PyTorch, run model "
    "training on Kubernetes with MLOps practice, and own model monitoring in production."
)
JOB = JobPosting(id="job-o", title="Senior ML Engineer", company="Orbit AI", description=JOB_TEXT)


def _ctx(master: MasterCV, **kw: object) -> AdditionContext:
    tailored = apply_plan(master, TailoringPlan(), JD)
    defaults: dict[str, object] = {
        "master": master,
        "cv_text": cv_to_text(master),
        "printed": printed_text(tailored),
        "jd_text": JOB_TEXT,
        "requirements": frozenset({"model monitoring"}),
        "allowed_numbers": frozenset({"8"}),
    }
    return AdditionContext(**{**defaults, **kw})  # type: ignore[arg-type]


def _proposal(**kw: object) -> AdditionCandidate:
    base: dict[str, object] = {
        "kind": "enrichment",
        "text": "with model monitoring dashboards",
        "anchor_id": "nimbus-2",
        "requirement": "model monitoring",
        "background_quote": "model training pipelines",
        "job_quote": "own model monitoring in production",
        "reason": "Pipelines in production usually come with monitoring; confirm you set it up.",
    }
    return AdditionCandidate.model_validate({**base, **kw})


# --- wording, keywords, years ------------------------------------------------------------


def test_jargon_is_rejected_only_when_the_rewrite_adds_it(master_cv: MasterCV) -> None:
    assert jargon("Built a service", "Leveraged Python to build a service")
    assert jargon("Leveraged Python to build a service", "Leveraged Python for a service") is None
    plan = TailoringPlan(
        rewritten_bullets=[
            RewrittenBullet(
                source_id="nimbus-3",
                text="Mentored 5 engineers and fostered robust code review standards for ML code.",
            )
        ]
    )
    change = apply_plan(master_cv, plan, JD).changes[0]
    assert not change.accepted and "jargon" in (change.reason or "")
    sources = {"nimbus-1": "Built a recommendation service in Python"}
    reason = _paragraph_reason(
        "I am a passionate engineer who built a recommendation service in Python.",
        ["nimbus-1"], sources, JOB_TEXT, ["Python"],
    )  # fmt: skip
    assert reason and "jargon" in reason  # flagged in red, the letter is still written


def test_cv_text_is_written_in_the_first_person(master_cv: MasterCV) -> None:
    name = master_cv.basics.name
    assert third_person("Built a service", "She built a service", name)
    assert third_person("", f"{name.split()[0]} is an ML engineer.", name)
    assert third_person("", "The candidate builds ML systems.", name)
    assert third_person("", "I build ML systems with PyTorch.", name) is None
    assert third_person("Supported her grant", "Supported her grant in Python") is None
    third = TailoringPlan(summary="The candidate is an ML engineer shipping PyTorch systems.")
    out = apply_plan(master_cv, third, JD)
    assert out.cv.basics.summary == master_cv.basics.summary
    assert "third person" in (out.changes[0].reason or "")
    first = TailoringPlan(summary="I am an ML engineer shipping PyTorch systems.")
    assert apply_plan(master_cv, first, JD).changes[0].accepted
    plan = TailoringPlan(
        rewritten_bullets=[RewrittenBullet(source_id="nimbus-3", text="He mentored 5 engineers.")]
    )
    change = apply_plan(master_cv, plan, JD).changes[0]
    assert not change.accepted and "third person" in (change.reason or "")


def test_a_job_term_is_brought_into_two_bullets_at_most(master_cv: MasterCV) -> None:
    jd = JDAnalysis(job_title="ML Engineer", hard_skills=["Python"], must_have=["model serving"])
    rows = [
        RequirementEvidence(
            requirement="model serving", priority="essential", importance="key", match="direct",
            evidence=["nimbus-1", "nimbus-2", "lexa-2"],
        )
    ]  # fmt: skip
    plan = TailoringPlan(
        requirements=rows,
        rewritten_bullets=[
            RewrittenBullet(
                source_id="nimbus-1",
                text="Built a real-time recommendation service for model serving in Python and "
                "PyTorch serving 4M users, lifting click-through rate by 18%.",
            ),
            RewrittenBullet(
                source_id="nimbus-2",
                text="Led migration of model training and model serving pipelines to Kubernetes, "
                "cutting training costs by 35%.",
            ),
            RewrittenBullet(
                source_id="lexa-2",
                text="Designed PostgreSQL feature store for model serving used by 3 data "
                "science teams.",
            ),
        ],
        bullet_order=[{"experience_id": "nimbus", "bullet_ids": ["nimbus-1", "nimbus-2"]}],  # type: ignore[list-item]
        term_bridges=[
            TermBridge(term="model serving", source_id=i, quote=q, requirement="model serving")
            for i, q in (
                ("nimbus-1", "real-time recommendation service"),
                ("nimbus-2", "model training pipelines"),
                ("lexa-2", "PostgreSQL feature store"),
            )
        ],
    )
    out = apply_plan(master_cv, plan, jd)
    assert len(out.term_bridges) == 3
    undone = [c for c in out.changes if not c.accepted]
    assert [c.source_id for c in undone] == ["lexa-2"]
    assert "repeats the job term" in (undone[0].reason or "")
    assert "model serving" not in out.cv.bullet_index()["lexa-2"].text


def test_years_of_experience_may_be_stated_but_no_other_new_number(master_cv: MasterCV) -> None:
    years = career_years(master_cv)
    assert years >= 9
    ok = apply_plan(
        master_cv,
        TailoringPlan(
            summary=f"Machine learning engineer with {years} years building recommendation "
            "and NLP systems in Python.",
            headline_options=[f"Senior ML Engineer | {years} years"],
        ),
        JD,
    )
    assert ok.cv.basics.summary and str(years) in ok.cv.basics.summary
    assert ok.headline_options == [f"Senior ML Engineer | {years} years"]
    bad = apply_plan(master_cv, TailoringPlan(summary="Engineer with 25 years in ML."), JD)
    assert bad.cv.basics.summary == master_cv.basics.summary


def test_the_skills_line_holds_only_skills_shown_in_use(master_cv: MasterCV) -> None:
    plan = TailoringPlan(role_skills=["PyTorch", "Go", "Docker", "Kubernetes", "AWS"])
    out = apply_plan(master_cv, plan, JD)
    assert "PyTorch" in out.role_skills and "Kubernetes" in out.role_skills
    assert not {"Go", "Docker", "AWS"} & set(out.role_skills)  # only in the skills list
    assert MAX_ROLE_SKILLS == 8  # a few skills (user's decision 2026-10-21)


def test_standard_revision_stops_asking_for_keywords_past_80_percent(master_cv: MasterCV) -> None:
    def feedback(jd: JDAnalysis, mode: str = "standard") -> str:
        out = apply_plan(master_cv, TailoringPlan(), jd)
        out.screening = screen(master_cv, out, jd, "Python PyTorch Kubernetes NLP Rust Scala")
        return reader_feedback(out, jd, mode)  # type: ignore[arg-type]

    many = JDAnalysis(job_title="x", hard_skills=["Python", "PyTorch", "Kubernetes", "NLP", "Rust"])
    assert "keywords the printed CV does not contain" not in feedback(many)  # 4 of 5 shown
    assert "Rust" in feedback(many, "optimised")
    few = JDAnalysis(job_title="x", hard_skills=["Python", "Rust", "Scala"])
    assert "Rust" in feedback(few)


def test_length_notes_advise_without_changing_anything(master_cv: MasterCV) -> None:
    cv = master_cv.model_copy(deep=True)
    lexa = cv.experience[1]
    lexa.end = "2012-01"
    lexa.start = "2009-01"
    long = " ".join(["Developed"] + ["models"] * 40)
    lexa.bullets[0].text = long
    lexa.bullets += [lexa.bullets[1].model_copy(update={"id": f"lexa-x{n}"}) for n in range(3)]
    notes = length_notes(cv)
    assert any("over 35 words" in n for n in notes)
    assert any("overlapping work" in n for n in notes)
    assert any("ended over 10 years ago" in n for n in notes)
    out = apply_plan(cv, TailoringPlan(), JD)
    assert out.length_notes and len(out.cv.bullet_index()) == len(cv.bullet_index())
    assert apply_plan(cv, TailoringPlan(), JD, mode="optimised").length_notes == []


# --- Word read-back ------------------------------------------------------------------------


def test_read_back_warns_about_layouts_parsers_miss(master_cv: MasterCV, tmp_path: Path) -> None:
    path = tmp_path / "cv.docx"
    doc = Document()
    doc.add_paragraph(master_cv.basics.name)
    doc.add_paragraph("Where I worked")
    doc.add_paragraph("Nimbus Analytics, Senior Machine Learning Engineer")
    doc.add_paragraph("Built a real-time recommendation service in Python.")
    doc.sections[0].footer.paragraphs[0].text = master_cv.basics.email or ""
    cols = doc.sections[0]._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        doc.sections[0]._sectPr.append(cols)
    cols.set(qn("w:num"), "2")
    doc.paragraphs[0].runs[0].font.name = "Papyrus"
    doc.save(str(path))
    warnings = " | ".join(layout_warnings(path, master_cv))
    assert "columns" in warnings and "footer" in warnings
    assert "No year next to the role: Nimbus Analytics" in warnings
    assert "experience, education" in warnings and "Papyrus" in warnings
    clean = tmp_path / "clean.docx"
    export_docx(master_cv, clean)
    assert layout_warnings(clean, master_cv) == []


# --- optimised proposals -------------------------------------------------------------------


def test_proposal_kinds_are_checked_where_they_belong(master_cv: MasterCV) -> None:
    cv = master_cv.model_copy(deep=True)
    cv.experience[0].bullets[2].text = "Introduced code review standards for ML code."
    ctx = _ctx(cv)
    kept = check_additions(
        [
            _proposal(),  # words inside nimbus-2
            _proposal(anchor_id="lexa-1"),  # quote from another line
            _proposal(kind="metric", text="how many models?", anchor_id="nimbus-2"),  # has one
            _proposal(
                kind="metric", text="how many reviews a month?", anchor_id="nimbus-3",
                background_quote="code review standards",
            ),
            _proposal(kind="skill", text="Model monitoring", anchor_id=None, skill_group="ML"),
            _proposal(kind="headline", text="Principal ML Engineer", anchor_id=None),  # too high
            _proposal(
                kind="summary", text="Leveraged dashboards for model monitoring.", anchor_id=None
            ),
        ],
        ctx,
    )  # fmt: skip
    assert [(a.kind, a.anchor_id) for a in kept] == [
        ("enrichment", "nimbus-2"),
        ("metric", "nimbus-3"),
        ("skill", None),
    ]
    assert kept[0].role_id == "nimbus" and kept[0].original_text.startswith("Led migration")
    assert kept[2].skill_group == "ML"


def test_other_cvs_ground_an_inference_and_declined_ones_never_return(master_cv: MasterCV) -> None:
    other = "Set up Grafana monitoring for the recommendation models at Nimbus Analytics."
    experience = _proposal(
        kind="experience", text="Set up monitoring for production recommendation models",
        anchor_id="nimbus-1", background_quote="Set up Grafana monitoring",
    )  # fmt: skip
    assert check_additions([experience], _ctx(master_cv)) == []  # not in this CV
    kept = check_additions([experience], _ctx(master_cv, background=other))
    assert kept and kept[0].role_id == "nimbus"
    said_no = ("Set up monitoring for production recommendation models.",)
    declined_ctx = _ctx(master_cv, background=other, declined=said_no)
    assert check_additions([experience], declined_ctx) == []


def test_caps_scale_with_the_gaps_and_likely_proposals_come_first(master_cv: MasterCV) -> None:
    words = ("alpha", "beta", "gamma", "delta", "epsilon", "zeta", "eta", "theta", "iota")
    items = [_proposal(text=f"with monitoring step {word}") for word in words]
    assert len(check_additions(items, _ctx(master_cv))) == 8  # at most eight per role
    ctx = _ctx(master_cv)
    assert ctx.limit() == 12 and replace(ctx, open_items=40).limit() == 30
    possible = [i.model_copy(update={"likelihood": "possible"}) for i in items[:6]]
    likely = [i.model_copy(update={"text": i.text + " now"}) for i in items[6:]]
    kept = check_additions([*possible, *likely], ctx)
    assert [k.likelihood for k in kept[:3]] == ["likely"] * 3  # likely first
    assert sum(k.likelihood == "possible" for k in kept) == 4  # a third of the limit (12)


def test_proposals_that_change_nothing_are_dropped_and_the_rest_ranked() -> None:
    def addition(id_: str, kind: str = "experience", text: str = "x y") -> ReviewAddition:
        return ReviewAddition(
            id=id_, kind=kind, text=text, requirement="r", background_quote="a b",  # type: ignore[arg-type]
            job_quote="c d", reason="e",
        )  # fmt: skip

    projection = ProjectedFit(
        confirmed=60, projected=80,
        gained=[{"requirement": "r1", "addition_id": "b"},
                {"requirement": "r2", "addition_id": "b"},
                {"requirement": "r3", "addition_id": "a"}],
    )  # type: ignore[list-item]  # fmt: skip
    ranked = rank_additions(
        [addition("a"), addition("b"), addition("c"), addition("d", "metric"),
         addition("e", text="uses Rust")],
        projection, ["Rust"],
    )  # fmt: skip
    assert [a.id for a in ranked] == ["b", "a", "e", "d"]


def test_the_reading_places_words_inside_their_line(master_cv: MasterCV) -> None:
    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    item = check_additions([_proposal()], _ctx(master_cv))[0]
    tailored.review_additions = [item]
    projected = with_proposals(tailored)
    assert projected is not None
    line = projected.cv.bullet_index()["nimbus-2"].text
    assert line.endswith("cutting training costs by 35%, with model monitoring dashboards.")


# --- compressions --------------------------------------------------------------------------


def _postdocs(master: MasterCV) -> MasterCV:
    cv = master.model_copy(deep=True)
    first = cv.experience[1].model_copy(
        update={"id": "pd1", "title": "Postdoctoral Research Fellow", "company": "University A",
                "start": "2010-01", "end": "2014-06"}
    )  # fmt: skip
    second = cv.experience[1].model_copy(
        update={"id": "pd2", "title": "Postdoctoral Research Fellow", "company": "Institute B",
                "start": "2014-07", "end": "2017-08"}
    )  # fmt: skip
    first.bullets = [b.model_copy(update={"id": f"pd1-{n}"}) for n, b in enumerate(first.bullets)]
    second.bullets = [b.model_copy(update={"id": f"pd2-{n}"}) for n, b in enumerate(second.bullets)]
    cv.experience = [cv.experience[0], second, first]
    return MasterCV.model_validate(cv.model_dump())


def test_compressions_keep_every_number_and_job_term(master_cv: MasterCV) -> None:
    vocab = [*JD.hard_skills, *master_cv.all_skills()]
    good = CompressionCandidate(
        kind="merge_bullets", role_ids=["nimbus"], source_ids=["nimbus-2", "nimbus-3"],
        bullets=["Led migration of training pipelines to Kubernetes (35% lower costs) and "
                 "mentored 5 engineers on ML code review."],
        reason="Two lines on the same platform work",
    )  # fmt: skip
    drops = good.model_copy(
        update={"bullets": ["Led migration to Kubernetes; mentored engineers."]}
    )
    adds = good.model_copy(
        update={"bullets": [good.bullets[0] + " Used Terraform for 12 clusters."]}
    )
    kept = check_compressions([drops, adds, good], master_cv, ["Kubernetes"], vocab)
    assert [k.source_ids for k in kept] == [["nimbus-2", "nimbus-3"]]
    out = apply_compression(master_cv, kept[0])
    role = out.experience[0]
    assert [b.id for b in role.bullets] == ["nimbus-1", "nimbus-2"]
    assert role.bullets[1].text == good.bullets[0]


def test_different_roles_are_never_combined(master_cv: MasterCV) -> None:
    cv = _postdocs(master_cv)
    merge = CompressionCandidate(
        kind="merge_roles", role_ids=["pd1", "pd2"],
        bullets=["Developed NLP models for document classification with 92% accuracy using "
                 "scikit-learn and spaCy, and designed a PostgreSQL feature store used by 3 data "
                 "science teams."],
        reason="Two postdocs with the same work",
    )  # fmt: skip
    assert check_compressions([merge], cv, [], [*cv.all_skills()]) == []
    assert "never combine different roles" in PLAN_SYSTEM


SHORTER_NIMBUS = CompressionCandidate(
    kind="shorter_role", role_ids=["nimbus"],
    bullets=["Built a real-time recommendation service in Python and PyTorch for 4M users "
             "(18% higher click-through rate).",
             "Moved model training pipelines to Kubernetes (35% lower costs) and introduced "
             "code review standards for ML code."],
    removed=["mentoring 5 engineers"], reason="The role's work in two lines",
)  # fmt: skip


def test_a_shorter_role_lists_what_it_leaves_out(master_cv: MasterCV) -> None:
    vocab = [*JD.hard_skills, *master_cv.all_skills()]
    unlisted = SHORTER_NIMBUS.model_copy(update={"removed": []})
    as_long = SHORTER_NIMBUS.model_copy(
        update={"bullets": [*SHORTER_NIMBUS.bullets, "Mentored 5 engineers."]}
    )
    no_term = SHORTER_NIMBUS.model_copy(
        update={"bullets": [SHORTER_NIMBUS.bullets[0], "Cut training costs by 35%."]}
    )
    tighten = CompressionCandidate(
        kind="tighten", role_ids=["nimbus"], source_ids=["nimbus-2"],
        bullets=["Moved training to Kubernetes, cutting costs by 35%."], reason="r",
    )  # fmt: skip
    kept = check_compressions(
        [tighten, unlisted, as_long, no_term, SHORTER_NIMBUS], master_cv, ["Kubernetes"], vocab
    )
    assert [k.kind for k in kept] == ["shorter_role"]  # it takes the role's lines first
    assert kept[0].source_ids == ["nimbus-1", "nimbus-2", "nimbus-3"]
    out = apply_compression(master_cv, kept[0])
    role = out.experience[0]
    assert [b.id for b in role.bullets] == ["nimbus-1", "nimbus-2"]
    assert [b.text for b in role.bullets] == SHORTER_NIMBUS.bullets
    assert len(master_cv.experience[0].bullets) == 3  # applied to a copy only


def test_a_shorter_role_replaces_the_code_built_entry(master_cv: MasterCV) -> None:
    cv = master_cv.model_copy(deep=True)
    old = cv.experience[1]
    old.start, old.end = "2009-01", "2013-06"
    old.bullets += [
        Bullet(id="lexa-3", text="Organised the team's weekly reading group."),
        Bullet(id="lexa-4", text="Wrote internal documentation for Kubernetes jobs."),
    ]
    shorter = CompressionCandidate(
        kind="shorter_role", role_ids=["lexa"],
        bullets=["Developed NLP models for document classification with 92% accuracy using "
                 "scikit-learn and spaCy.",
                 "Designed a PostgreSQL feature store for 3 data science teams and documented "
                 "Kubernetes jobs."],
        removed=["the weekly reading group"], reason="r",
    )  # fmt: skip
    standard = apply_plan(cv, TailoringPlan(compressions=[shorter]), JD)
    assert [c.kind for c in standard.compressions] == ["condense_role"]
    optimised = apply_plan(cv, TailoringPlan(compressions=[shorter]), JD, mode="optimised")
    assert [c.kind for c in optimised.compressions] == ["shorter_role"]  # no code-built cut
    # an earlier role the job relies on is not cut for its age, nor stripped of its skills
    relied = TailoringPlan(
        requirements=[
            RequirementEvidence(
                requirement="NLP", priority="essential", importance="key", match="direct",
                evidence=["lexa-1"],
            )
        ],
        compressions=[shorter.model_copy(update={"bullets": [
            "Developed NLP models for document classification with 92% accuracy.",
            "Designed a PostgreSQL feature store for 3 data science teams and documented "
            "Kubernetes jobs.",
        ]})],
    )  # fmt: skip
    assert apply_plan(cv, relied, JD).compressions == []
    assert apply_plan(cv, relied, JD, mode="optimised").compressions == []  # lost spaCy


NIMBUS_MORE = [  # each at least `MIN_LINE` characters, as content lines in a Word CV are
    "Reviewed experiment designs with the product team before launch.",
    "Wrote design notes for the recommendation service team.",
    "Ran weekly model quality reviews with the analytics team.",
    "Documented training pipeline runbooks for new engineers.",
    "Interviewed engineering candidates for the platform team.",
    "Presented model results to leadership in monthly reviews.",
    "Maintained the team's evaluation datasets and their labels.",
]
LEXA_TWO = [
    "Developed NLP models for document classification with 92% accuracy using "
    "scikit-learn and spaCy.",
    "Designed PostgreSQL feature store used by 3 data science teams.",
]
NIMBUS_SIX = [
    "Built a real-time recommendation service in Python and PyTorch serving 4M users, "
    "lifting click-through rate by 18%.",
    "Led migration of model training pipelines to Kubernetes, cutting training costs by 35%.",
    "Mentored 5 engineers and introduced code review standards for ML code.",
    "Reviewed experiment designs with the product team and presented model results.",
    "Wrote design notes and training pipeline runbooks.",
    "Ran weekly model quality reviews on the team's evaluation datasets.",
]


NIMBUS_SOURCES = [
    ["nimbus-1"], ["nimbus-2"], ["nimbus-3"], ["nimbus-4", "nimbus-9"],
    ["nimbus-5", "nimbus-7"], ["nimbus-6", "nimbus-10"],
]  # fmt: skip


def _version_lines(texts: list[str], sources: list[list[str]] | None = None) -> list[Any]:
    from src.cv.models import VersionLine

    return [VersionLine(text=t, sources=sources[n] if sources else []) for n, t in enumerate(texts)]


def _long_roles(master: MasterCV) -> MasterCV:
    cv = master.model_copy(deep=True)
    cv.experience[0].bullets += [
        Bullet(id=f"nimbus-{n}", text=t) for n, t in enumerate(NIMBUS_MORE, start=4)
    ]
    cv.experience[1].bullets += [
        Bullet(id="lexa-3", text="Organised the team's weekly reading group."),
        Bullet(id="lexa-4", text="Wrote internal documentation for model jobs."),
        Bullet(id="lexa-5", text="Reviewed pull requests for the data team."),
    ]
    return MasterCV.model_validate(cv.model_dump())


def test_each_long_role_gets_a_shorter_version_sized_by_relevance(master_cv: MasterCV) -> None:
    from src.cv.models import RoleVersion, RoleVersions, SkillGroup, SkillsLine
    from src.cv.role_versions import SYSTEM, shorter_cv

    cv = _long_roles(master_cv)
    business_items = ["OPEX", "Budgeting", "Vendor management", "Forecasting", "Procurement"]
    cv.skills.append(SkillGroup(category="Business", items=[*business_items, "Kubernetes"]))
    cv.skills[-1].items.insert(0, "Pricing")
    cited = RequirementEvidence(
        requirement="recommendation systems", priority="essential", importance="key",
        match="direct", evidence=["nimbus-1", "nimbus-2"],
    )  # fmt: skip
    draft = apply_plan(cv, TailoringPlan(requirements=[cited]), JD, mode="optimised")
    lexa = RoleVersion(
        role_id="lexa",
        bullets=_version_lines(LEXA_TWO),
        removed=["reading group", "documentation", "pull request reviews"], reason="r",
    )  # fmt: skip
    too_short = RoleVersion(role_id="nimbus", bullets=_version_lines(NIMBUS_SIX[:2]), reason="r")
    business = SkillsLine(category="Business", keep=["Kubernetes"], reason="not lab work")
    first = RoleVersions(roles=[too_short, lexa], skills=[business])
    second = RoleVersions(
        roles=[RoleVersion(role_id="nimbus", bullets=_version_lines(NIMBUS_SIX), reason="r")]
    )
    llm = FakeLLM({RoleVersions: lambda prompt: second if "<feedback>" in prompt else first})
    vocab = [*JD.hard_skills, *cv.all_skills()]
    job = JOB_TEXT + " Kubernetes."
    found = shorter_cv(llm, draft, JD, job, ["Kubernetes"], vocab)
    assert len(llm.calls) == 2  # one more call, for the version the checks rejected
    assert "central to this job with 10 lines needs 6 to 8" in llm.calls[1][0]
    roles = [(c.role_ids, len(c.bullets)) for c in found.proposals if c.kind == "shorter_role"]
    assert roles == [(["nimbus"], 6), (["lexa"], 2)]  # cited most: central; the other: 2 or 3
    assert found.proposals[0].reason.startswith("Central to this job")
    trim = found.proposals[-1]
    assert trim.kind == "trim_skills" and trim.bullets == ["Kubernetes"]
    assert "OPEX" in trim.removed and found.rejected == []
    applied = apply_compression(draft.cv, trim)
    assert next(g for g in applied.skills if g.category == "Business").items == ["Kubernetes"]
    assert "never combine roles" in SYSTEM


def test_the_shorter_cv_rules_hold_for_any_cv_and_job(master_cv: MasterCV) -> None:
    from src.cv.role_versions import central_roles, sizes

    cv = _long_roles(master_cv)
    bare = apply_plan(cv, TailoringPlan(), JD, mode="optimised")  # no requirement map
    assert central_roles(bare, ["Kubernetes", "PyTorch"]) == {"nimbus"}  # job terms decide
    assert central_roles(bare, ["Fortran"]) == {"nimbus"}  # no evidence: the latest role
    nimbus, lexa = bare.cv.experience
    assert sizes(nimbus, True) == (6, 8) and sizes(lexa, False) == (2, 3)
    assert sizes(lexa, True) == (5, 5)  # a central role of 7 lines or fewer keeps them
    nine = nimbus.model_copy(update={"bullets": nimbus.bullets[:9]})
    assert sizes(nine, True) == (6, 7)  # always at least two fewer


def test_the_optimiser_shortens_a_different_cv_end_to_end(master_cv: MasterCV) -> None:
    from src.cv.models import RoleVersion, RoleVersions, SkillGroup, SkillsLine

    cv = _long_roles(master_cv)
    cv.skills.append(
        SkillGroup(category="Office", items=["Excel", "PowerPoint", "Word", "Outlook",
                                             "Teams", "Visio", "SharePoint"])
    )  # fmt: skip
    plan = TailoringPlan(
        requirements=[
            RequirementEvidence(
                requirement="recommendation systems", priority="essential", importance="key",
                match="direct", evidence=["nimbus-1"],
            )
        ],
        compressions=[CompressionCandidate(
            kind="tighten", role_ids=["nimbus"], source_ids=["nimbus-2"],
            bullets=["Moved model training pipelines to Kubernetes, cutting costs by 35%."],
            reason="r",
        )],
    )  # fmt: skip
    lexa = RoleVersion(
        role_id="lexa", bullets=_version_lines(LEXA_TWO),
        removed=["reading group", "documentation", "reviews"], reason="r",
    )  # fmt: skip
    nimbus = RoleVersion(
        role_id="nimbus", bullets=_version_lines(NIMBUS_SIX, NIMBUS_SOURCES), reason="r"
    )
    versions = RoleVersions(
        roles=[nimbus, lexa],
        skills=[SkillsLine(category="office", keep=["excel"], reason="not asked for")],
    )
    llm = FakeLLM({TailoringPlan: plan, RoleVersions: versions})
    out = tailor(cv, JOB_TEXT, llm, review=False, jd=JD, mode="optimised")
    kinds = [(c.kind, c.role_ids or [c.category], c.status) for c in out.compressions]
    assert kinds == [
        ("shorter_role", ["nimbus"], "applied"), ("shorter_role", ["lexa"], "applied"),
        ("trim_skills", ["Office"], "applied"),
    ]  # the draft is the short version; the plan's tightened line gave way  # fmt: skip
    nimbus_role, lexa_role = out.cv.experience
    assert [b.text for b in nimbus_role.bullets] == NIMBUS_SIX
    assert [b.id for b in nimbus_role.bullets] == [f"nimbus-{n}" for n in range(1, 7)]
    assert [b.text for b in lexa_role.bullets] == LEXA_TWO
    assert next(g for g in out.cv.skills if g.category == "Office").items == ["Excel"]
    assert out.compressions[0].before_ids == [f"nimbus-{n}" for n in range(1, 11)]
    assert out.screening is not None  # scored on the short version


def test_a_skills_line_keeps_what_the_posting_names(master_cv: MasterCV) -> None:
    from src.cv.proposals import trimmed_skills_reason

    group = master_cv.skills[0]
    assert trimmed_skills_reason(group, group.items[:1], "") is None
    named = trimmed_skills_reason(group, group.items[:1], f"We use {group.items[1]} daily.")
    assert named and named.startswith("leaves out skills the posting names")
    assert trimmed_skills_reason(group, ["Cobol"], "") == "keeps a skill that is not on this line"
    assert trimmed_skills_reason(group, group.items, "") == "leaves nothing out"


def test_a_shorter_role_is_red_below_the_role_until_chosen(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    from tests.test_cv_writer import _red_texts

    original = tmp_path / "original.docx"
    export_docx(master_cv, original)
    vocab = [*JD.hard_skills, *master_cv.all_skills()]
    item = check_compressions([SHORTER_NIMBUS], master_cv, ["Kubernetes"], vocab)[0]
    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    pending = tmp_path / "pending.docx"
    write_like_original(original, master_cv, tailored.cv, pending, compressions=[item])
    text = [p.text for p in Document(str(pending)).paragraphs]
    label = next(i for i, t in enumerate(text) if t.startswith("[CONFIRM SHORTER VERSION"))
    assert "the 2 red lines below without a label replace your 3 lines above" in text[label]
    assert text[label + 1 : label + 3] == SHORTER_NIMBUS.bullets
    assert text[label - 1].startswith("Mentored 5 engineers")  # the role's lines stay
    assert {text[label], *SHORTER_NIMBUS.bullets} <= _red_texts(pending)
    applied_item = item.model_copy(update={"status": "applied"})
    applied = tmp_path / "applied.docx"
    write_like_original(
        original, master_cv, apply_compression(tailored.cv, applied_item), applied,
        compressions=[applied_item], mark_changes=False,
    )  # fmt: skip
    text = [p.text for p in Document(str(applied)).paragraphs]
    assert all(line in text for line in SHORTER_NIMBUS.bullets)
    assert not any(t.startswith("Mentored 5") or "[CONFIRM" in t for t in text)


def test_a_shorter_skills_line_is_red_until_chosen(master_cv: MasterCV, tmp_path: Path) -> None:
    from tests.test_cv_writer import _original_cv, _red_texts

    original = _original_cv(master_cv, tmp_path / "original.docx")
    first, second = master_cv.skills[0], master_cv.skills[1]
    trim = CompressionProposal(
        id="t1", kind="trim_skills", category=first.category, role_ids=[],
        bullets=first.items[:2], removed=first.items[2:], reason="r",
    )  # fmt: skip
    drop = trim.model_copy(
        update={"id": "t2", "category": second.category, "bullets": [], "removed": second.items}
    )
    pending = tmp_path / "pending.docx"
    write_like_original(original, master_cv, master_cv, pending, compressions=[trim, drop])
    red = _red_texts(pending)
    label = "[CONFIRM SHORTER SKILLS LINE BEFORE USE: replaces the line above]"
    assert f"{label} {first.category}: {', '.join(first.items[:2])}" in red
    assert any(t.startswith("[CONFIRM BEFORE USE: remove the skills line above") for t in red)
    cv = apply_compression(apply_compression(master_cv, trim), drop)
    applied = tmp_path / "applied.docx"
    done = [c.model_copy(update={"status": "applied"}) for c in (trim, drop)]
    write_like_original(original, master_cv, cv, applied, compressions=done, mark_changes=False)
    lines = [p.text for p in Document(str(applied)).paragraphs]
    shown = next(t for t in lines if first.items[0] in t and first.items[1] in t)
    assert first.items[-1] not in shown
    assert not any(all(item in t for item in second.items) for t in lines)  # the line went
    assert not any("[CONFIRM" in t for t in lines)


def _old_merge(cv: MasterCV) -> CompressionProposal:
    """A merged entry proposed by a draft made before roles were no longer combined."""
    roles = [r for r in cv.experience if r.id in ("pd1", "pd2")]
    return CompressionProposal(
        id="m1", kind="merge_roles", role_ids=["pd1", "pd2"], heading=merged_heading(roles),
        bullets=["Developed NLP models for document classification with 92% accuracy using "
                 "scikit-learn and spaCy, and designed a PostgreSQL feature store used by 3 "
                 "data science teams."],
        reason="r",
    )  # fmt: skip


# --- Word file -----------------------------------------------------------------------------


def test_red_layer_goes_where_a_reader_expects_it(master_cv: MasterCV, tmp_path: Path) -> None:
    cv = _postdocs(master_cv)
    original = tmp_path / "original.docx"
    export_docx(cv, original)
    tailored = apply_plan(cv, TailoringPlan(), JD)
    additions = [
        ReviewAddition(
            id="e1", kind="enrichment", text="with model monitoring", anchor_id="nimbus-2",
            role_id="nimbus", requirement="model monitoring", background_quote="a b",
            job_quote="c d", reason="r", original_text=cv.bullet_index()["nimbus-2"].text,
        ),
        ReviewAddition(
            id="s1", kind="skill", text="Grafana", skill_group="Infrastructure",
            requirement="model monitoring", background_quote="a b", job_quote="c d", reason="r",
        ),
    ]  # fmt: skip
    merge = [_old_merge(cv)]  # an older draft's pending merge still displays
    pending = tmp_path / "pending.docx"
    write_like_original(
        original, cv, tailored.cv, pending, review_additions=additions, compressions=merge
    )
    text = [p.text for p in Document(str(pending)).paragraphs]
    line = next(t for t in text if t.startswith("Led migration"))
    assert line.endswith("[CONFIRM ADDED WORDS: with model monitoring]")
    assert any(t.endswith("· [CONFIRM ADDED SKILL: Grafana]") for t in text)
    assert any(t.startswith("[CONFIRM MERGED ENTRY BEFORE USE") for t in text)
    assert sum("University A" in t for t in text) == 2  # the original entry stays
    report = check_docx(pending, tailored.cv, ["Grafana", "model monitoring"])
    assert report.keyword_coverage == 0  # red words are never counted
    applied_item = merge[0].model_copy(update={"status": "applied"})
    applied_cv = apply_compression(tailored.cv, applied_item)
    applied = tmp_path / "applied.docx"
    write_like_original(original, cv, applied_cv, applied, compressions=[applied_item])
    text = [p.text for p in Document(str(applied)).paragraphs]
    assert sum("University A" in t for t in text) == 1
    assert any(t.startswith("Postdoctoral Research Fellow, University A") for t in text)
    assert not any("[CONFIRM" in t for t in text)


# --- confirming, declining, applying -------------------------------------------------------


@pytest.fixture
def ws(settings: Settings, master_cv: MasterCV) -> Workspace:
    settings.master_cv_path.write_text(master_cv.model_dump_json())
    return Workspace(settings)


def _draft(ws: Workspace, master: MasterCV, tailored: TailoredCV) -> str:
    path = ws.output_dir / "cvs" / "draft.docx"
    document = tailored_documents.create(
        ws, JOB.id, master, JD, tailored, "original", path, posting=JOB,
        settings=TailoringSettings(mode="optimised"), reviewed=False,
    )  # fmt: skip
    return document.id


def test_confirmed_words_and_numbers_reword_the_source_line(
    ws: Workspace, master_cv: MasterCV
) -> None:
    cv = master_cv.model_copy(deep=True)
    cv.experience[0].bullets[2].text = "Introduced code review standards for ML code."
    ws.master_cv = cv
    ws.settings.master_cv_path.write_text(cv.model_dump_json())
    tailored = apply_plan(cv, TailoringPlan(), JD)
    tailored.review_additions = check_additions(
        [
            _proposal(),
            _proposal(
                kind="metric", text="how many reviews a month?", anchor_id="nimbus-3",
                background_quote="code review standards",
            ),
        ],
        _ctx(cv),
    )  # fmt: skip
    document_id = _draft(ws, cv, tailored)
    enrich, metric = tailored.review_additions
    done = tailored_documents.review_addition(ws, document_id, enrich.id, "confirm")
    line = done.tailored.cv.bullet_index()["nimbus-2"].text
    assert line.endswith("35%, with model monitoring dashboards.")
    assert ws.master_cv is not None and ws.master_cv.bullet_index()["nimbus-2"].text == line
    assert tailored_documents.as_tailored(done).bullet_index()["nimbus-2"].text.endswith("35%.")
    with pytest.raises(ValueError, match="with its number"):
        tailored_documents.review_addition(ws, document_id, metric.id, "confirm")
    numbered = "Introduced code review standards for ML code, across 40 reviews a month."
    done = tailored_documents.review_addition(ws, document_id, metric.id, "confirm", numbered)
    assert done.tailored.cv.bullet_index()["nimbus-3"].text == numbered
    assert done.reviewed


def test_a_deleted_proposal_is_remembered_for_that_cv(ws: Workspace, master_cv: MasterCV) -> None:
    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    tailored.review_additions = check_additions([_proposal()], _ctx(master_cv))
    document_id = _draft(ws, master_cv, tailored)
    item = tailored.review_additions[0]
    tailored_documents.review_addition(ws, document_id, item.id, "delete", reason="never did it")
    cv_id = tailored_documents.load(ws, document_id).cv_id
    assert declined.texts(ws, cv_id) == [item.text]
    assert declined.load(ws)[0].reason == "never did it"
    assert declined.texts(ws, "another-cv") == []


def test_only_the_user_applies_a_compressed_version(ws: Workspace, master_cv: MasterCV) -> None:
    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    tailored.compressions = check_compressions(
        [CompressionCandidate(
            kind="merge_bullets", role_ids=["nimbus"], source_ids=["nimbus-2", "nimbus-3"],
            bullets=["Led migration of training pipelines to Kubernetes (35% lower costs) and "
                     "mentored 5 engineers on ML code review."],
            reason="r",
        )], master_cv, ["Kubernetes"], [*master_cv.all_skills()],
    )  # fmt: skip
    document_id = _draft(ws, master_cv, tailored)
    saved = tailored_documents.load(ws, document_id)
    assert len(saved.tailored.cv.experience[0].bullets) == 3  # nothing removed by the app
    item = tailored.compressions[0]
    done = tailored_documents.review_compression(ws, document_id, item.id, "apply")
    assert [b.id for b in done.tailored.cv.experience[0].bullets] == ["nimbus-1", "nimbus-2"]
    assert done.tailored.compressions[0].status == "applied"
    assert len(done.source_cv.experience[0].bullets) == 3  # the source CV keeps its lines
    with pytest.raises(ValueError, match="Pending shorter version"):
        tailored_documents.review_compression(ws, document_id, item.id, "reject")


def _gap_and_shorter(ws: Workspace, master_cv: MasterCV) -> tuple[str, str, str]:
    """A draft with a gap filler on nimbus-2 and a shorter version of nimbus."""
    cv = _long_roles(master_cv)
    ws.master_cv = cv
    ws.settings.master_cv_path.write_text(cv.model_dump_json())
    tailored = apply_plan(cv, TailoringPlan(), JD, mode="optimised")
    tailored.review_additions = check_additions([_proposal()], _ctx(cv))
    shorter = CompressionCandidate(
        kind="shorter_role", role_ids=["nimbus"], bullets=NIMBUS_SIX,
        line_sources=NIMBUS_SOURCES, reason="r",
    )  # fmt: skip
    vocab = [*JD.hard_skills, *cv.all_skills()]
    tailored.compressions = check_compressions([shorter], cv, ["Kubernetes"], vocab)
    assert tailored.compressions  # the version passed its checks
    document_id = _draft(ws, cv, tailored)
    return document_id, tailored.review_additions[0].id, tailored.compressions[0].id


def test_a_confirmed_gap_filler_survives_a_shorter_role(ws: Workspace, master_cv: MasterCV) -> None:
    document_id, addition, version = _gap_and_shorter(ws, master_cv)
    tailored_documents.review_addition(ws, document_id, addition, "confirm")
    done = tailored_documents.review_compression(ws, document_id, version, "apply")
    role = done.tailored.cv.experience[0]
    assert [b.text for b in role.bullets[:6]] == NIMBUS_SIX
    assert len(role.bullets) == 7  # the confirmed line stays after the shorter version
    assert role.bullets[-1].text.endswith("35%, with model monitoring dashboards.")


def test_a_gap_filler_confirmed_after_a_shorter_role_is_added_to_it(
    ws: Workspace, master_cv: MasterCV
) -> None:
    document_id, addition, version = _gap_and_shorter(ws, master_cv)
    tailored_documents.review_compression(ws, document_id, version, "apply")
    done = tailored_documents.review_addition(ws, document_id, addition, "confirm")
    role = done.tailored.cv.experience[0]
    assert len(role.bullets) == 6  # the words join the shorter line that carries nimbus-2
    assert role.bullets[1].text == enriched(NIMBUS_SIX[1], "with model monitoring dashboards")
    assert [b.text for b in role.bullets[2:]] == NIMBUS_SIX[2:]  # the rest untouched
    assert ws.master_cv is not None  # the fact also went into the CV it came from
    assert ws.master_cv.bullet_index()["nimbus-2"].text.endswith("model monitoring dashboards.")
    word = [p.text for p in Document(str(ws.output_dir / "cvs" / done.filename)).paragraphs]
    assert all(b.text in word for b in role.bullets)
    back = tailored_documents.review_compression(ws, document_id, version, "restore")
    lines = back.tailored.cv.experience[0].bullets
    assert len(lines) == 10 and lines[1].text.endswith("model monitoring dashboards.")
    assert back.tailored.compressions[0].status == "pending"  # can be used again


# --- the pass toward the target score ------------------------------------------------------


def test_optimised_mode_proposes_toward_grade_a_scored_by_code(master_cv: MasterCV) -> None:
    jd = JDAnalysis(
        job_title="Senior Machine Learning Engineer",
        hard_skills=["Python", "PyTorch", "model monitoring", "on-call", "Airflow", "Spark"],
    )
    job = "Senior Machine Learning Engineer. " + ", ".join(jd.hard_skills) + "."
    plan = TailoringPlan(
        requirements=[
            RequirementEvidence(
                requirement="model monitoring", priority="essential", importance="key",
                match="not_evidenced",
            )
        ],
        review_candidates=[_proposal(job_quote="model monitoring, on-call")],
    )  # fmt: skip
    second = _proposal(
        kind="experience", text="Took on-call shifts for production recommendation models",
        anchor_id="nimbus-1", background_quote="real-time recommendation service",
        job_quote="model monitoring, on-call",
    )  # fmt: skip
    revised = plan.model_copy(update={"review_candidates": [second]})
    llm = FakeLLM({TailoringPlan: lambda prompt: revised if "<reading>" in prompt else plan})
    readings: list[TailoredCV] = []

    def check(draft: TailoredCV) -> TailoredCV:
        readings.append(draft)
        return draft

    out = tailor(master_cv, job, llm, jd=jd, mode="optimised", check=check, declined=["x"])
    assert len(readings) == 1  # one AI reading, at the end; the revision is scored by code
    assert [m for _, m in llm.calls] == [TailoringPlan, TailoringPlan]  # plan + one revision
    assert sorted(a.kind for a in out.review_additions) == ["enrichment", "experience"]
    sc = out.screening
    assert sc is not None and sc.if_confirmed is not None
    assert sc.before.score <= sc.after.score < sc.if_confirmed.score
    found = {c.term for c in sc.if_confirmed.keywords if c.found}
    assert {"model monitoring", "on-call"} <= found
    assert all("<declined>" in p for p, _ in llm.calls)


# --- screening match: literal keywords, one baseline, rewordings ----------------------------


def test_keywords_are_the_posting_s_own_words() -> None:
    text = (
        "You will hold a PhD in life sciences, with experience of grant management and funding "
        "management, KPIs and evaluation frameworks, and risk management."
    )
    jd = JDAnalysis(
        job_title="Principal Scientist",
        hard_skills=[
            "Grant and funding management", "KPIs and evaluation frameworks",
            "Life sciences research (PhD)", "Risk management", "Python",
        ],
        must_have=["Experience of managing funded projects from initiation to delivery."],
    )  # fmt: skip
    out = literal_terms(jd, text)
    assert out.hard_skills == [
        "Grant management", "funding management", "KPIs and evaluation frameworks", "PhD",
        "Risk management",
    ]  # fmt: skip
    assert out.must_have == jd.must_have  # requirement sentences stay as written


def test_rejected_rewrites_become_red_rewordings(master_cv: MasterCV) -> None:
    jd = JDAnalysis(job_title="ML Engineer", hard_skills=["Python", "Airflow"])
    text = "Led migration of Airflow model training pipelines to Kubernetes, cutting costs by 35%."
    plan = TailoringPlan(rewritten_bullets=[RewrittenBullet(source_id="nimbus-2", text=text)])
    out = apply_plan(master_cv, plan, jd, mode="optimised")
    assert not out.changes[0].accepted
    candidates = rewordings(out.changes, out.requirements)
    assert [(c.kind, c.anchor_id) for c in candidates] == [("rewording", "nimbus-2")]
    ctx = _ctx(master_cv, keywords=frozenset({"Airflow"}))
    kept = check_additions(candidates, ctx)
    assert kept and kept[0].original_text.startswith("Led migration")
    invented = RewrittenBullet(source_id="nimbus-2", text=text.replace("35%", "50%"))
    lost = apply_plan(master_cv, TailoringPlan(rewritten_bullets=[invented]), jd, mode="optimised")
    assert rewordings(lost.changes, lost.requirements) == []  # numbers are never the user's call


def test_every_draft_of_a_job_is_set_against_one_reading_of_the_original(
    master_cv: MasterCV,
) -> None:
    def reader(prompt: str) -> EvidencedReadings:
        line = "Built a real-time recommendation service in Python and PyTorch serving 4M users"
        reading = EvidencedReading(
            **{n: RatedLine(rating=2, quote=line) for n in WEIGHTS},
            requirements_met=[MetRequirement(number=1, quote=line)],
        )
        return EvidencedReadings(readings=[reading] * prompt.count("</cv_"))

    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    llm = FakeLLM({EvidencedReadings: reader})
    first = compare_fit(master_cv, tailored, JD, JOB_TEXT, llm)
    assert first.baseline is not None and len(llm.calls) == 1
    second_calls = len(llm.calls)
    stored = first.baseline.model_copy(update={"requirements": ["Python", "PyTorch"]})
    second = compare_fit(master_cv, tailored, JD, JOB_TEXT, llm, original=stored)
    assert second.baseline == stored and second.requirements_total == 2  # same baseline
    assert len(llm.calls) == second_calls + 1
    other = master_cv.model_copy(deep=True)
    other.basics.summary = "Another CV."
    third = compare_fit(other, tailored, JD, JOB_TEXT, llm, original=stored)
    assert third.baseline is not None and third.baseline.text_key != stored.text_key


def test_a_rewording_is_red_below_its_line_and_confirmed_in_the_draft_only(
    ws: Workspace, master_cv: MasterCV, tmp_path: Path
) -> None:
    item = ReviewAddition(
        id="r1", kind="rewording", anchor_id="nimbus-3", role_id="nimbus",
        text="Mentored 5 engineers and introduced ML code review standards as the team's "
        "technical lead.", requirement="technical leadership", background_quote="Mentored 5",
        job_quote="", reason="r", original_text=master_cv.bullet_index()["nimbus-3"].text,
    )  # fmt: skip
    original = tmp_path / "original.docx"
    export_docx(master_cv, original)
    target = tmp_path / "draft.docx"
    write_like_original(original, master_cv, master_cv, target, review_additions=[item])
    text = [p.text for p in Document(str(target)).paragraphs]
    at = text.index(item.original_text)
    assert text[at + 1].startswith("[CONFIRM REWORDED LINE BEFORE USE")
    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    tailored.review_additions = [item]
    document_id = _draft(ws, master_cv, tailored)
    done = tailored_documents.review_addition(ws, document_id, "r1", "confirm")
    assert done.tailored.cv.bullet_index()["nimbus-3"].text == item.text
    assert done.source_cv.bullet_index()["nimbus-3"].text == item.original_text


# --- the deterministic screening score -----------------------------------------------------


def test_the_screening_score_is_deterministic_and_matches_as_an_ats_search() -> None:
    jd = JDAnalysis(
        job_title="Principal Scientist, Preclinical External Innovation",
        hard_skills=["post-doctoral research", "life science", "rare diseases", "KPIs", "Spark"],
        soft_skills=["stakeholder management"],
    )
    posting = (
        "Principal Scientist. PhD in a life science. Post-doctoral research in rare diseases. "
        "Use KPIs; KPIs matter. Spark. Stakeholder management."
    )
    cv = (
        "Principal Scientist with a PhD. Postdoctoral research in Life Sciences on rare disease "
        "models; tracked KPI dashboards."
    )
    first = score_cv(cv, "Principal Scientist", jd, posting, jd.hard_skills)
    assert first == score_cv(cv, "Principal Scientist", jd, posting, jd.hard_skills)
    found = {c.term: c.found for c in first.keywords}
    assert found == {
        "post-doctoral research": True, "life science": True, "rare diseases": True,
        "KPIs": True, "Spark": False,
    }  # fmt: skip
    assert next(c.weight for c in first.keywords if c.term == "KPIs") == 2  # used twice
    assert first.title_match == 1 and [e.found for e in first.education] == [True]
    assert first.soft_skills[0].found is False
    # 65 keywords (5 of 6 weight), 15 title, 10 degree, 10 soft skills (none found)
    assert first.score == round(65 * 5 / 6 + 15 + 10) == 79 and first.grade == "B"
    assert grade(79) == "B" and grade(50) == "C" and grade(49) == "D"
    assert not term_found("start the race", "art") and term_found("A/B testing", "A B testing")


def test_a_draft_saved_before_the_score_existed_is_scored_when_shown(
    ws: Workspace, master_cv: MasterCV
) -> None:
    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    assert tailored.screening is None  # as drafts saved before the score
    document = tailored_documents.load(ws, _draft(ws, master_cv, tailored))
    scored = tailored_documents.with_screening(document)
    assert scored.screening is not None
    assert scored.screening == tailored_documents.with_screening(document).screening


def test_the_final_version_leaves_out_what_is_pending(ws: Workspace, master_cv: MasterCV) -> None:
    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    tailored.headline_options = ["Staff ML Engineer", "ML Engineer"]
    tailored.review_additions = check_additions([_proposal()], _ctx(master_cv))
    document_id = _draft(ws, master_cv, tailored)
    final = tailored_documents.finalise(ws, document_id)
    assert final.final_file and final.final_file.endswith("_final.docx")
    path = ws.output_dir / "cvs" / final.final_file
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)
    assert "CONFIRM" not in text and "Staff ML Engineer" not in text
    assert any("1 suggestion(s) still pending" in n for n in final.final_notes)
    edited = tailored_documents.edit(
        ws, document_id, TailoredCVEdits(summary=master_cv.basics.summary + " I build ML.")
    )
    assert edited.final_file is None  # the draft changed: write the final version again


def test_a_role_heading_above_the_employer_line_is_never_taken_or_removed(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    from src.cv.docx_original import align_to_document

    original = tmp_path / "original.docx"
    doc = Document()
    doc.add_paragraph("EXPERIENCE")
    for role in master_cv.experience:
        doc.add_paragraph(f"{role.title}\tMar 2017 – Feb 2021")  # title and dates first
        doc.add_paragraph(f"{role.company}, Lisbon, Portugal")
        for bullet in role.bullets:
            doc.add_paragraph(bullet.text, style="List Bullet")
    doc.save(str(original))
    aligned = align_to_document(original, master_cv)
    assert all("\t" not in b.text for r in aligned.experience for b in r.bullets)
    tailored = aligned.model_copy(deep=True)
    tailored.experience[0].bullets = tailored.experience[0].bullets[:2]  # one line left out
    out = tmp_path / "out.docx"
    write_like_original(original, aligned, tailored, out)
    text = [p.text for p in Document(str(out)).paragraphs]
    assert sum("\tMar 2017 – Feb 2021" in t for t in text) == len(master_cv.experience)


def test_the_assistant_sees_and_decides_the_same_red_layer(
    ws: Workspace, master_cv: MasterCV
) -> None:
    import asyncio

    from src.agents import actions_more
    from src.agents.definitions import ASSISTANT
    from src.agents.runtime import AgentContext

    async def _ignore(_event: str, _data: dict[str, object]) -> None:
        return None

    document_id, addition, version = _gap_and_shorter(ws, master_cv)
    ctx = AgentContext(ws=ws, emit=_ignore)
    read = asyncio.run(
        actions_more.read_tailored_cv(actions_more.TailoredRead(document_id=document_id), ctx)
    )
    assert [g["id"] for g in read["gap_fillers"]] == [addition]
    shorter = read["shorter_versions"][0]
    assert shorter["id"] == version and shorter["kind"] == "shorter_role"
    assert (shorter["lines_before"], shorter["lines_after"]) == (10, 6)
    done = asyncio.run(
        actions_more.review_cv_compression(
            actions_more.CVCompressionDecision(
                document_id=document_id, compression_id=version, decision="apply"
            ),
            ctx,
        )
    )
    assert [v["status"] for v in done["still_to_review"]["shorter_versions"]] == ["applied"]
    assert [g["id"] for g in done["still_to_review"]["gap_fillers"]] == [addition]
    assert {"read_tailored_cv", "review_cv_compression"} <= set(ASSISTANT.allowed_tools)


def test_the_short_version_is_written_red_and_labelled_then_clean(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    from src.cv.models import SkillGroup
    from src.cv.proposals import applied

    cv = _long_roles(master_cv)
    cv.skills.append(SkillGroup(category="Office", items=["Excel", "Visio"]))
    original = tmp_path / "original.docx"
    export_docx(cv, original)
    vocab = [*JD.hard_skills, *cv.all_skills()]
    shorter = CompressionCandidate(
        kind="shorter_role", role_ids=["nimbus"], bullets=NIMBUS_SIX,
        line_sources=NIMBUS_SOURCES, reason="r",
    )  # fmt: skip
    version = check_compressions([shorter], cv, ["Kubernetes"], vocab)[0]
    office = CompressionProposal(
        id="o1", kind="trim_skills", category="Office", role_ids=[], bullets=[],
        removed=["Excel", "Visio"], reason="r",
    )  # fmt: skip
    draft, version = applied(cv, version)
    draft, office = applied(draft, office)
    gap = check_additions([_proposal()], _ctx(cv))  # words for nimbus-2, now in line 2
    out = tmp_path / "draft.docx"
    write_like_original(
        original, cv, draft, out, review_additions=gap, compressions=[version, office]
    )
    paragraphs = Document(str(out)).paragraphs
    text = [p.text for p in paragraphs]
    label = (
        "[CONFIRM SHORTER VERSION OF THIS ROLE BEFORE USE: the 6 red lines below without a "
        "label replace your 10"
    )
    at = next(i for i, t in enumerate(text) if t.startswith(label))
    # The user's lines stay above the label and the shorter version follows it, so
    # deleting everything above the label leaves the short version.
    assert any(t.startswith("Mentored 5") for t in text[:at])
    assert [t.split(" [CONFIRM")[0] for t in text[at + 1 : at + 7]] == NIMBUS_SIX
    assert "List" not in paragraphs[at].style.name  # a label, not an empty bullet
    from tests.test_cv_writer import _red_texts

    assert {text[at], *text[at + 1 : at + 7]} <= _red_texts(out)
    assert any(t.startswith("[CONFIRM REMOVAL: not related to this job]") for t in text)
    assert text[at + 2].endswith(
        "[CONFIRM ADDED WORDS: with model monitoring dashboards]"
    )  # on line 2
    final = tmp_path / "final.docx"
    write_like_original(
        original, cv, draft, final, compressions=[version, office], mark_changes=False
    )
    text = [p.text for p in Document(str(final)).paragraphs]
    assert all(line in text for line in NIMBUS_SIX)
    assert not any("[CONFIRM" in t or "Visio" in t for t in text)


def test_a_removed_skills_line_is_found_when_the_file_lists_one_more_skill(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    """The Word design may list a skill the stored CV lacks; the line is still removed."""
    from src.cv.models import SkillGroup
    from src.cv.proposals import applied

    in_file = master_cv.model_copy(deep=True)
    in_file.skills.append(SkillGroup(category="Office", items=["Excel", "Visio", "Access"]))
    original = tmp_path / "original.docx"
    export_docx(in_file, original)
    cv = master_cv.model_copy(deep=True)
    cv.skills.append(SkillGroup(category="Office", items=["Excel", "Visio"]))
    office = CompressionProposal(
        id="o1", kind="trim_skills", category="Office", role_ids=[], bullets=[],
        removed=["Excel", "Visio"], reason="r",
    )  # fmt: skip
    draft, office = applied(cv, office)
    out = tmp_path / "draft.docx"
    notes = write_like_original(original, cv, draft, out, compressions=[office])
    text = [p.text for p in Document(str(out)).paragraphs]
    assert any(t.startswith("[CONFIRM REMOVAL: not related to this job] Office") for t in text)
    assert not any("Office kept" in n for n in notes)


def test_a_shorter_role_shows_the_users_own_lines_and_one_rewording_per_line(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    """Above the label, the role as the user wrote it; below, the shorter lines and each
    suggestion labelled with what it does (user's report 2026-10-21)."""
    from src.cv.models import ReviewAddition
    from src.cv.proposals import applied

    cv = _long_roles(master_cv)
    original = tmp_path / "original.docx"
    export_docx(cv, original)
    vocab = [*JD.hard_skills, *cv.all_skills()]
    shorter = CompressionCandidate(
        kind="shorter_role", role_ids=["nimbus"], bullets=NIMBUS_SIX,
        line_sources=NIMBUS_SOURCES, reason="r",
    )  # fmt: skip
    version = check_compressions([shorter], cv, ["Kubernetes"], vocab)[0]
    rewritten = apply_plan(cv, TailoringPlan(), JD).cv
    rewritten.experience[0].bullets[0].text = "A plan rewrite of the first line."
    draft, version = applied(rewritten, version)
    line = version.line_sources[0][0]
    rewordings = [
        ReviewAddition(
            id=f"r{n}",
            kind="rewording",
            text=f"Reworded {n}",
            anchor_id=line,
            role_id="nimbus",
            requirement="r",
            background_quote="b",
            job_quote="j",
            reason="r",
        )  # fmt: skip
        for n in (1, 2)
    ]
    out = tmp_path / "draft.docx"
    write_like_original(
        original, cv, draft, out, review_additions=rewordings, compressions=[version]
    )
    text = [p.text for p in Document(str(out)).paragraphs]
    at = next(i for i, t in enumerate(text) if t.startswith("[CONFIRM SHORTER VERSION"))
    own = [b.text for b in cv.experience[0].bullets]
    assert all(t in text[:at] for t in own)  # the user's lines, as written
    assert "A plan rewrite of the first line." not in text
    reworded = [t for t in text if "Reworded" in t]
    assert reworded == [
        "[CONFIRM REWORDED LINE BEFORE USE: replaces the shorter line above] Reworded 1"
    ]
