"""Tailoring written to score higher with ATS and AI screens, without new facts."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from src.cv.acronyms import known_pairs, pairs_in
from src.cv.ats import layout_warnings
from src.cv.ats_score import score_cv
from src.cv.fit_check import ats_match
from src.cv.models import (
    FitComparison,
    JDAnalysis,
    MasterCV,
    RequirementEvidence,
    RewrittenBullet,
    TailoredCV,
    TailoringPlan,
    TermBridge,
)
from src.cv.tailor import apply_plan, printed_text, tailor
from src.cv.terms import bridgeable
from tests.conftest import FakeLLM

TERM = "Personalisation systems"
JD = JDAnalysis(
    job_title="Senior ML Engineer",
    hard_skills=["Python", TERM, "AWS SageMaker"],
    must_have=["Personalisation systems"],
)
REQUIREMENT = "Build personalisation systems for millions of users"
REWRITE = (
    "Built a real-time recommendation service for personalisation systems in Python and "
    "PyTorch serving 4M users, lifting click-through rate by 18%."
)


def _plan(match: str = "direct", quote: str = "real-time recommendation service") -> TailoringPlan:
    return TailoringPlan(
        requirements=[
            RequirementEvidence(requirement=REQUIREMENT, match=match, evidence=["nimbus-1"])  # type: ignore[arg-type]
        ],
        term_bridges=[
            TermBridge(term=TERM, source_id="nimbus-1", quote=quote, requirement=REQUIREMENT)
        ],
        rewritten_bullets=[RewrittenBullet(source_id="nimbus-1", text=REWRITE)],
    )


def test_a_verified_job_term_may_name_what_the_bullet_shows(master_cv: MasterCV) -> None:
    out = apply_plan(master_cv, _plan(), JD)
    assert [b.term for b in out.term_bridges] == [TERM]
    assert out.changes[0].accepted
    assert TERM in out.matched_keywords
    assert out.cv.bullet_index()["nimbus-1"].text == REWRITE


def test_without_a_pairing_the_job_term_is_still_rejected(master_cv: MasterCV) -> None:
    plan = _plan().model_copy(update={"term_bridges": []})
    out = apply_plan(master_cv, plan, JD)
    assert not out.changes[0].accepted
    assert out.cv.bullet_index()["nimbus-1"].text == master_cv.bullet_index()["nimbus-1"].text


def test_standard_pairs_only_direct_evidence_optimised_also_transferable(
    master_cv: MasterCV,
) -> None:
    standard = apply_plan(master_cv, _plan("transferable"), JD)
    assert standard.term_bridges == [] and not standard.changes[0].accepted
    optimised = apply_plan(master_cv, _plan("transferable"), JD, mode="optimised")
    assert [b.term for b in optimised.term_bridges] == [TERM]
    assert optimised.changes[0].accepted
    adjacent = apply_plan(master_cv, _plan("adjacent"), JD, mode="optimised")
    assert adjacent.term_bridges == []


def test_a_pairing_needs_the_bullets_own_words(master_cv: MasterCV) -> None:
    out = apply_plan(master_cv, _plan(quote="large-scale personalisation platform"), JD)
    assert out.term_bridges == [] and not out.changes[0].accepted


def test_tools_acronyms_and_regulated_terms_need_the_cvs_own_word() -> None:
    assert bridgeable("Mammalian cell culture") and bridgeable("Drug discovery")
    for term in ("AWS SageMaker", "iPSC culture", "LC-MS", "Python", "GMP manufacturing",
                 "Project management", "5 years of research"):  # fmt: skip
        assert not bridgeable(term), term


def test_a_verified_term_no_line_prints_goes_on_the_relevant_skills_line(
    master_cv: MasterCV,
) -> None:
    out = apply_plan(master_cv, _plan().model_copy(update={"rewritten_bullets": []}), JD)
    assert TERM in out.role_skills
    assert f"Relevant skills for this role: {TERM} · Python" in printed_text(out)


def _reading(before: int, after: int, ats: tuple[int, int] = (40, 40)) -> FitComparison:
    return FitComparison(
        before=before, after=after, ats_before=ats[0], ats_after=ats[1],
        requirements_total=2, requirements_before=1, requirements_after=1,
        requirements_missing=[REQUIREMENT], impression_after="generic ML engineer",
    )  # fmt: skip


def _checker(*readings: FitComparison) -> tuple[list[TailoredCV], object]:
    seen: list[TailoredCV] = []
    queue = list(readings)

    def check(draft: TailoredCV) -> TailoredCV:
        seen.append(draft)
        return draft.model_copy(update={"fit_check": queue.pop(0)})

    return seen, check


JOB_TEXT = f"Senior ML Engineer. Python, {TERM} and AWS SageMaker for millions of users."


def _revising(first: TailoringPlan, second: TailoringPlan) -> FakeLLM:
    """The first plan, then (asked with a <reading>) the revised one."""
    return FakeLLM({TailoringPlan: lambda p: second if "<reading>" in p else first})


def test_a_draft_under_grade_a_is_revised_once_and_the_better_kept(master_cv: MasterCV) -> None:
    llm = _revising(TailoringPlan(), _plan())
    seen, check = _checker(_reading(55, 55))
    out = tailor(master_cv, JOB_TEXT, llm, jd=JD, check=check)  # type: ignore[arg-type]
    assert len(seen) == 1 and len(llm.calls) == 2  # one AI reading, at the end
    assert "<reading>" in llm.calls[1][0] and TERM in llm.calls[1][0]
    assert out.screening is not None and TERM in out.matched_keywords
    assert out.screening.after.score > out.screening.before.score
    assert out.revision.startswith("Revised once (screening score") and "kept." in out.revision


def test_the_first_draft_is_kept_when_the_revision_scores_no_better(master_cv: MasterCV) -> None:
    llm = _revising(_plan(), TailoringPlan())
    _, check = _checker(_reading(55, 57))
    out = tailor(master_cv, JOB_TEXT, llm, jd=JD, check=check)  # type: ignore[arg-type]
    assert TERM in out.matched_keywords
    assert "first draft scored as well or better" in out.revision


def test_no_revision_once_the_screening_score_is_grade_a(master_cv: MasterCV) -> None:
    jd = JDAnalysis(job_title="Senior Machine Learning Engineer", hard_skills=["Python"])
    llm = FakeLLM({TailoringPlan: lambda _p: TailoringPlan()})
    seen, check = _checker(_reading(60, 60))
    out = tailor(master_cv, "Senior Machine Learning Engineer, Python", llm, jd=jd, check=check)  # type: ignore[arg-type]
    assert out.screening is not None and out.screening.after.grade == "A"
    assert len(seen) == 1 and len(llm.calls) == 1 and out.revision == ""


def test_ats_match_weighs_keywords_and_requirements_shown() -> None:
    assert ats_match(0.5, 2, 4) == 50
    assert ats_match(1.0, 3, 4) == 88
    assert ats_match(0.4, 0, 0) == 40  # no requirements read: keywords alone


# --- a term and its abbreviation (user's decision 2026-10-15) -------------------------------


def test_abbreviations_pair_only_with_matching_initials() -> None:
    text = (
        "Experience with liquid chromatography–mass spectrometry (LC-MS) and Key Performance "
        "Indicators (KPIs). Uses AWS (Amazon Web Services). Python (3.11). Doctor of "
        "Philosophy (PhD)."
    )
    assert pairs_in(text) == [
        ("liquid chromatography–mass spectrometry", "LC-MS"),
        ("Key Performance Indicators", "KPIs"),
        ("Amazon Web Services", "AWS"),
    ]
    assert ("machine learning", "ML") in known_pairs()


def test_the_score_credits_the_form_the_posting_pairs_with_a_term() -> None:
    jd = JDAnalysis(job_title="Scientist", hard_skills=["LC-MS", "machine learning"])
    posting = "Scientist with liquid chromatography-mass spectrometry (LC-MS) and machine learning."
    cv = "Ran liquid chromatography mass spectrometry assays; built ML models."
    found = {c.term: c.found for c in score_cv(cv, "", jd, posting, jd.hard_skills).keywords}
    # the posting pairs LC-MS with its long form; "ML" alone is not what it searches for
    assert found == {"LC-MS": True, "machine learning": False}


def test_a_missing_job_term_is_written_in_both_forms_once(master_cv: MasterCV) -> None:
    cv = master_cv.model_copy(deep=True)
    cv.experience[0].bullets[1].text = "Tracked KPIs weekly and prepared 10 ml buffer stocks."
    jd = JDAnalysis(
        job_title="Data Scientist",
        hard_skills=[
            "natural language processing",
            "key performance indicators",
            "machine learning",
        ],
    )
    posting = "You will report key performance indicators (KPIs) and use NLP."
    out = apply_plan(cv, TailoringPlan(), jd, jd_text=posting)
    bullets = out.cv.bullet_index()
    assert bullets["lexa-1"].text.startswith("Developed natural language processing (NLP) models")
    assert bullets["nimbus-2"].text.startswith("Tracked key performance indicators (KPIs) weekly")
    assert "10 ml buffer" in bullets["nimbus-2"].text  # millilitres are not "ML"
    assert sum("natural language processing" in b.text for b in bullets.values()) == 1
    assert {c.source_id for c in out.changes if c.accepted} >= {"lexa-1", "nimbus-2"}
    assert "natural language processing" in out.matched_keywords
    again = apply_plan(out.cv, TailoringPlan(), jd, jd_text=posting)
    assert again.cv.bullet_index()["lexa-1"].text == bullets["lexa-1"].text  # once only


def test_read_back_flags_dates_parsers_may_misread(master_cv: MasterCV, tmp_path: Path) -> None:
    path = tmp_path / "cv.docx"
    doc = Document()
    for line in ["Experience", "Nimbus Analytics, Senior ML Engineer, 03/2021 to date",
                 "Lexa Labs, ML Engineer, Sep 2017 – Feb 2021", "Education", "Skills"]:  # fmt: skip
        doc.add_paragraph(line)
    doc.save(str(path))
    warnings = " | ".join(layout_warnings(path, master_cv))
    assert "Dates a parser may misread: Nimbus Analytics (" in warnings
    assert "Lexa Labs (" not in warnings and "No standard section heading" not in warnings
    doc.paragraphs[-1].text = "Tools"
    doc.save(str(path))
    assert "heading for: skills" in " | ".join(layout_warnings(path, master_cv))
