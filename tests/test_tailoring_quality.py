"""Evidence-led tailoring: the requirement map, positioning, level and academic steering, the
ownership/scope guard, the "Relevant skills for this role" line, academic section order, the
before/after fit check, letter bridge and review, and profile claim checks."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document

from src.cv.ats import check_docx
from src.cv.claims import escalation
from src.cv.cover_letter import apply_letter, letter_context, write_letter
from src.cv.docx_exporter import export_docx
from src.cv.docx_original import academic_order, write_like_original
from src.cv.fit_check import compare_fit
from src.cv.models import (
    AdditionCandidates,
    CoverLetterDraft,
    CVCritique,
    Education,
    EvidencedReading,
    EvidencedReadings,
    JDAnalysis,
    LetterOpening,
    LetterParagraph,
    LetterReview,
    LetterSections,
    MasterCV,
    MetRequirement,
    Positioning,
    RatedLine,
    RequirementEvidence,
    RewrittenBullet,
    TailoringPlan,
)
from src.cv.tailor import (
    ANALYZE_SYSTEM,
    PLAN_SYSTEM,
    apply_plan,
    check_requirements,
    level_note,
    printed_text,
    propose_review_additions,
    tailor,
)
from src.jobs.models import ExpansionRoute, ProfileSummary, RoleFamily
from src.jobs.profile_memory import ProfileMemory, check_claims, profile_brief
from src.jobs.screener import matcher_profile
from src.services import cv_service
from src.services.workspace import Workspace
from tests.conftest import FakeLLM

JD = JDAnalysis(
    job_title="Senior ML Engineer",
    company="Orbit AI",
    seniority="senior",
    hard_skills=["Python", "PyTorch", "Azure"],
    soft_skills=["stakeholder management"],
    must_have=["MLOps"],
)


def _plan(**kw: object) -> TailoringPlan:
    return TailoringPlan.model_validate(kw)


# ---------------------------------------------------------------- ownership and scope


@pytest.mark.parametrize(
    ("source", "new", "flagged"),
    [
        ("Mentored 5 engineers.", "Managed 5 engineers.", True),  # mentoring is not managing
        ("Contributed to a product launch.", "Led a product launch.", True),
        ("Built a feature store.", "Built a global feature store.", True),
        ("Led migration to Kubernetes.", "Drove the migration to Kubernetes.", False),
        ("Responsible for project management of trials.", "Managed trials.", False),
        ("Set the team's research strategy.", "Set the team's strategic research plan.", False),
        ("First author on 4 papers.", "Lead author on 4 papers.", False),
    ],
)
def test_escalation_keeps_ownership_and_scope_to_the_source(
    source: str, new: str, flagged: bool
) -> None:
    assert (escalation(source, new) is not None) is flagged


def test_rewrite_that_inflates_mentoring_into_management_is_rejected(master_cv: MasterCV) -> None:
    plan = _plan(
        rewritten_bullets=[
            {
                "source_id": "nimbus-3",
                "text": "Managed 5 engineers and introduced code review standards for ML code.",
            }
        ]
    )
    out = apply_plan(master_cv, plan, JD)
    assert out.cv.bullet_index()["nimbus-3"].text.startswith("Mentored")
    assert "leadership or ownership" in (out.changes[0].reason or "")


# ---------------------------------------------------------------- requirement map


def test_requirement_map_drops_unknown_ids_and_downgrades_unshown_skills(
    master_cv: MasterCV,
) -> None:
    rows = [
        RequirementEvidence(requirement="Azure", match="direct", evidence=["nimbus-2"]),
        RequirementEvidence(requirement="MLOps", match="direct", evidence=["nimbus-2", "ghost"]),
        RequirementEvidence(requirement="Team leadership", match="transferable", evidence=["x"]),
        RequirementEvidence(requirement="Spark", match="not_evidenced"),
    ]
    checked = check_requirements(rows, master_cv, JD)
    assert checked[0].match == "adjacent" and "Azure" in checked[0].note  # not AWS-for-Azure
    assert checked[1].match == "direct" and checked[1].evidence == ["nimbus-2"]
    assert checked[2].match == "not_evidenced" and checked[2].evidence == []
    assert checked[3].match == "not_evidenced"


def test_summary_keeps_the_only_mention_of_a_job_skill(master_cv: MasterCV) -> None:
    source = master_cv.model_copy(deep=True)
    source.basics.summary = "Experienced in technology transfer and experimental design."
    job = JDAnalysis(job_title="Principal Scientist", hard_skills=["Technology transfer"])
    missing = apply_plan(source, _plan(summary="Experienced in experimental design."), job)
    assert missing.cv.basics.summary == "Experienced in experimental design."
    assert missing.role_skills == ["Technology transfer"]
    assert "Technology transfer" in printed_text(missing)
    kept = apply_plan(
        source, _plan(summary="Experienced in experimental design and technology transfer."), job
    )
    assert kept.cv.basics.summary.endswith("technology transfer.")


def test_unverified_job_technique_is_red_and_not_counted_as_confirmed(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    source = master_cv.model_copy(deep=True)
    source.education.append(Education(institution="U", degree="PhD", field="Molecular Biology"))
    job = JDAnalysis(job_title="Scientist", hard_skills=["RT-PCR", "PyTorch", "cGMP"])
    plan = _plan(
        skills_to_confirm=[
            {
                "skill": "RT-PCR",
                "background_quote": "PhD Molecular Biology U",
                "reason": "Molecular biology training makes this worth confirming",
            },
            {
                "skill": "PyTorch",
                "background_quote": "PhD Molecular Biology U",
                "reason": "Already in the CV",
            },
            {
                "skill": "cGMP",
                "background_quote": "PhD Molecular Biology U",
                "reason": "Regulated capability must not be inferred",
            },
            {
                "skill": "Western blot",
                "background_quote": "PhD Molecular Biology U",
                "reason": "Not named in this posting",
            },
            {
                "skill": "RT-PCR",
                "background_quote": "Not in this CV",
                "reason": "The quote is invented",
            },
        ]
    )
    tailored = apply_plan(source, plan, job)
    assert [s.skill for s in tailored.skills_to_confirm] == ["RT-PCR"]
    assert "RT-PCR" not in printed_text(tailored)
    assert "RT-PCR" in tailored.missing_keywords

    from src.cv.docx_exporter import export_docx

    original = tmp_path / "original.docx"
    export_docx(source, original)
    for target, use_original in (
        (tmp_path / "copy.docx", True),
        (tmp_path / "classic.docx", False),
    ):
        if use_original:
            write_like_original(
                original,
                source,
                tailored.cv,
                target,
                skills_to_confirm=tailored.skills_to_confirm,
            )
        else:
            export_docx(tailored, target)
        paragraphs = Document(target).paragraphs
        line = next(p for p in paragraphs if "Skills to confirm before use" in p.text)
        assert "RT-PCR" in line.text
        assert all(str(run.font.color.rgb) == "C00000" for run in line.runs)
        assert check_docx(target, tailored.cv, ["RT-PCR"]).keyword_coverage == 0


def test_optimised_review_additions_are_anchored_red_and_excluded_from_fit(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    job_text = "Lead Azure deployment pipelines and model handover."
    jd = JDAnalysis(job_title="ML lead", hard_skills=["Azure"])
    tailored = apply_plan(master_cv, TailoringPlan(), jd)
    raw = AdditionCandidates.model_validate(
        {
            "items": [
                {
                    "kind": "skill",
                    "text": "Azure",
                    "requirement": "Azure deployment",
                    "background_quote": "Led migration of model training pipelines to Kubernetes",
                    "job_quote": "Azure deployment pipelines",
                    "reason": "Confirm hands-on use",
                },
                {
                    "kind": "experience",
                    "text": "Supported Azure deployment and model handover",
                    "role_id": "nimbus",
                    "requirement": "model handover",
                    "background_quote": "Led migration of model training pipelines to Kubernetes",
                    "job_quote": "model handover",
                    "reason": "Confirm the actual contribution",
                },
                {
                    "kind": "experience",
                    "text": "Led Azure deployment globally",
                    "role_id": "lexa",
                    "requirement": "Azure deployment",
                    "background_quote": "Led migration of model training pipelines to Kubernetes",
                    "job_quote": "Azure deployment pipelines",
                    "reason": "Wrong role anchor",
                },
            ]
        }
    )
    additions = propose_review_additions(
        master_cv, tailored, jd, job_text, FakeLLM({AdditionCandidates: raw})
    )
    assert [a.kind for a in additions] == ["skill", "experience"]
    tailored.review_additions = additions
    assert "Azure" not in printed_text(tailored)
    original = tmp_path / "original.docx"
    export_docx(master_cv, original)
    for target, copy_original in (
        (tmp_path / "same.docx", True),
        (tmp_path / "classic.docx", False),
    ):
        if copy_original:
            write_like_original(
                original,
                master_cv,
                tailored.cv,
                target,
                review_additions=tailored.review_additions,
            )
        else:
            export_docx(tailored, target)
        lines = [p for p in Document(target).paragraphs if p.text.startswith("[CONFIRM ")]
        assert len(lines) == 2
        assert all(str(run.font.color.rgb) == "C00000" for p in lines for run in p.runs)
        assert check_docx(target, tailored.cv, ["Azure"]).keyword_coverage == 0


def test_optimised_mode_writes_one_guarded_plan_with_review_questions(master_cv: MasterCV) -> None:
    job_text = "Lead Azure deployment pipelines and model handover."
    plan = TailoringPlan.model_validate(
        {
            "review_candidates": [
                {
                    "kind": "skill",
                    "text": "Azure",
                    "requirement": "Azure deployment",
                    "background_quote": "Led migration of model training pipelines to Kubernetes",
                    "job_quote": "Azure deployment pipelines",
                    "reason": "Confirm hands-on use",
                },
            ]
        }
    )
    llm = FakeLLM({TailoringPlan: plan})
    out = tailor(
        master_cv,
        job_text,
        llm,
        jd=JDAnalysis(job_title="ML lead", hard_skills=["Azure"]),
        guidance="Gap: no hands-on Azure delivery",
        review=False,
        mode="optimised",
    )
    assert [item.text for item in out.review_additions] == ["Azure"]
    assert "<optimised_draft>" in llm.calls[0][0]
    assert "<match_assessment>\nGap: no hands-on Azure delivery" in llm.calls[0][0]
    assert [model for _, model in llm.calls] == [TailoringPlan]


def test_optimised_mode_revises_once_with_rejections_and_gaps_together(
    master_cv: MasterCV,
) -> None:
    drafts = iter(
        [
            TailoringPlan(
                rewritten_bullets=[
                    RewrittenBullet(
                        source_id="nimbus-3",
                        text=(
                            "Managed 5 engineers and introduced code review standards for ML code."
                        ),
                    )
                ]
            ),
            TailoringPlan(),
        ]
    )
    llm = FakeLLM({TailoringPlan: lambda _prompt: next(drafts)})
    out = tailor(master_cv, "ML engineering leadership", llm, jd=JD, mode="optimised")
    assert out.cv.bullet_index()["nimbus-3"].text == master_cv.bullet_index()["nimbus-3"].text
    # the rejected wording and the open gaps go into one revision; no further passes
    assert [model for _, model in llm.calls] == [TailoringPlan, TailoringPlan]
    assert "Rejected rewrite of nimbus-3" in llm.calls[1][0]


def test_skills_to_confirm_come_from_the_plan_only_when_asked(master_cv: MasterCV) -> None:
    source = master_cv.model_copy(deep=True)
    source.education.append(Education(institution="U", degree="PhD", field="Cell Biology"))
    plan = TailoringPlan.model_validate(
        {
            "skills_to_confirm": [
                {
                    "skill": "RT-PCR",
                    "background_quote": "PhD Cell Biology U",
                    "job_quote": "cell characterisation methods",
                    "reason": "Worth confirming for the role",
                }
            ]
        }
    )
    llm = FakeLLM({TailoringPlan: plan})
    job, jd = "Develop cell characterisation methods.", JDAnalysis(job_title="Scientist")
    out = tailor(source, job, llm, jd=jd, suggest_unverified=True, review=False)
    assert [s.skill for s in out.skills_to_confirm] == ["RT-PCR"]
    assert [m for _, m in llm.calls] == [TailoringPlan]  # no separate suggestion pass
    assert tailor(source, job, llm, jd=jd, review=False).skills_to_confirm == []


def test_job_wording_needs_the_bullet_mapped_to_that_requirement(master_cv: MasterCV) -> None:
    text = (
        "Designed PostgreSQL feature store used by 3 data science teams, through stakeholder "
        "management across teams."
    )
    rewrite = [{"source_id": "lexa-2", "text": text}]
    unmapped = apply_plan(master_cv, _plan(rewritten_bullets=rewrite), JD)
    assert "stakeholder management" in (unmapped.changes[0].reason or "")
    row = {"requirement": "stakeholder management", "match": "transferable",
           "evidence": ["lexa-2"], "note": "served 3 teams"}  # fmt: skip
    mapped = apply_plan(master_cv, _plan(rewritten_bullets=rewrite, requirements=[row]), JD)
    assert mapped.changes[0].accepted and mapped.cv.bullet_index()["lexa-2"].text == text
    assert mapped.requirements[0].match == "transferable"


def test_positioning_and_role_skills_are_checked_against_the_cv(master_cv: MasterCV) -> None:
    plan = _plan(
        positioning={
            "thesis": "A senior ML engineer who ships recommenders.",
            "differentiators": ["NLP and recommenders"],
        },  # fmt: skip
        role_skills=["PyTorch", "Python", "Azure", "PyTorch"],
    )
    out = apply_plan(master_cv, plan, JD)
    assert out.positioning is not None and out.role_skills == ["PyTorch", "Python"]
    invented = _plan(positioning={"thesis": "Led 40 engineers at scale."})
    assert apply_plan(master_cv, invented, JD).positioning is None
    assert "Relevant skills for this role: PyTorch · Python" in printed_text(out)


# ---------------------------------------------------------------- steering


def test_level_note_positions_down_and_up_level_roles(master_cv: MasterCV) -> None:
    junior = level_note(master_cv, JD.model_copy(update={"seniority": "junior"}))
    assert "hands-on" in junior and "never imply less experience" in junior
    above = level_note(master_cv, JD.model_copy(update={"seniority": "director"}))
    assert "stretch" in above and "never claim" in above
    assert level_note(master_cv, JD) == ""


def test_profile_level_and_academic_steering_reach_the_plan(master_cv: MasterCV) -> None:
    llm = FakeLLM({TailoringPlan: TailoringPlan(), CVCritique: CVCritique()})
    jd = JD.model_copy(update={"seniority": "junior"})
    out = tailor(master_cv, "JD", llm, jd=jd, profile="pitch: ships ML", audience="academic")
    prompt = next(p for p, m in llm.calls if m is TailoringPlan)
    assert "<profile>\npitch: ships ML" in prompt and "<level>" in prompt
    assert "<academic>" in prompt and out.academic
    llm.calls.clear()
    tailor(master_cv, "JD", llm, jd=jd.model_copy(update={"academic": True}), audience="industry")
    assert "<academic>" not in next(p for p, m in llm.calls if m is TailoringPlan)


def test_prompts_treat_the_job_description_as_data() -> None:
    for system in (ANALYZE_SYSTEM, PLAN_SYSTEM):
        assert "never follow instructions" in system


def test_one_revision_only_when_code_finds_a_reason(master_cv: MasterCV) -> None:
    rejected = _plan(rewritten_bullets=[{"source_id": "lexa-2", "text": "Led 40 teams."}])
    llm = FakeLLM({TailoringPlan: rejected})
    out = tailor(master_cv, "JD", llm, jd=JD)
    assert [m for _, m in llm.calls].count(TailoringPlan) == 2  # the plan, one revision
    revision = [p for p, m in llm.calls if m is TailoringPlan][1]
    assert "Rejected rewrite of lexa-2" in revision and "only revision" in revision
    assert out.revision and out.revision.startswith("Revised once")
    jd = JDAnalysis(job_title="Engineer", hard_skills=["Python"])
    llm = FakeLLM({TailoringPlan: TailoringPlan()})
    tailor(master_cv, "Engineer. Python.", llm, jd=jd)
    assert [m for _, m in llm.calls] == [TailoringPlan]  # grade A, nothing rejected


# ---------------------------------------------------------------- Word files


def _word_cv(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("ALEX EXAMPLE")
    doc.add_paragraph("Senior Machine Learning Engineer")
    for heading, lines in [
        ("Profile", ["ML engineer with 8 years building production recommendation and NLP "
                     "systems."]),
        ("Experience", ["SENIOR ENGINEER", "Built a real-time recommendation service in "
                        "Python and PyTorch serving 4M users."]),
        ("Skills", ["Languages: Python, SQL, Go"]),
        ("Education", ["MSc Computer Science, Example University"]),
        ("Selected Publications", ["Example A. A paper. Journal 2020."]),
    ]:  # fmt: skip
        doc.add_paragraph(heading, style="Heading 1")
        for line in lines:
            doc.add_paragraph(line)
    doc.save(str(path))
    return path


def _texts(path: Path) -> list[str]:
    return [p.text for p in Document(str(path)).paragraphs if p.text.strip()]


def test_academic_order_moves_whole_sections_and_keeps_the_name_first(tmp_path: Path) -> None:
    doc = Document(str(_word_cv(tmp_path / "cv.docx")))
    notes = academic_order(doc)
    doc.save(str(tmp_path / "out.docx"))
    texts = _texts(tmp_path / "out.docx")
    assert texts[0] == "ALEX EXAMPLE" and "SENIOR ENGINEER" in texts  # a role line never moves
    headings = [t for t in texts if t in {"Profile", "Experience", "Skills", "Education",
                                          "Selected Publications"}]  # fmt: skip
    assert headings == ["Profile", "Education", "Experience", "Selected Publications", "Skills"]
    assert texts.index("MSc Computer Science, Example University") == texts.index("Education") + 1
    assert notes and "academic" in notes[0]


def test_word_copy_gets_the_role_skills_line_and_academic_order(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    original = _word_cv(tmp_path / "cv.docx")
    out = tmp_path / "tailored.docx"
    write_like_original(
        original, master_cv, master_cv, out, role_skills=["PyTorch", "Python"], academic=True
    )
    texts = _texts(out)
    line = "Relevant skills for this role: PyTorch · Python"
    assert texts[texts.index("Profile") + 2] == line
    assert texts.index("Education") < texts.index("Experience")
    industry = tmp_path / "industry.docx"
    write_like_original(original, master_cv, master_cv, industry, role_skills=["Python"])
    assert _texts(industry).index("Skills") < _texts(industry).index("Education")


def test_template_export_carries_role_skills_and_puts_education_first_for_academia(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    tailored = apply_plan(master_cv, _plan(role_skills=["Python"]), JD)
    tailored = tailored.model_copy(update={"academic": True})
    texts = _texts(export_docx(tailored, tmp_path / "t.docx", "classic"))
    assert "Relevant skills for this role: Python · PyTorch" in texts  # job keywords lead
    upper = [t.casefold() for t in texts]
    assert upper.index("education") < upper.index("experience")


# ---------------------------------------------------------------- before/after fit check


def test_fit_check_reads_both_cvs_blind_and_reports_the_change(master_cv: MasterCV) -> None:
    tailored = apply_plan(master_cv, _plan(role_skills=["PyTorch"]), JD)
    before = {"function": 2, "domain": 3, "seniority": 3, "leadership": 2, "sector": 2}
    after = {"function": 4, "domain": 3, "seniority": 3, "leadership": 1, "sector": 3}

    def reader(prompt: str) -> EvidencedReadings:
        out = []
        for _, text in re.findall(r"<cv_(\d+)>\n(.*?)\n</cv_\1>", prompt, re.DOTALL):
            line = next(ln for ln in text.splitlines() if len(ln.split()) >= 2)
            mine = "Relevant skills for this role" in text
            levels, met = (after, [1, 2, 9]) if mine else (before, [1])
            out.append(EvidencedReading(
                **{n: RatedLine(rating=r, quote=line) for n, r in levels.items()},
                requirements_met=[MetRequirement(number=n, quote=line) for n in met],
                impression="clear fit" if mine else "generic",
            ))  # fmt: skip
        return EvidencedReadings(readings=out)

    llm = FakeLLM({EvidencedReadings: reader})
    check = compare_fit(master_cv, tailored, JD, "We need a senior ML engineer.", llm)
    assert check.before < check.after and check.impression_after == "clear fit"
    assert check.requirements_total == 1 and check.requirements_after == 1  # "MLOps" only
    assert check.worse == ["leadership: 2 → 1"]
    prompt = llm.calls[0][0]
    assert "tailored" not in prompt.casefold() and "Relevant skills for this role" in prompt


def test_a_failed_fit_check_keeps_the_tailored_cv(master_cv: MasterCV) -> None:
    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    out = cv_service._fit_check(master_cv, tailored, JD, "JD", FakeLLM({}))  # type: ignore[arg-type]
    assert out.fit_check is None and "fit check could not run" in out.document_notes[0]


# ---------------------------------------------------------------- cover letter

OPENING = LetterOpening(
    interest="I am applying for the Senior ML Engineer role for its recommender work.",
    fit=LetterParagraph(text="My recommender work at Nimbus fits it.", source_ids=["nimbus"]),
)
EVIDENCE = [
    LetterParagraph(text="I built recommenders at Nimbus.", source_ids=["nimbus-1"]),
    LetterParagraph(text="I moved training to Kubernetes.", source_ids=["nimbus-2"]),
]
CLOSE = LetterParagraph(text="I would welcome a conversation.", source_ids=["nimbus"])


def test_bridge_paragraph_is_kept_before_the_conclusion_and_guarded(master_cv: MasterCV) -> None:
    bridge = LetterParagraph(
        text="My mentoring of engineers carries into this role.", source_ids=["nimbus-3"]
    )
    draft = CoverLetterDraft(opening=OPENING, evidence=EVIDENCE, bridge=bridge, conclusion=CLOSE)
    letter = apply_letter(master_cv, draft, JD, "JD")
    assert letter.paragraphs[3] == bridge.text and letter.paragraphs[-1] == CLOSE.text
    inflated = bridge.model_copy(update={"text": "I managed engineers."})
    rejected = apply_letter(master_cv, draft.model_copy(update={"bridge": inflated}), JD, "JD")
    assert "leadership or ownership" in (rejected.changes[3].reason or "")


def test_letter_review_drives_one_guarded_redraft(master_cv: MasterCV) -> None:
    first = LetterSections(opening=OPENING, evidence=EVIDENCE, conclusion=CLOSE)
    better = first.model_copy(update={"conclusion": CLOSE.model_copy(
        update={"text": "I would welcome a conversation about your ranking work."})})  # fmt: skip
    drafts = [first, better]
    llm = FakeLLM(
        {
            LetterSections: lambda _: drafts.pop(0),
            LetterReview: LetterReview(issues=["the close is generic"]),
        }
    )
    letter = write_letter(master_cv, JD, "JD", llm, context="<positioning>x</positioning>")
    assert letter.paragraphs[-1].endswith("ranking work.")
    prompts = [p for p, m in llm.calls if m is LetterSections]
    assert "Improve the letter: the close is generic" in prompts[1]
    assert all("<positioning>x</positioning>" in p for p in prompts)


def test_a_redraft_that_fails_the_guards_keeps_the_first_letter(master_cv: MasterCV) -> None:
    first = LetterSections(opening=OPENING, evidence=EVIDENCE, conclusion=CLOSE)
    bad = first.model_copy(update={"evidence": [EVIDENCE[0], LetterParagraph(
        text="I led 40 engineers.", source_ids=["nimbus-1"])]})  # fmt: skip
    drafts = [first, bad]
    llm = FakeLLM(
        {LetterSections: lambda _: drafts.pop(0), LetterReview: LetterReview(issues=["more"])}
    )
    assert write_letter(master_cv, JD, "JD", llm).paragraphs[2] == EVIDENCE[1].text


def test_letter_context_carries_positioning_and_requirements(master_cv: MasterCV) -> None:
    plan = _plan(
        positioning={"thesis": "Ships recommenders."},
        requirements=[{"requirement": "MLOps", "match": "direct", "evidence": ["nimbus-2"]}],
    )
    context = letter_context(apply_plan(master_cv, plan, JD))
    assert "Ships recommenders." in context and "- MLOps (essential): direct [nimbus-2]" in context
    assert letter_context(None) == ""


# ---------------------------------------------------------------- profile


def _summary(**kw: object) -> ProfileSummary:
    base: dict[str, object] = {
        "headline": "ML engineer", "seniority": "senior", "years_experience": 8,
        "core_expertise": ["recommenders"], "key_skills": [], "target_roles": [],
        "not_a_fit": [], "summary": "An ML engineer.",
    }  # fmt: skip
    return ProfileSummary.model_validate(base | kw)


def test_profile_claims_are_checked_against_the_cv(master_cv: MasterCV) -> None:
    summary = _summary(
        achievements=["Served 4M users with a recommender", "Grew revenue 300%"],
        pitch="A world-class engineer.",
        expansion=[
            ExpansionRoute(capability="NLP", opens=["legal tech"], evidence=["[lexa-1]"]),
            ExpansionRoute(capability="Sales", opens=["sales"], evidence=["ghost"]),
        ],
    )
    checked = check_claims(summary, master_cv)
    assert checked.achievements == ["Served 4M users with a recommender"]
    assert checked.pitch == ""
    assert [r.capability for r in checked.expansion] == ["NLP"]
    assert checked.expansion[0].evidence == ["lexa-1"]


def test_matcher_does_not_see_presentation_fields() -> None:
    family = RoleFamily(name="Core", titles=["ML Engineer"], positioning="Lead with recsys.")
    summary = _summary(pitch="Ships recommenders.", role_families=[family],
                       expansion=[ExpansionRoute(capability="NLP")])  # fmt: skip
    shown = matcher_profile(summary)
    assert "Ships recommenders." not in shown and "Lead with recsys." not in shown
    assert "expansion" not in shown


def test_stored_profile_feeds_tailoring_with_the_jobs_family(
    master_cv: MasterCV, settings: object, tmp_path: Path
) -> None:
    family = RoleFamily(name="Recsys", titles=["ML Engineer"], positioning="Lead with recsys.",
                        gap="none", evidence=["nimbus-1"])  # fmt: skip
    summary = _summary(pitch="Ships recommenders.", role_families=[family])
    brief = profile_brief(summary, "Recsys")
    assert "pitch: Ships recommenders." in brief and "Lead with recsys." in brief
    assert "role family" not in profile_brief(summary, "Other")
    ws = Workspace(settings)  # type: ignore[arg-type]
    ws.memory = ProfileMemory(tmp_path / "profiles.json")
    ws.active_cv_id = "cv1"
    assert cv_service.stored_profile(ws) == ""
    ws.memory.put("cv1", "any", summary, "fp")
    assert "pitch: Ships recommenders." in cv_service.stored_profile(ws, "Recsys")


def test_rewritten_bullet_model_still_matches(master_cv: MasterCV) -> None:
    # A plan from an older model (no requirement map) still applies its safe rewrites.
    text = (
        "Led the migration of model training pipelines to Kubernetes, cutting training costs "
        "by 35%."
    )
    plan = TailoringPlan(rewritten_bullets=[RewrittenBullet(source_id="nimbus-2", text=text)])
    out = apply_plan(master_cv, plan, JD)
    assert out.changes[0].accepted and out.positioning is None and out.requirements == []
    assert Positioning(thesis="x").differentiators == []


def test_proposal_pass_is_given_essential_gaps_first(master_cv: MasterCV) -> None:
    def row(requirement: str, priority: str, importance: str, match: str) -> RequirementEvidence:
        return RequirementEvidence.model_validate(
            {"requirement": requirement, "priority": priority, "importance": importance,
             "match": match}
        )  # fmt: skip

    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    tailored.requirements = [
        row("Pivotal desirable", "desirable", "key", "adjacent"),
        row("Routine desirable", "desirable", "supporting", "not_evidenced"),
        row("Essential knock-out", "essential", "supporting", "not_evidenced"),
        row("Shown essential", "essential", "key", "direct"),
    ]
    llm = FakeLLM({AdditionCandidates: AdditionCandidates()})
    propose_review_additions(master_cv, tailored, JD, "job text", llm)
    gaps = llm.calls[0][0].split("<gaps>\n")[1].split("\n</gaps>")[0].splitlines()
    assert [g.split(" (")[0] for g in gaps] == ["- Essential knock-out", "- Pivotal desirable"]


def test_other_cvs_profiles_stay_listed_until_the_user_deletes_them(
    settings: object, tmp_path: Path, master_cv: MasterCV
) -> None:
    from src.services import search_service

    ws = Workspace(settings)  # type: ignore[arg-type]
    ws.memory = ProfileMemory(tmp_path / "profiles.json")
    ws.memory.put("old-cv", "any", _summary(headline="Earlier profile"), "fp-old")
    ws.active_cv_id = "new-cv"
    ws.memory.put("new-cv", "any", _summary(headline="New profile"), "fp-new")
    own = search_service.list_profiles(ws, master_cv)
    others = search_service.other_profiles(ws, master_cv)
    assert [r.summary.headline for r in own] == ["New profile"]
    assert [r.summary.headline for r in others] == ["Earlier profile"]
    assert others[0].cv_name == "Another CV"  # its CV is not in this library
    search_service.delete_profile(ws, "old-cv:any")  # only on the user's request
    assert search_service.other_profiles(ws, master_cv) == []
    with pytest.raises(ValueError, match="not found"):
        search_service.delete_profile(ws, "old-cv:any")
