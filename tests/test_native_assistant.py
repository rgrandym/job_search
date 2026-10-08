"""The assistant run natively (Claude Code's own loop over the app's MCP tools), chat memory,
and CV edits reaching the Word document. Offline: a fake native model stands in for the CLI."""

from __future__ import annotations

import asyncio
from collections.abc import Awaitable, Callable
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient

from src.agents import chat, mcp
from src.agents.definitions import ASSISTANT
from src.agents.runtime import RUNS, AgentContext, NativeRun, run_agent
from src.core.config import Settings
from src.core.llm import ChatMessage, LLMError, NativeTurn, ToolCall
from src.core.llm import codex_backend as codex
from src.core.llm.claude_code_backend import claude_events
from src.cv.edits import CVEdit, apply_edits
from src.cv.models import MasterCV
from src.jobs.models import SearchQuery
from src.services import cv_document, cv_service
from src.services.workspace import Workspace


@pytest.fixture
def ws(settings: Settings, master_cv: MasterCV, monkeypatch: pytest.MonkeyPatch) -> Workspace:
    w = Workspace(settings)
    w.master_cv = master_cv
    w.active_cv_id = "master"
    monkeypatch.setattr(cv_service, "save_selected_cv", lambda w, cv: setattr(w, "master_cv", cv))
    return w


async def _ignore(kind: str, payload: dict[str, Any]) -> None:
    return None


class FakeClaudeCode:
    """Stands in for `claude -p`: calls the app's tools over MCP like the CLI, then answers."""

    def __init__(self, calls: list[tuple[str, dict[str, Any]]]) -> None:
        self.calls = calls
        self.seen: list[dict[str, Any]] = []

    async def chat(self, **_: Any) -> Any:
        raise AssertionError("a native model must not be driven turn by turn")

    async def run_native(
        self,
        *,
        system: str,
        prompt: str,
        mcp_url: str,
        session_id: str | None,
        workdir: Path,
        on_text: Callable[[str], Awaitable[None]],
        cancelled: Callable[[], bool],
    ) -> NativeTurn:
        token = mcp_url.rsplit("/", 1)[-1]
        self.seen.append({"prompt": prompt, "session_id": session_id})
        listed = await mcp.handle(token, {"jsonrpc": "2.0", "id": 1, "method": "tools/list"})
        assert listed and {"edit_cv", "update_profile"} <= {
            t["name"] for t in listed["result"]["tools"]
        }
        await on_text("Writing it now.")
        for name, arguments in self.calls:
            params = {"name": name, "arguments": arguments}
            reply = await mcp.handle(
                token, {"jsonrpc": "2.0", "id": 2, "method": "tools/call", "params": params}
            )
            assert reply and not reply["result"]["isError"], reply
        return NativeTurn(text="Done.", session_id=session_id or "session-1", input_tokens=5)


def test_mcp_server_lists_and_runs_tools_only_during_a_run(ws: Workspace) -> None:
    ctx = AgentContext(ws=ws, emit=_ignore)
    RUNS["tok"] = NativeRun(ASSISTANT, ctx, 0)
    try:
        init = asyncio.run(mcp.handle("tok", {"jsonrpc": "2.0", "id": 0, "method": "initialize"}))
        assert init and init["result"]["capabilities"]["tools"] is not None
        note = {"jsonrpc": "2.0", "method": "notifications/initialized"}
        assert asyncio.run(mcp.handle("tok", note)) is None
        params = {"name": "read_cv", "arguments": {}}
        call = {"jsonrpc": "2.0", "id": 3, "method": "tools/call", "params": params}
        reply = asyncio.run(mcp.handle("tok", call))
        assert reply and not reply["result"]["isError"]
        assert "experience" in reply["result"]["content"][0]["text"]
    finally:
        RUNS.pop("tok")
    ended = asyncio.run(mcp.handle("tok", call))
    assert ended and ended["error"]["code"] == -32001


def test_native_model_runs_its_own_loop_and_resumes_its_session(
    ws: Workspace, master_cv: MasterCV, monkeypatch: pytest.MonkeyPatch
) -> None:
    edit = {
        "action": "add_bullet",
        "role_id": master_cv.experience[0].id,
        "text": "Ran flow cytometry experiments",
    }
    fake = FakeClaudeCode([("edit_cv", {"edits": [edit], "reason": "ok"})])
    monkeypatch.setattr(ws, "chat", lambda *a: fake)
    events: list[tuple[str, dict[str, Any]]] = []

    async def emit(kind: str, payload: dict[str, Any]) -> None:
        events.append((kind, payload))

    session = chat.ChatSession()
    args = (ws, session, "Yes, add it", SearchQuery(), True, emit)
    asyncio.run(chat.handle_user_message(*args, mcp_url="http://127.0.0.1:8000/api/mcp"))
    assert ws.master_cv and ws.master_cv.experience[0].bullets[-1].text.startswith("Ran flow")
    provider = ws.llm.provider
    assert session.native_session == f"{provider}:session-1" and not RUNS
    kinds = [kind for kind, _ in events]
    assert {"tool_call", "tool_result", "cv_updated", "done"} <= set(kinds)
    final = {"agent": "assistant", "depth": 0, "text": "Done.", "final": True}
    assert ("agent_message", final) in events
    assert [m.content for m in session.messages if m.role == "assistant"] == ["Done."]

    fake.calls = []
    args = (ws, session, "Thanks", SearchQuery(), True, emit)
    asyncio.run(chat.handle_user_message(*args, mcp_url="http://127.0.0.1:8000/api/mcp"))
    assert fake.seen[1]["session_id"] == "session-1"
    assert "earlier_conversation" not in fake.seen[1]["prompt"]  # the session remembers
    assert chat._native_path(ws, session.id).read_text() == f"{provider}:session-1"
    assert chat._native_for(session, "codex" if provider != "codex" else "claude_code") is None
    chat.clear_session(ws, session)
    assert session.native_session is None and not chat._native_path(ws, session.id).exists()


def test_without_an_mcp_endpoint_the_app_loop_still_runs(ws: Workspace) -> None:
    from tests.test_webapp import ScriptedChat

    script = {"assistant": [ChatMessage(role="assistant", content="Hello.")]}
    ws.chat = lambda *a: ScriptedChat(script)  # type: ignore[method-assign]
    messages = [ChatMessage(role="user", content="hi")]
    assert asyncio.run(run_agent(ASSISTANT, messages, AgentContext(ws=ws, emit=_ignore))) == (
        "Hello."
    )


def _feed(on_event: Any, lines: list[dict[str, Any]]) -> None:
    async def run() -> None:
        for line in lines:
            await on_event(line)

    asyncio.run(run())


def test_claude_stream_passes_on_text_before_tool_use_and_holds_the_answer() -> None:
    seen: list[str] = []

    async def on_text(text: str) -> None:
        seen.append(text)

    events: list[dict[str, Any]] = []
    _feed(
        claude_events(events, on_text),
        [
            {"type": "assistant", "message": {"content": [{"type": "text", "text": "Reading."}]}},
            {"type": "assistant", "message": {"content": [{"type": "tool_use", "name": "x"}]}},
            {"type": "assistant", "message": {"content": [{"type": "text", "text": "Done."}]}},
            {"type": "result", "subtype": "success", "result": "Done.", "session_id": "s"},
        ],
    )
    assert seen == ["Reading."] and events[-1]["type"] == "result"


def test_codex_runs_natively_with_only_the_apps_tools_and_resumes_its_thread() -> None:
    from src.core.llm import LLMConfig

    cfg = LLMConfig(provider="codex", quality_model="m", screening_model="m")
    fresh = codex._native_command("codex", cfg, "m", "Be brief.", "http://x/api/mcp/t", None)
    resumed = codex._native_command("codex", cfg, "m", "Be brief.", "http://x/api/mcp/t", "th-1")
    assert fresh[:3] == ["codex", "exec", "--json"] and fresh[-1] == "-"
    assert resumed[:3] == ["codex", "exec", "resume"] and resumed[-2:] == ["th-1", "-"]
    assert 'mcp_servers.jobsearch.url="http://x/api/mcp/t"' in fresh
    assert "features.shell_tool=false" in fresh and 'web_search="disabled"' in fresh

    seen: list[str] = []

    async def on_text(text: str) -> None:
        seen.append(text)

    events: list[dict[str, Any]] = []
    on_event, relay = codex.codex_events(events, on_text)
    _feed(
        on_event,
        [
            {"type": "thread.started", "thread_id": "th-1"},
            {"type": "item.completed", "item": {"type": "agent_message", "text": "Reading."}},
            {"type": "item.started", "item": {"type": "mcp_tool_call", "tool": "read_cv"}},
            {"type": "item.completed", "item": {"type": "agent_message", "text": "Added."}},
            {"type": "turn.completed", "usage": {"input_tokens": 9, "output_tokens": 2}},
        ],
    )
    turn = codex._native_result(events, relay, 0, "")
    assert seen == ["Reading."] and turn.text == "Added." and turn.session_id == "th-1"
    with pytest.raises(LLMError, match="quota"):
        codex._native_result([{"type": "turn.failed", "error": {"message": "quota"}}], relay, 1, "")


def test_recent_turns_keep_tool_results_and_older_turns_keep_answers() -> None:
    messages: list[ChatMessage] = []
    for turn in range(chat.TOOL_TURNS + 2):
        call = ToolCall(id=f"c{turn}", name="read_cv")
        messages += [
            ChatMessage(role="user", content=f"q{turn}"),
            ChatMessage(role="assistant", tool_calls=[call]),
            ChatMessage(role="tool", tool_call_id=f"c{turn}", content="cv"),
            ChatMessage(role="assistant", content=f"a{turn}"),
        ]
    kept = chat._compact(messages)
    assert [m.role for m in kept[:4]] == ["user", "assistant", "user", "assistant"]
    assert sum(m.role == "tool" for m in kept) == chat.TOOL_TURNS


def test_bullet_edits_reach_the_word_copy(
    ws: Workspace, master_cv: MasterCV, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    role = master_cv.experience[0]
    path = tmp_path / "cv.docx"
    doc = Document()
    doc.add_paragraph(f"{role.title}, {role.company}")
    for bullet in role.bullets:
        doc.add_paragraph(bullet.text, style="List Bullet")
    doc.add_paragraph("Education")
    doc.save(str(path))
    monkeypatch.setattr(cv_document, "_document", lambda w: path)

    data = master_cv.model_dump(mode="json")
    bullets = data["experience"][0]["bullets"]
    bullets[0]["text"] = "Reworded first bullet"
    removed = bullets.pop(1)["text"]
    bullets.append({"id": "added-1", "text": "Ran flow cytometry experiments"})
    out = cv_document.sync_bullets(ws, master_cv, MasterCV.model_validate(data))

    texts = [p.text for p in Document(str(path)).paragraphs]
    assert texts[1] == "Reworded first bullet" and removed not in texts
    added = texts.index("Ran flow cytometry experiments")
    assert texts[added + 1] == "Education"  # placed at the end of that role's bullets
    assert Document(str(path)).paragraphs[added].style.name == "List Bullet"
    assert len(out.written) == 3 and not out.not_written


def test_mcp_endpoint_over_http(client: TestClient) -> None:
    gone = client.post("/api/mcp/unknown", json={"jsonrpc": "2.0", "id": 1, "method": "tools/list"})
    assert gone.status_code == 200 and gone.json()["error"]["code"] == -32001
    note = client.post("/api/mcp/unknown", json={"jsonrpc": "2.0", "method": "x"})
    assert note.status_code == 202
    assert client.get("/api/mcp/unknown").status_code == 405


@pytest.fixture
def client(ws: Workspace, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    from src.web import app as web_app

    monkeypatch.setattr(web_app, "get_workspace", lambda: ws)
    return TestClient(web_app.app)


def test_targeted_cv_edits_change_only_what_they_name(master_cv: MasterCV) -> None:
    role, other = master_cv.experience[0], master_cv.experience[1]
    first, second = role.bullets[0], role.bullets[1]
    edits = [
        CVEdit(action="add_bullet", role_id=role.id, text="New line", after=first.id),
        CVEdit(action="edit_bullet", bullet_id=second.id, text="Reworded"),
        CVEdit(action="remove_bullet", bullet_id=other.bullets[0].id),
        CVEdit(action="edit_role", role_id=other.id, fields={"title": "Lead"}),
        CVEdit(action="set_section", section="basics", value={"headline": "Scientist"}),
    ]
    out = apply_edits(master_cv, edits)
    texts = [b.text for b in out.experience[0].bullets]
    assert texts[:3] == [first.text, "New line", "Reworded"]
    assert out.experience[1].title == "Lead"
    assert len(out.experience[1].bullets) == len(other.bullets) - 1
    assert out.basics.headline == "Scientist" and out.basics.name == master_cv.basics.name
    assert out.experience[2:] == master_cv.experience[2:] and out.skills == master_cv.skills
    with pytest.raises(ValueError, match="No bullet"):
        apply_edits(master_cv, [CVEdit(action="edit_bullet", bullet_id="nope", text="x")])


def _docx_bytes(lines: list[str]) -> bytes:
    import io

    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_an_uploaded_cv_stays_one_cv_with_one_id_when_edited(
    settings: Settings, master_cv: MasterCV
) -> None:
    import os

    ws = Workspace(settings)
    role = master_cv.experience[0]
    asset = cv_service.store_cv(
        ws, "Master CV.docx", _docx_bytes([role.title, *[b.text for b in role.bullets]])
    )
    original = settings.data_dir / "cvs" / "Master CV.docx"
    copy = ws.output_dir / "cvs" / "Master CV.docx"
    cv_service._parsed_path(ws, asset.id).parent.mkdir(parents=True, exist_ok=True)
    cv_service._parsed_path(ws, asset.id).write_text(master_cv.model_dump_json())

    # Saved in Word (the output/cvs copy): still one CV, same id, original updated.
    copy.write_bytes(_docx_bytes([role.title, "Edited in Word"]))
    later = original.stat().st_mtime + 5
    os.utime(copy, (later, later))
    listed = cv_service.list_cvs(ws)
    assert [(a.id, a.filename, a.kind) for a in listed] == [
        (asset.id, "Master CV.docx", "uploaded")
    ]
    assert original.read_bytes() == copy.read_bytes()

    # Edited by the assistant: the original changes, the copy follows, the id holds.
    copy.write_bytes(original.read_bytes())
    ws.active_cv_id = asset.id
    before = MasterCV.model_validate_json(cv_service._parsed_path(ws, asset.id).read_text())
    data = before.model_dump(mode="json")
    data["experience"][0]["bullets"] = [{"id": "x-1", "text": "Edited in Word"}]
    edited = MasterCV.model_validate(data)
    data["experience"][0]["bullets"].append({"id": "x-2", "text": "Ran flow cytometry"})
    out = cv_document.sync_bullets(ws, edited, MasterCV.model_validate(data))
    assert out.written and "Ran flow cytometry" in [
        p.text for p in Document(str(original)).paragraphs
    ]
    assert original.read_bytes() == copy.read_bytes()
    assert [a.id for a in cv_service.list_cvs(ws)] == [asset.id]
    assert cv_service.cached_selected_cv(ws) is not None  # its parse is still attached


def test_tailoring_a_word_cv_writes_a_copy_of_it_and_saved_edits_keep_its_design(
    settings: Settings, master_cv: MasterCV, monkeypatch: pytest.MonkeyPatch
) -> None:
    from src.cv.models import JDAnalysis, TailoredCVEdits, TailoringPlan
    from src.cv.tailor import apply_plan
    from src.jobs.models import JobPosting
    from src.services import tailored_documents
    from tests.test_cv_writer import _original_cv

    ws = Workspace(settings)
    original = _original_cv(master_cv, settings.data_dir / "upload.docx")
    asset = cv_service.store_cv(ws, "My CV.docx", original.read_bytes())
    cv_service._parsed_path(ws, asset.id).parent.mkdir(parents=True, exist_ok=True)
    cv_service._parsed_path(ws, asset.id).write_text(master_cv.model_dump_json())
    cv_service.select_cv(ws, asset.id)
    jd = JDAnalysis(job_title="ML Engineer", seniority="senior")
    plan = TailoringPlan(bullet_order={"nimbus": ["nimbus-3"]})
    monkeypatch.setattr(ws, "structured", lambda *a: None)
    monkeypatch.setattr(cv_service, "analyze_jd", lambda text, llm: jd)
    monkeypatch.setattr(cv_service, "tailor", lambda cv, *a: apply_plan(cv, plan, jd))
    job = JobPosting(id="j1", title="ML Engineer", company="Orbit", description="Build ML.")

    tailored, path = asyncio.run(cv_service.tailor_to_job(ws, job))
    texts = [p.text for p in Document(str(path)).paragraphs]
    assert any(t.startswith("Example A. et al.") for t in texts)  # the original, not a template
    nimbus = texts.index("Senior Machine Learning Engineer, Nimbus Analytics")
    assert texts[nimbus + 1].startswith("Mentored") and len(texts) == len(
        Document(str(original)).paragraphs
    )
    document = tailored_documents.load(ws, tailored.document_id or "")
    assert document.template == "original" and document.original_file

    edited = tailored_documents.edit(
        ws, document.id, TailoredCVEdits(bullets={"nimbus-3": "Mentored 5 engineers."})
    )
    edited_texts = [
        p.text for p in Document(str(ws.output_dir / "cvs" / edited.filename)).paragraphs
    ]
    assert "Mentored 5 engineers." in edited_texts
    assert any(t.startswith("Example A. et al.") for t in edited_texts)

    # Picking a suggested headline writes it into the Word file, in the headline's own place.
    picked = tailored_documents.edit(
        ws, document.id, TailoredCVEdits(headline="Senior Machine Learning Engineer, NLP")
    )
    picked_doc = Document(str(ws.output_dir / "cvs" / picked.filename)).paragraphs
    assert picked_doc[1].text == "Senior Machine Learning Engineer, NLP"
    assert picked_doc[0].text == master_cv.basics.name


def test_a_pdf_cv_is_tailored_in_its_word_versions_design_unless_a_design_is_chosen(
    settings: Settings, master_cv: MasterCV, monkeypatch: pytest.MonkeyPatch
) -> None:
    from src.cv.models import JDAnalysis, TailoringPlan
    from src.cv.tailor import apply_plan
    from src.jobs.models import JobPosting
    from tests.test_cv_writer import _original_cv

    ws = Workspace(settings)
    word = _original_cv(master_cv, settings.data_dir / "w.docx")
    cv_service.store_cv(ws, "Candidate CV.docx", word.read_bytes())
    pdf = cv_service.store_cv(ws, "Candidate CV.pdf", b"%PDF-1.4 exported from Word")
    cv_service._parsed_path(ws, pdf.id).parent.mkdir(parents=True, exist_ok=True)
    cv_service._parsed_path(ws, pdf.id).write_text(master_cv.model_dump_json())
    cv_service.select_cv(ws, pdf.id)
    assert cv_service.original_docx(ws) == settings.data_dir / "cvs" / "Candidate CV.docx"

    jd = JDAnalysis(job_title="ML Engineer")
    monkeypatch.setattr(ws, "structured", lambda *a: None)
    monkeypatch.setattr(cv_service, "analyze_jd", lambda text, llm: jd)
    monkeypatch.setattr(cv_service, "tailor", lambda cv, *a: apply_plan(cv, TailoringPlan(), jd))
    job = JobPosting(id="j2", title="ML Engineer", company="Orbit", description="Build ML.")
    tailored, path = asyncio.run(cv_service.tailor_to_job(ws, job))
    assert any(p.text.startswith("Example A.") for p in Document(str(path)).paragraphs)
    assert tailored.document_notes[0] == "Written in the design of Candidate CV.docx"

    chosen, path = asyncio.run(cv_service.tailor_to_job(ws, job, template="modern"))
    assert not any(p.text.startswith("Example A.") for p in Document(str(path)).paragraphs)
    assert chosen.document_notes == ["Written in the modern design, as you chose"]


def test_the_structured_master_cv_is_tailored_in_its_word_cvs_design(
    settings: Settings, master_cv: MasterCV, monkeypatch: pytest.MonkeyPatch
) -> None:
    from src.cv.models import JDAnalysis, TailoringPlan
    from src.cv.tailor import apply_plan
    from src.jobs.models import JobPosting
    from tests.test_cv_writer import _original_cv

    ws = Workspace(settings)
    word = _original_cv(master_cv, settings.data_dir / "w.docx")
    cv_service.store_cv(ws, "Master_CV_2026.docx", word.read_bytes())
    ws.master_cv, ws.active_cv_id = master_cv, "master"
    assert cv_service.original_docx(ws) == settings.data_dir / "cvs" / "Master_CV_2026.docx"

    jd = JDAnalysis(job_title="ML Engineer")
    monkeypatch.setattr(ws, "structured", lambda *a: None)
    monkeypatch.setattr(cv_service, "analyze_jd", lambda text, llm: jd)
    monkeypatch.setattr(cv_service, "tailor", lambda cv, *a: apply_plan(cv, TailoringPlan(), jd))
    job = JobPosting(id="j3", title="ML Engineer", company="Orbit", description="Build ML.")
    tailored, path = asyncio.run(cv_service.tailor_to_job(ws, job))
    assert any(p.text.startswith("Example A.") for p in Document(str(path)).paragraphs)
    assert tailored.document_notes[0] == "Written in the design of Master_CV_2026.docx"
