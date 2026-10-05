"""CV writing quality: tighter no-fabrication guards (per-bullet skills, headline, summary),
restoring dropped keywords, matcher guidance, the critic -> revise pass, the ATS read-back
and guarded cover letters."""

from __future__ import annotations

import asyncio
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient

from src.cv.ats import check_docx
from src.cv.cover_letter import apply_letter, draft_letter, write_letter
from src.cv.docx_exporter import export_cover_letter, export_docx
from src.cv.models import (
    Bullet,
    BulletIssue,
    CoverLetterDraft,
    CVCritique,
    Experience,
    JDAnalysis,
    LetterOpening,
    LetterParagraph,
    LetterSections,
    MasterCV,
    MissedRequirement,
    RewrittenBullet,
    TailoredCVEdits,
    TailoringPlan,
)
from src.cv.tailor import apply_plan, tailor
from src.jobs.models import JobPosting, SearchIntent
from src.services import cover_letters, cv_service, intent, tailored_documents
from src.services.workspace import Workspace
from tests.conftest import FakeLLM

JD = JDAnalysis(
    job_title="Senior ML Engineer",
    company="Orbit AI",
    hard_skills=["Python", "PyTorch", "Kubernetes", "Spark"],
    must_have=["MLOps"],
)


def _rewrite(source_id: str, text: str) -> TailoringPlan:
    return TailoringPlan(rewritten_bullets=[RewrittenBullet(source_id=source_id, text=text)])


# ---------------------------------------------------------------- guards


def test_skill_from_another_bullet_is_rejected(master_cv: MasterCV) -> None:
    # Kubernetes is in the CV (nimbus-2), but nimbus-1 never used it.
    plan = _rewrite("nimbus-1", "Built a Kubernetes recommendation service serving 4M users.")
    change = apply_plan(master_cv, plan, JD).changes[0]
    assert not change.accepted and "Kubernetes" in (change.reason or "").title()


def test_headline_is_capped_at_the_level_and_roles_held(master_cv: MasterCV) -> None:
    def headline(text: str) -> Any:
        out = apply_plan(master_cv, TailoringPlan(headline=text), JD)
        return out.cv.basics.headline, out.changes[0]

    kept, change = headline("Staff Machine Learning Engineer")
    assert kept == master_cv.basics.headline and not change.accepted
    assert change.source_id == "headline" and "higher level (staff)" in (change.reason or "")
    kept, change = headline("Senior Data Scientist")
    assert kept == master_cv.basics.headline and "scientist" in (change.reason or "")
    kept, change = headline("Senior Machine Learning Engineer, NLP and recommenders")
    assert change.accepted and kept == "Senior Machine Learning Engineer, NLP and recommenders"


def test_summary_may_not_name_skills_the_cv_lacks(master_cv: MasterCV) -> None:
    plan = TailoringPlan(summary="ML engineer shipping PyTorch and Spark systems.")
    out = apply_plan(master_cv, plan, JD)
    assert out.cv.basics.summary == master_cv.basics.summary
    assert out.changes[0].source_id == "summary" and "spark" in (out.changes[0].reason or "")
    ok = apply_plan(master_cv, TailoringPlan(summary="ML engineer shipping PyTorch systems."), JD)
    assert ok.changes[0].accepted


def test_new_cv_claims_reject_stock_superlatives(master_cv: MasterCV) -> None:
    plan = TailoringPlan(
        headline="Visionary Senior Machine Learning Engineer",
        summary="World-class ML engineer with PyTorch experience.",
        rewritten_bullets=[
            RewrittenBullet(source_id="nimbus-1", text="Built a best-in-class recommender.")
        ],
    )
    out = apply_plan(master_cv, plan, JD)
    assert all(not change.accepted for change in out.changes)
    assert all("superlative" in (change.reason or "") for change in out.changes)
    assert out.cv.bullet_index()["nimbus-1"].text == master_cv.bullet_index()["nimbus-1"].text


def test_a_dropped_bullet_carrying_a_jd_keyword_is_put_back(master_cv: MasterCV) -> None:
    plan = TailoringPlan(bullet_order={"nimbus": ["nimbus-1", "nimbus-3"]})  # drops nimbus-2
    out = apply_plan(master_cv, plan, JD)
    assert [b.id for b in out.cv.experience[0].bullets] == ["nimbus-1", "nimbus-3", "nimbus-2"]
    assert out.restored_keywords == ["MLOps"]  # Kubernetes is still in the skills section
    assert out.missing_keywords == ["Spark"]  # the only true gap


def test_tailoring_keeps_two_bullets_per_role(master_cv: MasterCV) -> None:
    out = apply_plan(master_cv, TailoringPlan(bullet_order={"nimbus": ["nimbus-1"]}), JD)
    assert len(out.cv.experience[0].bullets) >= 2
    assert len({b.id for b in out.cv.experience[0].bullets}) == len(out.cv.experience[0].bullets)


# ---------------------------------------------------------------- guidance and review


def test_matcher_guidance_reaches_the_plan_and_review_drives_one_revision(
    master_cv: MasterCV,
) -> None:
    review = CVCritique(
        missed_requirements=[
            MissedRequirement(requirement="MLOps", source_id="nimbus-2"),
            MissedRequirement(requirement="Spark", source_id="ghost"),  # not a CV fact
        ],
        weak_bullets=[BulletIssue(source_id="nimbus-1", issue="result buried")],
    )
    plans = iter(
        [TailoringPlan(), TailoringPlan(bullet_order={"nimbus": ["nimbus-2", "nimbus-1"]})]
    )
    llm = FakeLLM({JDAnalysis: JD, TailoringPlan: lambda _: next(plans), CVCritique: review})
    out = tailor(master_cv, "JD text", llm, guidance="reasons: recsys at scale")
    prompts = [p for p, m in llm.calls if m is TailoringPlan]
    assert len(prompts) == 2 and "<match_assessment>\nreasons: recsys at scale" in prompts[0]
    assert "<review>" in prompts[1] and "ghost" not in prompts[1]
    assert out.critique == ["Surface MLOps (nimbus-2)", "Strengthen nimbus-1: result buried"]
    assert out.cv.experience[0].bullets[0].id == "nimbus-2"


def test_tailoring_emphasis_reaches_initial_and_revised_plans(master_cv: MasterCV) -> None:
    plans = iter([TailoringPlan(), TailoringPlan()])
    llm = FakeLLM({
        JDAnalysis: JD,
        TailoringPlan: lambda _: next(plans),
        CVCritique: CVCritique(order_notes=["Retain evidenced leadership in the summary"]),
    })
    tailor(master_cv, "JD text", llm, emphasis="leadership", level="senior")
    prompts = [prompt for prompt, model in llm.calls if model is TailoringPlan]
    assert len(prompts) == 2
    assert all("Foreground evidenced people" in prompt for prompt in prompts)
    assert all("do not claim a higher title" in prompt for prompt in prompts)


# ---------------------------------------------------------------- ATS read-back


def test_ats_check_reads_contact_keywords_and_length(master_cv: MasterCV, tmp_path: Path) -> None:
    path = export_docx(master_cv, tmp_path / "cv.docx")
    report = check_docx(path, master_cv, ["PyTorch", "Spark"])
    assert report.contact_missing == [] and report.missing_keywords == ["Spark"]
    assert report.keyword_coverage == 0.5 and 0 < report.est_pages < 2 and report.warnings == []
    doc = Document(str(path))
    doc.add_table(rows=1, cols=2)
    doc.save(str(path))
    assert any("table" in w for w in check_docx(path, master_cv, []).warnings)


# ---------------------------------------------------------------- cover letters


def test_cover_letter_keeps_only_backed_paragraphs(master_cv: MasterCV, tmp_path: Path) -> None:
    draft = CoverLetterDraft(
        paragraphs=[
            LetterParagraph(
                text="At Nimbus I built a PyTorch recommender serving 4M users.",
                source_ids=["nimbus-1"],
            ),
            LetterParagraph(text="Your team of 40 engineers ships weekly.", source_ids=[]),
            LetterParagraph(text="I cut costs by 50%.", source_ids=["nimbus-2"]),
            LetterParagraph(text="I ran Spark jobs daily.", source_ids=["nimbus"]),
            LetterParagraph(text="Proven leader.", source_ids=["ghost"]),
        ]
    )
    jd_text = "Join our team of 40 engineers."
    letter = apply_letter(master_cv, draft, JD, jd_text)
    assert letter.paragraphs == [
        "At Nimbus I built a PyTorch recommender serving 4M users.",
        "Your team of 40 engineers ships weekly.",  # a fact about the employer, from the JD
    ]
    reasons = [c.reason or "" for c in letter.changes if not c.accepted]
    assert "50%" in reasons[0] and "spark" in reasons[1] and "ghost" in reasons[2]
    path = export_cover_letter(letter, master_cv, tmp_path / "letter.docx")
    paragraphs = [p.text for p in Document(str(path)).paragraphs]
    text = "\n".join(paragraphs)
    assert paragraphs[0] == "Dear Hiring Manager,"
    assert master_cv.basics.email not in text
    assert master_cv.basics.headline not in text
    assert "4M users" in text and "Alex Example" in text and "50%" not in text


def test_letter_requires_role_introduction_and_one_page_length(master_cv: MasterCV) -> None:
    def generate(draft: LetterSections) -> None:
        llm = FakeLLM({LetterSections: draft})
        write_letter(master_cv, JD, "Senior ML Engineer role", llm)

    evidence = LetterParagraph(text="I built recommenders at Nimbus.", source_ids=["nimbus"])
    close = LetterParagraph(
        text="I would welcome a conversation about the role.", source_ids=["nimbus"]
    )
    with pytest.raises(ValueError, match="name the exact role"):
        generate(
            LetterSections(
                opening=LetterOpening(interest="I am interested in this work.", fit=evidence),
                evidence=[evidence, evidence],
                conclusion=close,
            )
        )
    opening = LetterOpening(
        interest="I am interested in the Senior ML Engineer role because of its ML work.",
        fit=LetterParagraph(text="My experience at Nimbus fits.", source_ids=["nimbus"]),
    )
    with pytest.raises(ValueError, match="under 300 words"):
        generate(
            LetterSections(
                opening=opening,
                evidence=[
                    evidence,
                    LetterParagraph(text="I built ML. " * 170, source_ids=["nimbus"]),
                ],
                conclusion=close,
            )
        )
    with pytest.raises(ValueError, match="cite CV source_ids"):
        generate(
            LetterSections(
                opening=opening,
                evidence=[evidence, LetterParagraph(text="I built recommenders.")],
                conclusion=close,
            )
        )
    with pytest.raises(ValueError, match="exactly two evidence paragraphs"):
        generate(LetterSections(opening=opening, evidence=[evidence], conclusion=close))


def test_letter_rejects_skill_not_evidenced_by_approved_cv(master_cv: MasterCV) -> None:
    jd = JD.model_copy(
        update={"job_title": "Principal Scientist, Functional Genomics"},
    )
    jd.hard_skills.append("Functional Genomics")
    draft = LetterSections(
        opening=LetterOpening(
            interest="The Functional Genomics Principal Scientist role appeals to me.",
            fit=LetterParagraph(
                text="My screening work is functional genomics in practice.",
                source_ids=["nimbus"],
            ),
        ),
        evidence=[
            LetterParagraph(text="I built recommenders at Nimbus.", source_ids=["nimbus-1"]),
            LetterParagraph(text="I moved training to Kubernetes.", source_ids=["nimbus-2"]),
        ],
        conclusion=LetterParagraph(
            text="I would welcome a conversation about your genomics work.", source_ids=["nimbus"]
        ),
    )
    with pytest.raises(ValueError, match="names a skill its sources do not evidence"):
        write_letter(master_cv, jd, "Functional genomics role", FakeLLM({LetterSections: draft}))


def test_letter_retries_rejected_section_with_specific_feedback(master_cv: MasterCV) -> None:
    opening = LetterOpening(
        interest="I am interested in the Senior ML Engineer role because of its ML work.",
        fit=LetterParagraph(text="My work at Nimbus fits.", source_ids=["nimbus"]),
    )
    good = LetterParagraph(text="I built recommenders at Nimbus.", source_ids=["nimbus"])
    kube = LetterParagraph(text="I moved training to Kubernetes.", source_ids=["nimbus-2"])
    unsupported = LetterParagraph(text="I led 40 engineers.", source_ids=["nimbus-1"])
    close = LetterParagraph(
        text="I would welcome a conversation about the role.", source_ids=["nimbus"]
    )
    drafts = [
        LetterSections(opening=opening, evidence=[good, unsupported], conclusion=close),
        LetterSections(opening=opening, evidence=[good, kube], conclusion=close),
    ]
    llm = FakeLLM({LetterSections: lambda prompt: drafts.pop(0)})

    letter = write_letter(master_cv, JD, "Senior ML Engineer role", llm)

    assert len(letter.paragraphs) == 4
    assert [model for _, model in llm.calls].count(LetterSections) == 2
    assert "evidence: introduces numbers" in llm.calls[1][0]
    assert "40" in llm.calls[1][0]


def test_letter_can_use_older_research_from_the_full_cv(
    master_cv: MasterCV,
) -> None:
    cv = master_cv.model_copy(deep=True)
    cv.experience.extend(
        [
            Experience(
                id="northfield",
                company="Northfield University",
                title="Postdoctoral Researcher",
                start="2016-06",
                end="2020-06",
                bullets=[
                    Bullet(
                        id="northfield-1",
                        text=(
                            "Imaged patient-derived organoids by confocal microscopy and measured "
                            "their growth with custom scripts."
                        ),
                    )
                ],
            ),
            Experience(
                id="riverside",
                company="Riverside Institute",
                title="Postdoctoral Researcher",
                start="2012-06",
                end="2016-06",
                bullets=[
                    Bullet(
                        id="riverside-1",
                        text="Prepared proteomics samples for mass spectrometry runs.",
                    )
                ],
            ),
        ]
    )
    opening = LetterOpening(
        interest="I am interested in the Senior Scientist role developing cell therapies.",
        fit=LetterParagraph(text="My research background fits.", source_ids=["northfield"]),
    )
    close = LetterParagraph(
        text="I would welcome a conversation about this work.", source_ids=["northfield"]
    )
    draft = LetterSections(
        opening=opening,
        evidence=[
            LetterParagraph(
                text=(
                    "At Northfield I imaged patient-derived organoids "
                    "by confocal microscopy and measured their growth with custom scripts."
                ),
                source_ids=["northfield-1"],
            ),
            LetterParagraph(
                text=(
                    "At Riverside I prepared proteomics samples for "
                    "mass spectrometry runs on the core facility instruments."
                ),
                source_ids=["riverside-1"],
            ),
        ],
        conclusion=close,
    )
    llm = FakeLLM({LetterSections: draft})
    jd = JDAnalysis(job_title="Senior Scientist", company="Cell Therapies")

    letter = write_letter(cv, jd, "Develop new cell therapies.", llm)

    assert "organoids" in letter.paragraphs[1]
    assert "proteomics" in letter.paragraphs[2]
    draft_prompts = [prompt for prompt, model in llm.calls if model is LetterSections]
    assert len(draft_prompts) == 1
    assert "northfield-1" in draft_prompts[0] and "riverside-1" in draft_prompts[0]


def test_letter_accepts_valid_draft_without_an_evidence_review(master_cv: MasterCV) -> None:
    draft = LetterSections(
        opening=LetterOpening(
            interest="I am interested in the Senior ML Engineer role because of its ML work.",
            fit=LetterParagraph(text="My work at Nimbus fits.", source_ids=["nimbus"]),
        ),
        evidence=[
            LetterParagraph(text="I built recommenders at Nimbus.", source_ids=["nimbus-1"]),
            LetterParagraph(text="I improved recommendations at Nimbus.", source_ids=["nimbus-1"]),
        ],
        conclusion=LetterParagraph(
            text="I would welcome a conversation about the role.", source_ids=["nimbus"]
        ),
    )
    llm = FakeLLM({LetterSections: draft})

    letter = write_letter(master_cv, JD, "Senior ML Engineer role", llm)

    assert len(letter.paragraphs) == 4
    assert [model for _, model in llm.calls] == [LetterSections]


def test_cover_letter_rejects_unsourced_personal_claims_and_job_numbers(
    master_cv: MasterCV,
) -> None:
    draft = CoverLetterDraft(
        paragraphs=[
            LetterParagraph(text="I managed multiple teams.", source_ids=[]),
            LetterParagraph(text="I led 40 engineers.", source_ids=["nimbus-1"]),
            LetterParagraph(text="I am a proven leader.", source_ids=["nimbus"]),
            LetterParagraph(text="Your team has 40 engineers.", source_ids=[]),
        ]
    )
    letter = apply_letter(master_cv, draft, JD, "Your team has 40 engineers.")
    assert letter.paragraphs == ["Your team has 40 engineers."]
    assert "source_id" in (letter.changes[0].reason or "")
    assert "40" in (letter.changes[1].reason or "")
    assert "superlative" in (letter.changes[2].reason or "")


def test_letter_motivation_comes_only_from_the_career_intent(
    master_cv: MasterCV, settings: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    ws = Workspace(settings)
    ws.master_cv, ws.active_cv_id = master_cv, "master"
    intent.save_intent(ws, SearchIntent(direction="Lead applied ML in health"))
    draft = LetterSections(
        opening=LetterOpening(
            interest="I am interested in the Senior ML Engineer role because of its ML work.",
            fit=LetterParagraph(
                text="My experience building recommenders at Nimbus fits.", source_ids=["nimbus"]
            ),
        ),
        evidence=[
            LetterParagraph(text="I built recommenders at Nimbus.", source_ids=["nimbus"]),
            LetterParagraph(text="I moved training to Kubernetes.", source_ids=["nimbus-2"]),
        ],
        conclusion=LetterParagraph(
            text="I would welcome a conversation about the role.", source_ids=["nimbus"]
        ),
    )
    llm = FakeLLM(
        {JDAnalysis: JD, LetterSections: draft}
    )
    monkeypatch.setattr(ws, "structured", lambda *a, **k: llm)
    job = JobPosting(id="j", title="Senior ML Engineer", company="Orbit AI", description="ML")
    letter, path = asyncio.run(cv_service.write_cover_letter(ws, job))
    assert path.name == "Alex_Example_Orbit_AI_Senior_ML_Engineer_cover_letter.docx"
    assert letter.paragraphs
    assert letter.paragraphs[0].startswith("I am interested in the Senior ML Engineer role")
    assert "My experience building recommenders" in letter.paragraphs[0]
    prompt = next(p for p, m in llm.calls if m is LetterSections)
    assert "<motivation>\nLead applied ML in health\n</motivation>" in prompt
    assert draft_letter  # the drafting step is the only LLM call besides the JD analysis


def test_saved_tailored_cv_edits_feed_letter_after_restart(
    master_cv: MasterCV, settings: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    ws = Workspace(settings)
    ws.master_cv, ws.active_cv_id = master_cv, "master"
    job = JobPosting(id="saved-role", title="Senior ML Engineer", company="Orbit AI")
    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    original_path = export_docx(tailored, settings.output_dir / "first.docx")
    document = tailored_documents.create(
        ws, job.id, master_cv, JD, tailored, "classic", original_path, posting=job
    )
    new_text = "Built a recommendation service in Python and PyTorch serving 4M users."
    with pytest.raises(ValueError, match="40"):
        tailored_documents.edit(
            ws, document.id, TailoredCVEdits(bullets={"nimbus-1": "I led 40 engineers."})
        )
    edited = tailored_documents.edit(
        ws, document.id, TailoredCVEdits(bullets={"nimbus-1": new_text})
    )
    assert edited.filename != document.filename
    assert edited.tailored.cv.bullet_index()["nimbus-1"].text == new_text

    restarted = Workspace(settings)
    restarted.active_cv_id = None
    assert tailored_documents.list_for_job(restarted, job.id)[0].id == document.id
    with pytest.raises(ValueError, match="does not belong"):
        tailored_documents.load(restarted, document.id, "another-role")
    draft = LetterSections(
        opening=LetterOpening(
            interest="I am interested in the Senior ML Engineer role because of its ML work.",
            fit=LetterParagraph(text="My recommendation work fits.", source_ids=["nimbus"]),
        ),
        evidence=[
            LetterParagraph(text=new_text, source_ids=["nimbus-1"]),
            LetterParagraph(text="I moved training to Kubernetes.", source_ids=["nimbus-2"]),
        ],
        conclusion=LetterParagraph(
            text="I would welcome a conversation about the role.", source_ids=["nimbus-1"]
        ),
    )
    llm = FakeLLM(
        {JDAnalysis: JD, LetterSections: draft}
    )
    monkeypatch.setattr(restarted, "structured", lambda *a, **k: llm)
    letter, path = asyncio.run(
        cv_service.write_cover_letter(restarted, job, tailored_cv_id=document.id)
    )
    assert new_text in letter.paragraphs[1] and path.exists()
    assert new_text in next(prompt for prompt, model in llm.calls if model is LetterSections)
    from src.web import app as webapp

    monkeypatch.setattr(webapp, "get_workspace", lambda: restarted)
    client = TestClient(webapp.app)
    assert any(item["id"] == document.id for item in client.get("/api/tailored-cvs").json())
    response = client.post(
        f"/api/jobs/{job.id}/cover-letter", json={"tailored_cv_id": document.id}
    )
    assert response.status_code == 200, response.json()


def test_older_word_cv_requires_review_before_it_can_feed_letter(
    master_cv: MasterCV, settings: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    ws = Workspace(settings)
    ws.master_cv, ws.active_cv_id = master_cv, "master"
    job = JobPosting(id="older-role", title="Senior ML Engineer", company="Orbit AI")
    changed = master_cv.model_copy(deep=True)
    changed.bullet_index()[
        "nimbus-1"
    ].text = "Built a recommendation service in Python and PyTorch serving 4M users."
    source = export_docx(changed, settings.output_dir / "older_tailored.docx")
    draft = LetterSections(
        opening=LetterOpening(
            interest="I am interested in the Senior ML Engineer role because of its ML work.",
            fit=LetterParagraph(text="My recommendation work fits.", source_ids=["nimbus"]),
        ),
        evidence=[
            LetterParagraph(text=changed.bullet_index()["nimbus-1"].text, source_ids=["nimbus-1"]),
            LetterParagraph(text="I moved training to Kubernetes.", source_ids=["nimbus-2"]),
        ],
        conclusion=LetterParagraph(
            text="I would welcome a conversation about the role.", source_ids=["nimbus"]
        ),
    )
    llm = FakeLLM(
        {
            MasterCV: changed,
            JDAnalysis: JD,
            LetterSections: draft,
        }
    )
    monkeypatch.setattr(ws, "structured", lambda *a, **k: llm)

    imported = asyncio.run(cv_service.attach_existing_cv(ws, job, f"generated:{source.name}"))
    assert imported.imported and not imported.reviewed
    assert changed.bullet_index()["nimbus-1"].text in llm.calls[0][0]
    with pytest.raises(ValueError, match="Review and save"):
        asyncio.run(cv_service.write_cover_letter(ws, job, tailored_cv_id=imported.id))
    tailored_documents.edit(ws, imported.id, TailoredCVEdits())

    restarted = Workspace(settings)
    monkeypatch.setattr(restarted, "structured", lambda *a, **k: llm)
    assert tailored_documents.list_for_job(restarted, job.id)[0].reviewed
    letter, _ = asyncio.run(
        cv_service.write_cover_letter(restarted, job, tailored_cv_id=imported.id)
    )
    assert changed.bullet_index()["nimbus-1"].text in letter.paragraphs[1]

    from src.web import app as webapp

    monkeypatch.setattr(webapp, "get_workspace", lambda: ws)
    client = TestClient(webapp.app)
    attached = client.post(
        "/api/tailored-cvs/import",
        json={
            "asset_id": f"generated:{source.name}",
            "title": job.title,
            "company": job.company,
            "description": "Build recommendation systems with Python and PyTorch.",
        },
    )
    assert attached.status_code == 200, attached.json()
    linked = attached.json()
    assert linked["job_id"].startswith("linked:") and not linked["reviewed"]
    assert client.put(f"/api/tailored-cvs/{linked['id']}", json={}).json()["reviewed"]
    response = client.post(
        f"/api/jobs/{linked['job_id']}/cover-letter", json={"tailored_cv_id": linked["id"]}
    )
    assert response.status_code == 200, response.json()


def test_general_cv_exports_all_roles_and_keeps_prior_versions(
    master_cv: MasterCV, settings: Any
) -> None:
    ws = Workspace(settings)
    ws.master_cv, ws.active_cv_id = master_cv, "master"
    first, roles = asyncio.run(cv_service.export_general_cv(ws))
    second, _ = asyncio.run(cv_service.export_general_cv(ws))
    assert roles == len(master_cv.experience)
    assert first != second and first.exists() and second.exists()
    body = "\n".join(p.text for p in Document(str(first)).paragraphs)
    assert all(role.title in body for role in master_cv.experience)


def test_cover_letters_have_their_own_library_editor_and_exports(
    master_cv: MasterCV, settings: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    from src.web import app as webapp

    ws = Workspace(settings)
    ws.master_cv, ws.active_cv_id = master_cv, "master"
    cv_path, _ = asyncio.run(cv_service.export_general_cv(ws))
    assert cv_path.parent == settings.output_dir / "cvs"
    job = JobPosting(id="letter-job", title="Senior ML Engineer", company="Orbit AI")
    draft = CoverLetterDraft(
        paragraphs=[
            LetterParagraph(text="My work at Nimbus fits this role.", source_ids=["nimbus"]),
            LetterParagraph(text="I built recommenders at Nimbus.", source_ids=["nimbus-1"]),
            LetterParagraph(text="I moved training to Kubernetes.", source_ids=["nimbus-2"]),
            LetterParagraph(text="I would welcome a conversation.", source_ids=["nimbus"]),
        ]
    )
    letter = apply_letter(master_cv, draft, JD, "Senior ML Engineer role")
    letter_path = export_cover_letter(
        letter, master_cv, settings.output_dir / "cover_letters" / "example_cover_letter.docx"
    )
    saved = cover_letters.create(
        ws, job, letter, master_cv, "Senior ML Engineer role", "classic", letter_path
    )
    legacy = export_cover_letter(
        letter, master_cv, settings.output_dir / "older_cover_letter.docx"
    )
    older_cv = export_docx(master_cv, settings.output_dir / "older_cv.docx")
    assert legacy.exists() and older_cv.exists()
    assets = cv_service.list_cvs(ws)
    assert all("cover_letter" not in asset.filename for asset in assets)
    assert any(asset.id == "generated:older_cv.docx" for asset in assets)
    assert list(settings.output_dir.glob("*.docx")) == []
    assert (settings.output_dir / "cvs" / older_cv.name).exists()
    assert (settings.output_dir / "cover_letters" / legacy.name).exists()
    assert {item.filename for item in cover_letters.list_all(ws)} == {
        letter_path.name,
        legacy.name,
    }

    monkeypatch.setattr(webapp, "get_workspace", lambda: ws)
    client = TestClient(webapp.app)
    assert client.get(f"/api/files/cvs/{cv_path.name}").status_code == 200
    assert client.get(f"/api/files/{legacy.name}").status_code == 200
    listed = client.get("/api/cover-letters").json()
    assert len(listed) == 2
    assert next(item for item in listed if item["id"] == saved.id)["paragraphs"][1] == (
        "I built recommenders at Nimbus."
    )
    edits = {
        "greeting": letter.greeting,
        "paragraphs": [*letter.paragraphs],
        "closing": letter.closing,
    }
    edits["paragraphs"][1] = "I built recommenders at Nimbus and led 40 engineers."
    assert client.put(f"/api/cover-letters/{saved.id}", json=edits).status_code == 422
    edits["paragraphs"][1] = "At Nimbus I built recommenders."
    updated = client.put(f"/api/cover-letters/{saved.id}", json=edits)
    assert updated.status_code == 200, updated.json()
    assert updated.json()["paragraphs"][1] == "At Nimbus I built recommenders."
    assert client.get(updated.json()["docx_url"]).status_code == 200
    text_export = client.get(updated.json()["txt_url"])
    assert text_export.status_code == 200
    assert b"At Nimbus I built recommenders." in text_export.content
    text_path = (settings.output_dir / "cover_letters" / updated.json()["filename"]).with_suffix(
        ".txt"
    )
    assert text_path.exists()
    legacy_view = next(item for item in listed if item["filename"] == legacy.name)
    legacy_edit = client.put(
        f"/api/cover-letters/{legacy_view['id']}",
        json={
            "greeting": legacy_view["greeting"],
            "paragraphs": legacy_view["paragraphs"],
            "closing": "Kind regards,",
        },
    )
    assert legacy_edit.status_code == 200, legacy_edit.json()
    assert legacy_edit.json()["closing"] == "Kind regards,"
    assert len(cover_letters.list_all(Workspace(settings))) == 2
    assert all("cover_letter" not in asset.filename for asset in cv_service.list_cvs(ws))

    deleted = client.delete(f"/api/cover-letters/{saved.id}")
    assert deleted.status_code == 200 and deleted.json() == {"deleted": 1}
    assert not letter_path.exists()
    assert not text_path.exists()
    assert not (settings.output_dir / "cover_letters" / updated.json()["filename"]).exists()
    assert len(client.get("/api/cover-letters").json()) == 1
    assert client.delete(f"/api/cover-letters/{saved.id}").status_code == 404

    extra_path = export_cover_letter(
        letter, master_cv, settings.output_dir / "cover_letters" / "another_cover_letter.docx"
    )
    cover_letters.create(
        ws, job, letter, master_cv, "Senior ML Engineer role", "classic", extra_path
    )
    removed_all = client.delete("/api/cover-letters")
    assert removed_all.status_code == 200 and removed_all.json() == {"deleted": 2}
    assert client.get("/api/cover-letters").json() == []
    assert client.delete("/api/cover-letters").json() == {"deleted": 0}
    assert not extra_path.exists()
    assert not (settings.output_dir / "cover_letters" / legacy.name).exists()
    assert not (settings.output_dir / "cover_letters" / legacy_edit.json()["filename"]).exists()
    assert older_cv.name in {path.name for path in (settings.output_dir / "cvs").glob("*.docx")}
