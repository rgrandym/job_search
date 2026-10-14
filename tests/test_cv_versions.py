"""Tailored CV versions for one job, read side by side against the original."""

from __future__ import annotations

import asyncio
import re
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from src.core.config import Settings
from src.cv.fit_check import (
    NOISE_MARGIN,
    anchored_score,
    best_version,
    compare_fit,
    consensus,
    read_versions,
    requirements_for,
    verdict,
    verify,
)
from src.cv.models import (
    EvidencedReading,
    EvidencedReadings,
    FitReading,
    JDAnalysis,
    MasterCV,
    MetRequirement,
    RatedLine,
    RequirementEvidence,
    ReviewAddition,
    SearchFitBaseline,
    TailoredDocument,
    TailoringPlan,
    TailoringSettings,
    VersionReading,
)
from src.cv.tailor import apply_plan, printed_text
from src.jobs.models import FitRatings, JobAssessment, JobPosting, MatchResult
from src.jobs.screener import finalize
from src.services import cv_service, cv_versions, tailored_documents
from src.services.workspace import Workspace
from tests.conftest import FakeLLM

JD = JDAnalysis(
    job_title="Senior ML Engineer",
    company="Orbit AI",
    hard_skills=["Python", "PyTorch", "Azure"],
    must_have=["MLOps", "Python"],
)
JOB = JobPosting(id="job-1", title="Senior ML Engineer", company="Orbit AI", description="MLOps")
CV_BLOCK = re.compile(r"<cv_(\d+)>\n(.*?)\n</cv_\1>", re.DOTALL)
WEAK_RATINGS = {"function": 2, "domain": 2, "seniority": 2, "leadership": 2, "sector": 2}
STRONG_RATINGS = {"function": 4, "domain": 2, "seniority": 3, "leadership": 1, "sector": 2}


def _evidenced(
    text: str, ratings: dict[str, int], met: list[int], impression: str
) -> EvidencedReading:
    """A reading whose quotes are real lines of `text` (so code accepts them)."""
    line = next(ln for ln in text.splitlines() if len(ln.split()) >= 2)
    return EvidencedReading(
        **{name: RatedLine(rating=r, quote=line) for name, r in ratings.items()},
        requirements_met=[MetRequirement(number=n, quote=line) for n in met],
        impression=impression,
    )


def _reader(prompt: str) -> EvidencedReadings:
    """The role-skills line marks the strong version; every other CV reads weak."""
    out = []
    for _, text in CV_BLOCK.findall(prompt):
        strong = "Relevant skills for this role: PyTorch" in text
        ratings, met = (STRONG_RATINGS, [1, 2, 7]) if strong else (WEAK_RATINGS, [1])
        out.append(_evidenced(text, ratings, met, "clear fit" if strong else "generic"))
    return EvidencedReadings(readings=out)


def _version(
    ws: Workspace, master_cv: MasterCV, settings: TailoringSettings, **plan: object
) -> TailoredDocument:
    tailored = apply_plan(master_cv, TailoringPlan.model_validate(plan), JD)
    path = ws.output_dir / "cvs" / f"{len(tailored_documents.list_all(ws))}.docx"
    return tailored_documents.create(
        ws, JOB.id, master_cv, JD, tailored, "original", path, posting=JOB, settings=settings
    )


@pytest.fixture
def ws(settings: Settings, master_cv: MasterCV, monkeypatch: pytest.MonkeyPatch) -> Workspace:
    settings.master_cv_path.write_text(master_cv.model_dump_json())  # the library's Master CV
    workspace = Workspace(settings)
    llm = FakeLLM({EvidencedReadings: _reader})
    monkeypatch.setattr(workspace, "structured", lambda *_args, **_kw: llm)
    return workspace


def test_versions_are_read_side_by_side_and_set_against_their_original(
    ws: Workspace, master_cv: MasterCV
) -> None:
    senior = _version(ws, master_cv, TailoringSettings(level="senior"), role_skills=["PyTorch"])
    auto = _version(ws, master_cv, TailoringSettings())
    comparison = asyncio.run(cv_versions.compare(ws, JOB.id, []))
    by_id = {r.id: r for r in comparison.readings}
    assert [r.kind for r in comparison.readings] == ["original", "tailored", "tailored"]
    assert by_id[senior.id].label == "V1 · Level: more senior"
    assert by_id[auto.id].label == "V2 · Automatic"
    assert comparison.best_id == senior.id
    assert comparison.summary.startswith("V1 · Level: more senior shows the fit most clearly")
    areas = {a.area: a for a in by_id[senior.id].areas}
    assert (areas["Function"].before, areas["Function"].after) == ("2/4", "4/4")
    assert areas["Function"].change == "better" and areas["Leadership"].change == "worse"
    assert areas["Requirements shown"].after == "2 of 2"  # 7 is not a listed requirement
    assert areas["Requirement: Python"].change == "better"
    assert all(a.change == "same" for a in by_id[auto.id].areas if a.area != "Fit score")


def test_saved_comparison_goes_stale_when_a_version_changes(
    ws: Workspace, master_cv: MasterCV
) -> None:
    assert cv_versions.saved(ws, JOB.id) == (None, False)
    version = _version(ws, master_cv, TailoringSettings())
    asyncio.run(cv_versions.compare(ws, JOB.id, [version.id]))
    saved, stale = cv_versions.saved(ws, JOB.id)
    assert saved is not None and not stale
    later = version.model_copy(update={"updated_at": "9999-01-01T00:00:00+00:00"})
    tailored_documents._save(ws, later)
    assert cv_versions.saved(ws, JOB.id)[1]


def test_old_comparison_method_is_marked_stale(ws: Workspace, master_cv: MasterCV) -> None:
    version = _version(ws, master_cv, TailoringSettings())
    comparison = asyncio.run(cv_versions.compare(ws, JOB.id, [version.id]))
    old = comparison.model_copy(update={"method_version": 1})
    cv_versions._path(ws, JOB.id).write_text(old.model_dump_json())
    assert cv_versions.saved(ws, JOB.id)[1]


def test_too_many_versions_or_a_foreign_one_is_refused(ws: Workspace, master_cv: MasterCV) -> None:
    for _ in range(6):
        _version(ws, master_cv, TailoringSettings())
    with pytest.raises(ValueError, match="at most 6"):
        asyncio.run(cv_versions.compare(ws, JOB.id, []))
    with pytest.raises(ValueError, match="does not belong"):
        asyncio.run(cv_versions.compare(ws, JOB.id, ["0" * 24]))


def test_one_original_from_the_library_never_a_deleted_cv(
    ws: Workspace, master_cv: MasterCV
) -> None:
    from src.services import cv_service

    main = cv_service.store_cv(ws, "main.txt", b"Main CV\nML engineer")
    other = cv_service.store_cv(ws, "vertex.txt", b"Vertex CV\nPython engineer")
    ws.active_cv_id = other.id
    from_other = _version(ws, master_cv, TailoringSettings())  # tailored from the other CV
    cv_service.select_cv(ws, main.id)
    _version(ws, master_cv, TailoringSettings())
    comparison = asyncio.run(cv_versions.compare(ws, JOB.id, []))
    originals = [r for r in comparison.readings if r.kind == "original"]
    assert [r.id for r in originals] == [f"original:{main.id}"]  # the selected CV, only it
    assert all(r.baseline_id == originals[0].id for r in comparison.readings[1:])

    picked = asyncio.run(cv_versions.compare(ws, JOB.id, [from_other.id], baseline_id=other.id))
    assert picked.readings[0].label == "Original · vertex.txt"
    assert not cv_versions.saved(ws, JOB.id)[1]
    cv_service.delete_cv(ws, other.id)
    assert cv_versions.saved(ws, JOB.id)[1]  # its original is gone: compare again
    with pytest.raises(ValueError, match="Choose a CV from your library"):
        asyncio.run(cv_versions.compare(ws, JOB.id, [], baseline_id=other.id))


def test_comparing_again_gives_the_same_readings_without_rereading(
    ws: Workspace, master_cv: MasterCV
) -> None:
    llm = ws.structured("cv")
    assert isinstance(llm, FakeLLM)
    senior = _version(ws, master_cv, TailoringSettings(level="senior"), role_skills=["PyTorch"])
    first = asyncio.run(cv_versions.compare(ws, JOB.id, [senior.id]))
    assert len(llm.calls) == 3  # three independent reads
    again = asyncio.run(cv_versions.compare(ws, JOB.id, [senior.id]))
    assert len(llm.calls) == 3  # the same CVs are not read (and re-scored) again
    assert again.readings == first.readings and again.best_id == first.best_id

    auto = _version(ws, master_cv, TailoringSettings(), role_skills=["Python"])
    added = asyncio.run(cv_versions.compare(ws, JOB.id, [senior.id, auto.id]))
    assert len(llm.calls) == 6  # the new version needed reading; all were shown again
    by_id = {r.id: r for r in added.readings}
    assert by_id[added.readings[0].id].score == first.readings[0].score  # original kept
    assert by_id[senior.id].score == {r.id: r for r in first.readings}[senior.id].score


def test_a_reader_that_skips_a_cv_is_an_error() -> None:
    one = _evidenced("Python engineer", WEAK_RATINGS, [], "")
    llm = FakeLLM({EvidencedReadings: EvidencedReadings(readings=[one])})
    with pytest.raises(ValueError, match="1 readings for 2 CVs"):
        read_versions(["one", "two"], ["Python"], "job", llm)


def test_api_lists_labels_and_compares(
    ws: Workspace, master_cv: MasterCV, monkeypatch: pytest.MonkeyPatch
) -> None:
    from src.web import app as webapp

    monkeypatch.setattr(webapp, "get_workspace", lambda: ws)
    client = TestClient(webapp.app)
    version = _version(ws, master_cv, TailoringSettings(length="full"))
    listed = client.get(f"/api/jobs/{JOB.id}/tailored-cvs").json()
    assert listed[0]["label"] == "V1 · Full length" and listed[0]["settings"]["length"] == "full"
    assert client.get(f"/api/jobs/{JOB.id}/cv-comparison").json() == {
        "comparison": None,
        "stale": False,
        "baseline": "master",
    }
    body = client.post(f"/api/jobs/{JOB.id}/cv-comparison", json={"document_ids": [version.id]})
    assert body.status_code == 200 and body.json()["readings"][1]["id"] == version.id
    assert client.get(f"/api/jobs/{JOB.id}/cv-comparison").json()["comparison"] is not None
    refused = client.post("/api/jobs/other/cv-comparison", json={"document_ids": []})
    assert refused.status_code == 422


def test_a_rating_or_requirement_counts_only_with_a_quote_found_in_that_cv() -> None:
    cv = "Jane Doe\nLed a team of 6 ML engineers shipping PyTorch models\nPython, SQL"
    raw = EvidencedReading(
        function=RatedLine(rating=4, quote="led a team of 6 ML engineers... PyTorch models"),
        domain=RatedLine(rating=4, quote="Built the company's MLOps platform"),  # not in the CV
        seniority=RatedLine(rating=1, quote=""),
        leadership=RatedLine(rating=3, quote="Python"),  # one word is not evidence
        sector=RatedLine(rating=0, quote=""),
        requirements_met=[
            MetRequirement(number=1, quote="Shipping PyTorch models"),
            MetRequirement(number=2, quote="Ran MLOps on Azure"),  # invented
            MetRequirement(number=9, quote="Python, SQL"),  # not a listed requirement
        ],
    )
    reading = verify(raw, cv, total=2)
    assert (reading.function, reading.domain, reading.leadership) == (4, 1, 1)
    assert reading.requirements_shown == [1]
    assert set(reading.evidence) == {"function"}


def test_the_median_of_the_reads_counts() -> None:
    def run(function: int, shown: list[int]) -> FitReading:
        return FitReading(function=function, domain=2, seniority=2, leadership=2, sector=2,
                          requirements_shown=shown,
                          requirement_quotes={str(n): "a line" for n in shown})  # fmt: skip

    merged = consensus([run(4, [1, 2]), run(2, [1]), run(3, [1, 3])])
    assert merged.function == 3  # one high and one low read do not move it
    assert merged.requirements_shown == [1]  # found by most reads


def test_verdicts_and_setting_effects_use_fixed_margins(ws: Workspace, master_cv: MasterCV) -> None:
    auto = _version(ws, master_cv, TailoringSettings())
    senior = _version(ws, master_cv, TailoringSettings(level="senior"), role_skills=["PyTorch"])
    both = _version(
        ws, master_cv, TailoringSettings(level="junior", length="full"), role_skills=["PyTorch"]
    )
    comparison = asyncio.run(cv_versions.compare(ws, JOB.id, []))
    by_id = {r.id: r for r in comparison.readings}
    assert by_id[auto.id].verdict == "same"  # reads as the original does
    assert by_id[auto.id].verdict_reason.startswith("No meaningful change")
    assert by_id[senior.id].verdict == "better"
    assert "1 more requirement shown" in by_id[senior.id].verdict_reason
    assert comparison.reads == 3 and comparison.margin == NOISE_MARGIN
    effects = comparison.setting_effects
    assert len(effects) == 1  # the junior+full version differs in two settings: not attributed
    level = effects[0]
    assert (level.setting, level.from_value, level.to_value) == (
        "Level",
        "Automatic",
        "Level: more senior",
    )
    assert level.effect == "clear" and level.score_change > 0
    assert "seniority 2/4 → 3/4" in level.text
    assert by_id[both.id].verdict == "better"


def test_equal_scores_are_decided_by_checked_requirements_and_say_so() -> None:
    def cv(label: str, score: int, shown: list[int], keywords: list[str]) -> VersionReading:
        return VersionReading(id=label, kind="tailored", label=label, baseline_id="original:m",
                              score=score, requirements_shown=shown, keywords_found=keywords,
                              keyword_coverage=len(keywords) / 11)  # fmt: skip

    original = cv("Original", 42, [1, 2, 6, 13], []).model_copy(update={"kind": "original"})
    v1 = cv("V1", 50, [1, 2, 6], ["Project management"])
    v2 = cv("V2", 50, [1, 2, 6, 13], [])
    best, summary = best_version([original, v2, v1], total=15)
    assert best == "V2"
    assert "score about the same (50 vs 50/100" in summary
    assert "more of the job's requirements (4 vs 3 of 15)" in summary
    assert "V1 has more job keywords (1 vs 0: Project management)" in summary
    # Keyword differences remain visible, but they do not decide overall fit alone.
    judged, reason = verdict(v1, original)
    assert judged == "mixed" and "job keyword added: Project management" in reason


def test_keyword_loss_alone_does_not_make_the_tailored_cv_worse() -> None:
    original = VersionReading(
        id="original:cv",
        kind="original",
        label="Original",
        score=86,
        requirements_shown=[1, 2],
        keywords_found=["Technology transfer"],
    )
    tailored = VersionReading(
        id="tailored",
        kind="tailored",
        label="Tailored",
        baseline_id=original.id,
        score=86,
        requirements_shown=[1, 2],
        keywords_found=[],
    )
    judged, reason = verdict(tailored, original)
    assert judged == "same" and "Technology transfer" in reason
    best, summary = best_version([original, tailored], 2)
    assert best is None and "No version shows a measurable improvement" in summary


def test_search_score_is_the_exact_baseline_and_only_cv_dimensions_can_move_it() -> None:
    original = FitReading(function=2, domain=3, seniority=3, leadership=2, sector=3)
    improved = original.model_copy(update={"function": 3})
    baseline = SearchFitBaseline(
        score=89,
        ratings={
            "function": 3,
            "domain": 4,
            "seniority": 3,
            "leadership": 3,
            "sector": 4,
            "practicality": 2,
        },
    )
    assert anchored_score(baseline, original, original) == 89
    assert anchored_score(baseline, original, improved) == 96
    assert anchored_score(baseline, original, original.model_copy(update={"function": 1})) == 81
    blocked = baseline.model_copy(update={"score": 55, "cap_reason": "core requirement unmet: GMP"})
    assert anchored_score(blocked, original, improved) == 55


def test_unmet_core_cap_lifts_only_for_a_cv_that_newly_shows_it_with_a_quote() -> None:
    ratings = {"function": 3, "domain": 3, "seniority": 3, "leadership": 3, "sector": 3}
    blocked = SearchFitBaseline(
        score=55,
        ratings={**ratings, "practicality": 3},
        cap_reason="core requirement unmet: GMP manufacturing; Cell therapy",
    )
    requirements = ["GMP manufacturing", "Cell therapy", "Python"]
    original = FitReading(**ratings, requirements_shown=[3])
    shows = FitReading(
        **ratings,
        requirements_shown=[1, 2, 3],
        requirement_quotes={"1": "Ran GMP batches", "2": "CAR-T process", "3": "Python"},
    )
    # The same levels uncapped: 77. The cap lifts for the CV that now shows both.
    assert anchored_score(blocked, original, shows, requirements) == 77
    one = shows.model_copy(update={"requirements_shown": [1, 3]})
    assert anchored_score(blocked, original, one, requirements) == 55  # one still unmet
    unquoted = shows.model_copy(update={"requirement_quotes": {}})
    assert anchored_score(blocked, original, unquoted, requirements) == 55
    already = original.model_copy(update={"requirements_shown": [1, 2, 3]})
    assert anchored_score(blocked, already, shows, requirements) == 55  # nothing newly shown
    assert anchored_score(blocked, original, shows) == 55  # requirements not read


def test_core_requirements_the_search_found_unmet_are_read_first(master_cv: MasterCV) -> None:
    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    blocked = SearchFitBaseline(
        score=55,
        ratings={**WEAK_RATINGS, "practicality": 2},
        cap_reason="core requirement unmet: Python; Model serving",
    )
    assert requirements_for(tailored, JD, blocked) == ["Python", "Model serving", "MLOps"]


def test_proposals_that_show_the_unmet_core_requirement_lift_the_projection(
    master_cv: MasterCV,
) -> None:
    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    tailored.review_additions = [HANDOVER]  # quoted by the reader for requirement 2
    blocked = SearchFitBaseline(
        score=55,
        ratings={**STRONG_RATINGS, "practicality": 3},
        cap_reason="core requirement unmet: Python",
    )
    check = compare_fit(
        master_cv, tailored, JD, "MLOps and Python", FakeLLM({EvidencedReadings: _swapped}),
        blocked,
    )  # fmt: skip
    projection = check.if_confirmed
    assert check.before == check.after == 55  # the factual draft does not show it
    assert projection is not None and projection.confirmed == 55
    # The proposal shows the core requirement: only the matcher's weak-leadership cap remains.
    assert projection.projected == 65


def test_capped_search_scores_move_as_the_matcher_would_move_them() -> None:
    original = FitReading(function=3, domain=3, seniority=3, leadership=3, sector=3)
    weaker = original.model_copy(update={"function": 2})
    # raw 79, capped at 65 for weak practicality: one level less on function is raw 71,
    # which the matcher still caps at 65, so the version is not shown as worse.
    weak = SearchFitBaseline(
        score=65,
        ratings={
            "function": 3,
            "domain": 4,
            "seniority": 4,
            "leadership": 3,
            "sector": 3,
            "practicality": 0,
        },
        cap_reason="weak on practicality",
    )
    assert anchored_score(weak, original, weaker) == 65
    # an unread posting below its cap has no cap reason; the cap still applies upward
    unread = SearchFitBaseline(
        score=60,
        ratings={
            "function": 2,
            "domain": 3,
            "seniority": 3,
            "leadership": 2,
            "sector": 2,
            "practicality": 2,
        },
        requirements_checked=False,
    )
    better = original.model_copy(update={"function": 4, "domain": 4})
    assert anchored_score(unread, original, better) == 62


def test_matcher_and_fit_check_share_one_capped_total() -> None:
    from src.core.fit_rubric import capped_score, points_for

    levels = {
        "function": 3,
        "domain": 3,
        "seniority": 3,
        "leadership": 3,
        "sector": 3,
        "practicality": 0,
    }
    assessment = JobAssessment(job_id="j", fit_summary="x", ratings=FitRatings(**levels))
    assert finalize(assessment, threshold=60).fit_score == capped_score(points_for(levels)) == 65
    unread = finalize(assessment, threshold=60, checked=False)
    baseline = cv_service.search_fit_baseline(MatchResult(job=JOB, verdict=unread), "master")
    assert baseline is not None and baseline.score == 62 and not baseline.requirements_checked


def test_comparison_reuses_saved_search_score_for_the_source_cv(
    ws: Workspace, master_cv: MasterCV, monkeypatch: pytest.MonkeyPatch
) -> None:
    version = _version(ws, master_cv, TailoringSettings(), role_skills=["PyTorch"])
    assessment = JobAssessment(
        job_id=JOB.id,
        fit_summary="Relevant ML work",
        ratings=FitRatings(
            function=3, domain=4, seniority=3, leadership=3, sector=4, practicality=2
        ),
    )
    result = MatchResult(job=JOB, verdict=finalize(assessment, threshold=70))
    monkeypatch.setattr(ws, "result", lambda _job_id: result)
    comparison = asyncio.run(cv_versions.compare(ws, JOB.id, [version.id]))
    original, tailored = comparison.readings
    assert comparison.search_fit_score == result.verdict.fit_score  # type: ignore[union-attr]
    assert original.score == comparison.search_fit_score
    assert original.dimensions[-1].name == "practicality"
    assert tailored.dimensions[-1].after == original.dimensions[-1].after == 2
    assert tailored.score > original.score


def test_a_verdict_for_another_cv_is_never_the_original_score(
    ws: Workspace, master_cv: MasterCV, monkeypatch: pytest.MonkeyPatch
) -> None:
    version = _version(ws, master_cv, TailoringSettings(), role_skills=["PyTorch"])
    ratings = FitRatings(function=3, domain=4, seniority=3, leadership=3, sector=4, practicality=2)
    judged = finalize(JobAssessment(job_id=JOB.id, fit_summary="x", ratings=ratings), 70)
    other = judged.model_copy(update={"cv_id": "someone-else.docx"})
    assert cv_service.search_fit_baseline(MatchResult(job=JOB, verdict=other), "master") is None
    monkeypatch.setattr(ws, "result", lambda _id: MatchResult(job=JOB, verdict=other))
    comparison = asyncio.run(cv_versions.compare(ws, JOB.id, [version.id]))
    assert comparison.search_fit_score is None and "another CV" in comparison.search_fit_note
    same = judged.model_copy(update={"cv_id": version.cv_id})
    monkeypatch.setattr(ws, "result", lambda _id: MatchResult(job=JOB, verdict=same))
    comparison = asyncio.run(cv_versions.compare(ws, JOB.id, [version.id]))
    assert comparison.search_fit_score == same.fit_score and not comparison.search_fit_note


def test_a_confirmed_suggestion_is_scored_against_the_cv_the_search_judged(
    ws: Workspace, master_cv: MasterCV
) -> None:
    standard = _version(ws, master_cv, TailoringSettings())
    optimised = _version(ws, master_cv, TailoringSettings(mode="optimised"))
    baseline = SearchFitBaseline(
        score=70,
        ratings={"function": 3, "domain": 3, "seniority": 3,
                 "leadership": 3, "sector": 3, "practicality": 3},
    )  # fmt: skip
    proposal = ReviewAddition(
        id="skill1",
        kind="skill",
        text="Azure",
        requirement="Azure deployment",
        background_quote="Led migration of model training pipelines to Kubernetes",
        job_quote="Azure deployment",
        reason="Ask about hands-on use",
    )
    tailored = optimised.tailored.model_copy(update={"review_additions": [proposal]})
    tailored_documents._save(
        ws, optimised.model_copy(update={"tailored": tailored, "search_fit": baseline})
    )
    tailored_documents.review_addition(ws, optimised.id, "skill1", "confirm")
    documents = [tailored_documents.load(ws, d.id) for d in (optimised, standard)]
    _items, texts = cv_versions._entries(ws, documents, {}, None)
    assert "Azure" not in texts[0] and "Azure" in texts[1]  # original: the CV as searched
    comparison = asyncio.run(cv_versions.compare(ws, JOB.id, [standard.id, optimised.id]))
    assert comparison.search_fit_score == 70  # the confirmation keeps the search anchor


def test_review_additions_confirm_to_master_or_delete_without_counting_pending(
    ws: Workspace, master_cv: MasterCV
) -> None:
    version = _version(ws, master_cv, TailoringSettings(mode="optimised"))
    proposals = [
        ReviewAddition(
            id="skill1",
            kind="skill",
            text="Azure",
            requirement="Azure deployment",
            background_quote="Led migration of model training pipelines to Kubernetes",
            job_quote="Azure deployment",
            reason="Ask about hands-on use",
        ),
        ReviewAddition(
            id="work1",
            kind="experience",
            role_id="nimbus",
            text="Supported model handover",
            requirement="model handover",
            background_quote="Led migration of model training pipelines to Kubernetes",
            job_quote="model handover",
            reason="Ask about handover",
        ),
    ]
    version = version.model_copy(
        update={
            "tailored": version.tailored.model_copy(update={"review_additions": proposals}),
            "reviewed": False,
        }
    )
    tailored_documents._save(ws, version)
    from src.services import cv_service

    with pytest.raises(ValueError, match="Resolve the CV's red suggestions"):
        asyncio.run(cv_service.write_cover_letter(ws, JOB, tailored_cv_id=version.id))
    confirmed = tailored_documents.review_addition(ws, version.id, "skill1", "confirm", "Azure")
    assert not confirmed.reviewed
    assert confirmed.search_fit == version.search_fit  # still the search score of the original
    assert "Azure" in confirmed.source_cv.all_skills()
    assert "Azure" not in tailored_documents.as_tailored(confirmed).all_skills()
    assert "Azure" in ws.master_cv.all_skills()  # type: ignore[union-attr]
    assert confirmed.tailored.fit_check is None
    from src.cv.ats import check_docx

    path = ws.output_dir / "cvs" / confirmed.filename
    assert check_docx(path, confirmed.tailored.cv, ["Azure"]).keyword_coverage == 1
    from src.web import app as webapp

    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(webapp, "get_workspace", lambda: ws)
        response = TestClient(webapp.app).put(
            f"/api/tailored-cvs/{version.id}/additions/work1", json={"decision": "delete"}
        )
    assert response.status_code == 200
    finished = tailored_documents.load(ws, version.id)
    assert finished.reviewed
    from docx import Document

    lines = Document(ws.output_dir / "cvs" / finished.filename).paragraphs
    assert not any("[CONFIRM " in p.text for p in lines)


def test_confirmed_experience_is_added_to_source_and_tailored_cv(
    ws: Workspace, master_cv: MasterCV
) -> None:
    version = _version(ws, master_cv, TailoringSettings(mode="optimised"))
    proposal = ReviewAddition(
        id="work2",
        kind="experience",
        role_id="nimbus",
        text="Coordinated model handover",
        requirement="model handover",
        background_quote="Led migration of model training pipelines to Kubernetes",
        job_quote="model handover",
        reason="Confirm the scope",
    )
    version = version.model_copy(
        update={
            "tailored": version.tailored.model_copy(update={"review_additions": [proposal]}),
            "reviewed": False,
        }
    )
    tailored_documents._save(ws, version)
    updated = tailored_documents.review_addition(
        ws, version.id, proposal.id, "confirm", "Coordinated model handover with platform teams"
    )
    assert updated.reviewed
    assert any("Coordinated model handover" in b.text for b in ws.master_cv.bullet_index().values())  # type: ignore[union-attr]
    assert any(
        "Coordinated model handover" in b.text for b in updated.tailored.cv.bullet_index().values()
    )


def test_confirmed_skill_is_written_into_selected_word_cv(
    settings: Settings, master_cv: MasterCV, tmp_path: Path
) -> None:
    from docx import Document

    from src.cv.docx_exporter import export_docx
    from src.services import cv_document, cv_service

    ws = Workspace(settings)
    original = export_docx(master_cv, tmp_path / "source.docx")
    asset = cv_service.store_cv(ws, "source.docx", original.read_bytes())
    cv_service.select_cv(ws, asset.id)
    result = cv_document.sync_confirmed_skill(ws, "Azure")
    assert result.written == ["Additional confirmed skill: Azure"]
    path = cv_service.source_file(ws, asset.id)
    assert any("Additional confirmed skill: Azure" in p.text for p in Document(path).paragraphs)


def test_key_requirements_are_read_before_long_desirable_lists(master_cv: MasterCV) -> None:
    rows = [
        RequirementEvidence(
            requirement=f"Desirable {i}", priority="desirable", match="not_evidenced"
        )
        for i in range(15)
    ]
    rows.append(
        RequirementEvidence(
            requirement="Core process development",
            priority="desirable",
            importance="key",
            match="direct",
            evidence=["nimbus-2"],
        )
    )
    tailored = apply_plan(master_cv, TailoringPlan(requirements=rows), JD)
    assert requirements_for(tailored, JD)[0] == "Core process development"


HANDOVER = ReviewAddition(
    id="work1",
    kind="experience",
    role_id="nimbus",
    text="Supported Azure deployment and model handover",
    requirement="Python",
    background_quote="Led migration of model training pipelines to Kubernetes",
    job_quote="model handover",
    reason="Ask about the actual contribution",
)


def _proposal_reader(prompt: str) -> EvidencedReadings:
    """Reads strong only with the proposal in place, and quotes it for requirement 2."""
    out = []
    for _, text in CV_BLOCK.findall(prompt):
        line = next(ln for ln in text.splitlines() if len(ln.split()) >= 2)
        if HANDOVER.text in text:
            out.append(
                EvidencedReading(
                    **{n: RatedLine(rating=r, quote=line) for n, r in STRONG_RATINGS.items()},
                    requirements_met=[
                        MetRequirement(number=1, quote=line),
                        MetRequirement(number=2, quote=HANDOVER.text),
                    ],
                    impression="hands-on deployment",
                )
            )
        else:
            out.append(_evidenced(text, WEAK_RATINGS, [1], "generic"))
    return EvidencedReadings(readings=out)


def _swapped(prompt: str) -> EvidencedReadings:
    """`_proposal_reader` with the proposal shown for requirement 1 (the core one)."""
    out = _proposal_reader(prompt)
    for reading in out.readings:
        reading.requirements_met = [
            MetRequirement(number=1 if m.number == 2 else 2, quote=m.quote)
            for m in reading.requirements_met
        ]
    return out


def test_pending_proposals_are_scored_separately_as_if_confirmed(master_cv: MasterCV) -> None:
    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    tailored.review_additions = [HANDOVER]
    llm = FakeLLM({EvidencedReadings: _proposal_reader})
    check = compare_fit(master_cv, tailored, JD, "MLOps and Python", llm)
    assert len(llm.calls) == 1  # a tailoring's check is one read
    assert HANDOVER.text not in printed_text(tailored)  # the factual draft is unchanged
    assert check.after == check.before and check.requirements_after == 1
    projection = check.if_confirmed
    assert projection is not None and projection.proposals == 1
    assert projection.confirmed == check.after < projection.projected
    assert (projection.requirements_confirmed, projection.requirements_projected) == (1, 2)
    assert [(g.requirement, g.addition_id) for g in projection.gained] == [("Python", "work1")]
    assert projection.no_effect == [] and projection.lost == []
    assert projection.keywords_added == ["Azure"]  # searchable only with the proposal
    assert projection.keywords_confirmed < projection.keywords_projected
    function = next(d for d in projection.dimensions if d.name == "function")
    assert (function.before, function.after) == (2, 4)


def test_without_pending_proposals_there_is_no_projection(master_cv: MasterCV) -> None:
    tailored = apply_plan(master_cv, TailoringPlan(), JD)
    tailored.review_additions = [HANDOVER.model_copy(update={"status": "deleted"})]
    llm = FakeLLM({EvidencedReadings: _reader})
    assert compare_fit(master_cv, tailored, JD, "MLOps and Python", llm).if_confirmed is None


def test_comparison_reads_an_optimised_version_with_its_proposals_in_place(
    ws: Workspace, master_cv: MasterCV, monkeypatch: pytest.MonkeyPatch
) -> None:
    llm = FakeLLM({EvidencedReadings: _proposal_reader})
    monkeypatch.setattr(ws, "structured", lambda *_args, **_kw: llm)
    standard = _version(ws, master_cv, TailoringSettings())
    optimised = _version(ws, master_cv, TailoringSettings(mode="optimised"))
    tailored = optimised.tailored.model_copy(update={"review_additions": [HANDOVER]})
    tailored_documents._save(ws, optimised.model_copy(update={"tailored": tailored}))
    comparison = asyncio.run(cv_versions.compare(ws, JOB.id, []))
    by_id = {r.id: r for r in comparison.readings}
    assert len(comparison.readings) == 3  # the projection is not a version of its own
    assert by_id[standard.id].if_confirmed is None
    mine = by_id[optimised.id]
    assert mine.score == by_id[standard.id].score  # its own score stays factual
    assert mine.if_confirmed is not None
    assert mine.if_confirmed.confirmed == mine.score < mine.if_confirmed.projected
    assert [g.addition_id for g in mine.if_confirmed.gained] == ["work1"]


def test_comparison_projects_changes_when_all_cv_slots_are_used(
    ws: Workspace, master_cv: MasterCV, monkeypatch: pytest.MonkeyPatch
) -> None:
    llm = FakeLLM({EvidencedReadings: _proposal_reader})
    monkeypatch.setattr(ws, "structured", lambda *_args, **_kw: llm)
    for _ in range(4):
        _version(ws, master_cv, TailoringSettings())
    optimised = _version(ws, master_cv, TailoringSettings(mode="optimised"))
    tailored = optimised.tailored.model_copy(update={"review_additions": [HANDOVER]})
    tailored_documents._save(ws, optimised.model_copy(update={"tailored": tailored}))

    comparison = asyncio.run(cv_versions.compare(ws, JOB.id, []))

    assert len(comparison.readings) == cv_versions.MAX_VERSIONS
    mine = next(r for r in comparison.readings if r.id == optimised.id)
    assert mine.if_confirmed is not None
    assert mine.if_confirmed.confirmed == mine.score < mine.if_confirmed.projected


def test_confirming_works_whichever_cv_is_selected_and_keeps_one_file(
    ws: Workspace, master_cv: MasterCV
) -> None:
    version = _version(ws, master_cv, TailoringSettings(mode="optimised"))
    proposal = ReviewAddition(
        id="skill1", kind="skill", text="Azure", requirement="Azure deployment",
        background_quote="Led migration of model training pipelines to Kubernetes",
        job_quote="Azure deployment", reason="Ask about hands-on use",
    )  # fmt: skip
    tailored = version.tailored.model_copy(update={"review_additions": [proposal]})
    tailored_documents._save(ws, version.model_copy(update={"tailored": tailored}))
    other = master_cv.model_copy(
        update={"basics": master_cv.basics.model_copy(update={"name": "Other"})}
    )
    ws.active_cv_id, ws.master_cv = "another-cv", other  # the user selected another CV since
    confirmed = tailored_documents.review_addition(ws, version.id, "skill1", "confirm", "Azure")
    assert (ws.active_cv_id, ws.master_cv) == ("another-cv", other)  # selection untouched
    assert confirmed.filename == version.filename  # the draft's own file, rewritten
    saved = MasterCV.model_validate_json(ws.settings.master_cv_path.read_text())
    assert "Azure" in saved.all_skills()  # into the CV the draft came from
