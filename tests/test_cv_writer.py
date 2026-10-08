"""cv_writer: Master CV management, guarded tailoring, .docx export."""

from __future__ import annotations

import asyncio
import json
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.shared import Pt
from pydantic import ValidationError

from src.cv import master_cv_manager as mgr
from src.cv.docx_exporter import TEMPLATES, export_docx, fmt_date
from src.cv.docx_original import restore_normalized_indents
from src.cv.models import (
    CVCritique,
    JDAnalysis,
    MasterCV,
    RewrittenBullet,
    TailoredCV,
    TailoringPlan,
)
from src.cv.tailor import apply_plan, keyword_coverage, tailor
from src.services import cv_service
from src.services.workspace import Workspace
from tests.conftest import FakeLLM

JD = JDAnalysis(
    job_title="Staff ML Engineer",
    company="Orbit AI",
    hard_skills=["Python", "PyTorch", "Kubernetes", "Spark"],
    must_have=["MLOps"],
)


# ---------------------------------------------------------------- master CV


def test_committed_schema_matches_model() -> None:
    committed = json.loads(mgr.SCHEMA_PATH.read_text())
    assert committed == mgr.json_schema(), "Run: python -m src.cv.master_cv_manager export-schema"


def test_example_cv_validates_against_json_schema(master_cv: MasterCV) -> None:
    jsonschema = pytest.importorskip("jsonschema")
    jsonschema.validate(master_cv.model_dump(mode="json"), mgr.json_schema())


def test_invalid_dates_and_duplicate_ids_rejected(master_cv: MasterCV) -> None:
    data = master_cv.model_dump(mode="json")
    data["experience"][0]["start"] = "2021/03"
    with pytest.raises(ValidationError):
        MasterCV.model_validate(data)
    data = master_cv.model_dump(mode="json")
    data["experience"][1]["bullets"][0]["id"] = "nimbus-1"
    with pytest.raises(ValidationError, match="Duplicate ids"):
        MasterCV.model_validate(data)


def test_save_load_roundtrip_keeps_backup(master_cv: MasterCV, tmp_path: Path) -> None:
    path = tmp_path / "cv.json"
    mgr.save(master_cv, path)
    mgr.save(master_cv, path)
    assert mgr.load(path) == master_cv
    assert path.with_suffix(".json.bak").exists()


def test_update_merge_patch(master_cv: MasterCV) -> None:
    updated = mgr.update(master_cv, {"basics": {"phone": "+351 900 000 000", "summary": None}})
    assert updated.basics.phone == "+351 900 000 000"
    assert updated.basics.summary is None
    assert updated.basics.name == master_cv.basics.name


def test_from_text_uses_llm(master_cv: MasterCV) -> None:
    llm = FakeLLM({MasterCV: master_cv})
    assert mgr.from_text("raw cv text", llm) == master_cv
    assert llm.calls[0][0] == "raw cv text"


def test_upload_stores_without_llm_then_imports_on_demand(
    master_cv: MasterCV, settings: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    ws = Workspace(settings)
    llm = FakeLLM({MasterCV: master_cv})
    monkeypatch.setattr(ws, "structured", lambda role="worker", *_: llm)
    monkeypatch.setattr(ws, "llm_ready", lambda: True)
    monkeypatch.setattr(ws, "role_ready", lambda role: True)
    raw = ("Alex Example\nSenior engineer\n" + "Built reliable systems.\n" * 8).encode()

    asset = cv_service.store_cv(ws, "resume.txt", raw)

    assert asset.filename == "resume.txt"
    stored = settings.data_dir / "cvs" / "resume.txt"
    assert stored.read_bytes() == raw
    assert ws.master_cv is None
    assert llm.calls == []

    imported = asyncio.run(cv_service.ensure_selected_cv(ws))
    assert imported == master_cv
    assert llm.calls[0][0].startswith("Alex Example")
    assert len(llm.calls) == 1
    edited = imported.model_copy(
        update={"basics": imported.basics.model_copy(update={"headline": "Edited headline"})}
    )
    cv_service.save_selected_cv(ws, edited)
    ws.master_cv = None
    assert asyncio.run(cv_service.ensure_selected_cv(ws)).basics.headline == "Edited headline"
    assert len(llm.calls) == 1
    with pytest.raises(ValueError, match="empty"):
        cv_service.store_cv(ws, "empty.txt", b"")
    with pytest.raises(ValueError, match="10 MB"):
        cv_service.store_cv(ws, "large.txt", b"x" * (cv_service.MAX_UPLOAD_BYTES + 1))


def test_upload_deduplicates_content_and_preserves_same_filename(settings: Any) -> None:
    ws = Workspace(settings)
    original = b"A sufficiently detailed CV"

    first = cv_service.store_cv(ws, "My CV.txt", original)
    duplicate = cv_service.store_cv(ws, "Copy of My CV.txt", original)

    assert duplicate.id == first.id
    assert duplicate.filename == "My CV.txt"
    assert [path.name for path in (settings.data_dir / "cvs").iterdir()] == ["My CV.txt"]

    replacement = cv_service.store_cv(ws, "My CV.txt", b"An updated and sufficiently detailed CV")

    assert replacement.id != first.id
    assert replacement.filename == "My CV (2).txt"
    assert [item.filename for item in cv_service.list_cvs(ws)] == ["My CV (2).txt", "My CV.txt"]


def test_upload_hidden_filename_remains_selectable(settings: Any) -> None:
    ws = Workspace(settings)

    asset = cv_service.store_cv(ws, ".resume.txt", b"A sufficiently detailed hidden CV")

    assert asset.filename == "resume.txt"
    assert cv_service.list_cvs(ws)[0].id == asset.id


def test_uploaded_file_replaces_legacy_master_library_entry(
    master_cv: MasterCV, settings: Any
) -> None:
    ws = Workspace(settings)
    ws.save_master_cv(master_cv)

    uploaded = cv_service.store_cv(ws, "My CV.txt", b"A sufficiently detailed CV")

    assert [(item.filename, item.selected) for item in cv_service.list_cvs(ws)] == [
        ("My CV.txt", True)
    ]
    assert ws.active_cv_id == uploaded.id


def test_legacy_hashed_uploads_migrate_to_real_filename(settings: Any) -> None:
    ws = Workspace(settings)
    directory = settings.data_dir / "cvs"
    directory.mkdir()
    data = b"A sufficiently detailed CV"
    real_id = "1" * 24
    duplicate_id = "2" * 24
    (directory / f"{real_id}.source.pdf").write_bytes(data)
    (directory / f"{real_id}.json").write_text(
        json.dumps(
            {
                "id": real_id,
                "filename": "My Real CV.pdf",
                "size": len(data),
                "kind": "uploaded",
            }
        )
    )
    (directory / f"{duplicate_id}.source.pdf").write_bytes(data)
    (directory / f"{duplicate_id}.json").write_text(
        json.dumps(
            {
                "id": duplicate_id,
                "filename": f"{real_id}.source.pdf",
                "size": len(data),
                "kind": "uploaded",
            }
        )
    )

    assets = cv_service.list_cvs(ws)

    assert [(item.filename, item.kind, item.selected) for item in assets] == [
        ("My Real CV.pdf", "uploaded", True)
    ]
    assert [path.name for path in directory.iterdir()] == ["My Real CV.pdf"]


def test_cv_library_lists_and_selects_generated_files(master_cv: MasterCV, settings: Any) -> None:
    ws = Workspace(settings)
    ws.save_master_cv(master_cv)
    ws.active_cv_id = "master"
    settings.output_dir.mkdir(parents=True)
    generated = settings.output_dir / "Tailored.docx"
    generated.write_bytes(b"word document")

    assets = cv_service.list_cvs(ws)

    assert [(item.id, item.selected) for item in assets] == [
        ("master", True),
        ("generated:Tailored.docx", False),
    ]
    selected = cv_service.select_cv(ws, "generated:Tailored.docx")
    assert selected.selected and selected.kind == "generated" and ws.master_cv is None


# ---------------------------------------------------------------- tailoring


def test_faithful_rewrite_accepted(master_cv: MasterCV) -> None:
    plan = TailoringPlan(
        rewritten_bullets=[
            RewrittenBullet(
                source_id="nimbus-2",
                text="Drove MLOps migration of training pipelines to Kubernetes, cutting cost 35%.",
            )
        ]
    )
    out = apply_plan(master_cv, plan, JD)
    assert out.changes[0].accepted
    assert out.cv.bullet_index()["nimbus-2"].text.startswith("Drove MLOps")
    assert master_cv.bullet_index()["nimbus-2"].text.startswith("Led migration")  # untouched


def test_invented_metric_rejected(master_cv: MasterCV) -> None:
    plan = TailoringPlan(
        rewritten_bullets=[
            RewrittenBullet(
                source_id="nimbus-2",
                text="Migrated pipelines to Kubernetes, cutting costs by 60%.",
            )
        ]
    )
    out = apply_plan(master_cv, plan, JD)
    assert not out.changes[0].accepted
    assert "60%" in (out.changes[0].reason or "")
    assert out.cv.bullet_index()["nimbus-2"].text == master_cv.bullet_index()["nimbus-2"].text


def test_unevidenced_jd_skill_rejected(master_cv: MasterCV) -> None:
    plan = TailoringPlan(
        rewritten_bullets=[
            RewrittenBullet(
                source_id="nimbus-2",
                text="Migrated Spark pipelines to Kubernetes, cutting costs by 35%.",
            )
        ]
    )
    out = apply_plan(master_cv, plan, JD)
    assert not out.changes[0].accepted
    assert "Spark" in (out.changes[0].reason or "")


def test_unknown_source_id_rejected(master_cv: MasterCV) -> None:
    plan = TailoringPlan(rewritten_bullets=[RewrittenBullet(source_id="ghost", text="x")])
    assert not apply_plan(master_cv, plan, JD).changes[0].accepted


def test_bullet_order_and_skill_priority(master_cv: MasterCV) -> None:
    plan = TailoringPlan(
        bullet_order={"nimbus": ["nimbus-2", "nimbus-1"]},
        skills_priority=["Kubernetes", "PyTorch"],
    )
    cv = apply_plan(master_cv, plan, JD).cv
    assert [b.id for b in cv.experience[0].bullets] == ["nimbus-2", "nimbus-1", "nimbus-3"]
    assert cv.skills[0].category == "Infrastructure"
    assert cv.skills[0].items[0] == "Kubernetes"
    assert master_cv.all_skills() == cv.all_skills()  # reordered, nothing removed


def test_keyword_coverage(master_cv: MasterCV) -> None:
    matched, missing = keyword_coverage(master_cv, JD)
    assert set(matched) == {"Python", "PyTorch", "Kubernetes", "MLOps"}
    assert missing == ["Spark"]


def test_requirement_sentences_are_not_reported_as_missing_keywords(master_cv: MasterCV) -> None:
    jd = JD.model_copy(
        update={"must_have": ["MLOps", "Experience leading and developing engineering teams."]}
    )
    matched, missing = keyword_coverage(master_cv, jd)
    assert "MLOps" in matched and missing == ["Spark"]


def test_tailor_end_to_end(master_cv: MasterCV) -> None:
    plan = TailoringPlan(headline_options=["Senior Machine Learning Engineer | MLOps"])
    llm = FakeLLM({JDAnalysis: JD, TailoringPlan: plan, CVCritique: CVCritique()})
    out = tailor(master_cv, "We need a Staff ML Engineer...", llm)
    assert out.cv.basics.headline == master_cv.basics.headline  # the CV's own headline stays
    assert out.headline_options == ["Senior Machine Learning Engineer | MLOps"]
    assert out.trim is not None
    assert out.target_company == "Orbit AI"
    assert 0 < out.keyword_coverage < 1
    assert [m for _, m in llm.calls].count(TailoringPlan) == 1  # no issues: no revision


# ---------------------------------------------------------------- docx export


@pytest.mark.parametrize("template", sorted(TEMPLATES))
def test_export_docx(master_cv: MasterCV, tmp_path: Path, template: str) -> None:
    path = export_docx(master_cv, tmp_path / f"cv_{template}.docx", template)
    text = "\n".join(p.text for p in Document(str(path)).paragraphs)
    assert "Alex Example" in text
    assert "Nimbus Analytics" in text
    assert "Mar 2021 – Present" in text
    assert "4M users" in text


def test_export_unknown_template(master_cv: MasterCV, tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="Unknown template"):
        export_docx(master_cv, tmp_path / "x.docx", "neon")


def test_fmt_date() -> None:
    assert fmt_date("2019-12") == "Dec 2019"
    assert fmt_date(None) == "Present"
    assert fmt_date("2018") == "2018"


# ---------------------------------------------------------------- the original's design


def test_restore_viewer_indent_drift_keeps_edited_wording(tmp_path: Path) -> None:
    original = tmp_path / "original.docx"
    edited = tmp_path / "tailored.docx"
    doc = Document()
    doc.styles["Normal"].paragraph_format.left_indent = Pt(10)
    doc.styles["Normal"].paragraph_format.first_line_indent = Pt(-10)
    for index in range(30):
        paragraph = doc.add_paragraph(f"Line {index}")
        if index % 2:
            paragraph.paragraph_format.first_line_indent = Pt(-9)
        else:
            paragraph.paragraph_format.left_indent = Pt(0)
            paragraph.paragraph_format.first_line_indent = Pt(0)
    doc.save(str(original))
    damaged = Document(str(original))
    damaged.paragraphs[4].runs[0].text = "Edited line four"
    for index, paragraph in enumerate(damaged.paragraphs):
        if index % 2:
            paragraph.paragraph_format.left_indent = Pt(15)
        else:
            paragraph.paragraph_format.first_line_indent = None
    damaged.save(str(edited))

    assert restore_normalized_indents(original, edited)
    repaired = Document(str(edited))
    source = Document(str(original))
    assert repaired.paragraphs[4].text == "Edited line four"
    assert all(
        (before.paragraph_format.left_indent, before.paragraph_format.first_line_indent)
        == (after.paragraph_format.left_indent, after.paragraph_format.first_line_indent)
        for before, after in zip(source.paragraphs, repaired.paragraphs, strict=True)
    )
    assert not restore_normalized_indents(original, edited)
    repaired.paragraphs[0].paragraph_format.left_indent = Pt(3)
    repaired.save(str(edited))
    assert not restore_normalized_indents(original, edited)  # a deliberate local edit stays


def _original_cv(master: MasterCV, path: Path) -> Path:
    """A Word CV with its own design: styled name and title, headings, bullets, a
    publication and labelled skills lines."""
    from docx.shared import Pt

    doc = Document()
    name = doc.add_paragraph()
    name.add_run(master.basics.name).font.size = Pt(26)
    doc.add_paragraph(master.basics.headline or "")
    doc.add_heading("Profile", level=1)
    doc.add_paragraph(master.basics.summary or "")
    doc.add_heading("Experience", level=1)
    for role in master.experience:
        doc.add_paragraph(f"{role.title}, {role.company}")
        for bullet in role.bullets:
            doc.add_paragraph(bullet.text, style="List Bullet")
    doc.add_heading("Publications", level=1)
    doc.add_paragraph("Example A. et al. (2020) Recommender systems at scale. J. ML 12:1-9.")
    doc.add_heading("Skills", level=1)
    for group in master.skills:
        line = doc.add_paragraph()
        line.add_run(f"{group.category}: ").bold = True
        line.add_run(", ".join(group.items))
    doc.save(str(path))
    return path


def test_tailored_cv_is_a_copy_of_the_original_with_only_content_changed(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    from src.cv.docx_original import write_like_original

    original = _original_cv(master_cv, tmp_path / "original.docx")
    plan = TailoringPlan(
        headline_options=["Senior Machine Learning Engineer, MLOps"],  # suggested only
        summary="ML engineer with 8 years building production recommendation systems.",
        rewritten_bullets=[
            RewrittenBullet(
                source_id="nimbus-2",
                text="Led MLOps migration of model training pipelines to Kubernetes, cutting "
                "training costs by 35%.",
            )
        ],
        bullet_order={"nimbus": ["nimbus-2"]},  # the others follow, nothing is removed
        skills_priority=["Kubernetes", "Docker", "PyTorch"],
    )
    tailored = apply_plan(master_cv, plan, JD)
    notes = write_like_original(original, master_cv, tailored.cv, tmp_path / "tailored.docx")
    before = Document(str(original)).paragraphs
    after = Document(str(tmp_path / "tailored.docx")).paragraphs
    texts = [p.text for p in after]

    assert notes == [] and len(after) == len(before)  # nothing lost, nothing added
    assert after[0].runs[0].font.size == before[0].runs[0].font.size  # same design
    assert texts[:3] == [master_cv.basics.name, master_cv.basics.headline, "Profile"]
    assert texts[3] == plan.summary
    nimbus = texts.index("Senior Machine Learning Engineer, Nimbus Analytics")
    assert texts[nimbus + 1].startswith("Led MLOps migration")
    assert [p.style.name for p in after[nimbus + 1 : nimbus + 4]] == ["List Bullet"] * 3
    assert texts[nimbus + 2].startswith("Built a real-time") and texts[nimbus + 3].startswith(
        "Mentored"
    )
    assert any(t.startswith("Example A. et al. (2020)") for t in texts)  # publications kept
    infrastructure = next(p for p in after if p.text.startswith("Infrastructure:"))
    assert infrastructure.text == "Infrastructure: Kubernetes, Docker, AWS, PostgreSQL"
    assert infrastructure.runs[0].bold and not infrastructure.runs[1].bold  # label kept bold
    headings = [p.text for p in after if p.style.name.startswith("Heading")]
    assert headings == ["Profile", "Experience", "Publications", "Skills"]


def test_trimming_removes_lines_from_the_original_only_when_allowed(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    from src.cv.docx_original import write_like_original
    from src.cv.models import TrimAssessment

    original = _original_cv(master_cv, tmp_path / "original.docx")
    plan = TailoringPlan(bullet_order={"nimbus": ["nimbus-1", "nimbus-2"]})
    trim = TrimAssessment(allowed=True, reason="a junior-level role")
    tailored = apply_plan(master_cv, plan, JD, trim)
    write_like_original(original, master_cv, tailored.cv, tmp_path / "trimmed.docx")
    texts = [p.text for p in Document(str(tmp_path / "trimmed.docx")).paragraphs]
    assert not any(t.startswith("Mentored") for t in texts)
    assert any(t.startswith("Example A. et al.") for t in texts)


def test_tailoring_starts_from_the_documents_own_wording(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    from src.cv.docx_original import align_to_document

    original = _original_cv(master_cv, tmp_path / "original.docx")
    doc = Document(str(original))
    # The user's file has one more line than the parsed CV knows about.
    mentoring = next(p for p in doc.paragraphs if p.text.startswith("Mentored"))
    extra = mentoring.insert_paragraph_before(
        "Organised the weekly reading group on recommender systems research.",
        style="List Bullet",
    )
    assert extra.text
    doc.save(str(original))
    condensed = master_cv.model_copy(deep=True)
    condensed.experience[0].bullets[0].text = "Built a real-time recommendation service."

    aligned = align_to_document(original, condensed)
    nimbus = aligned.experience[0]
    assert nimbus.bullets[0].id == "nimbus-1"  # same id, the document's full wording
    assert nimbus.bullets[0].text == master_cv.experience[0].bullets[0].text
    assert [b.text for b in nimbus.bullets].count(extra.text) == 1  # the missed line is added
    assert len(nimbus.bullets) == 4 and aligned.experience[1] == master_cv.experience[1]


def test_bullet_order_survives_a_strict_output_schema(master_cv: MasterCV) -> None:
    from src.core.llm.codex_backend import strict_schema

    schema = strict_schema(TailoringPlan.model_json_schema())
    order = schema["$defs"]["RoleOrder"]["properties"]
    assert set(order) == {"experience_id", "bullet_ids"}  # not an empty, keyless object
    plan = TailoringPlan.model_validate(
        {"bullet_order": [{"experience_id": "nimbus", "bullet_ids": ["nimbus-3"]}]}
    )
    tailored = apply_plan(master_cv, plan, JD)
    assert tailored.cv.experience[0].bullets[0].id == "nimbus-3"


def test_a_profile_of_several_paragraphs_is_read_and_rewritten_in_place(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    from src.cv.docx_original import align_to_document, write_like_original

    original = _original_cv(master_cv, tmp_path / "original.docx")
    doc = Document(str(original))
    first = next(p for p in doc.paragraphs if p.text == master_cv.basics.summary)
    second = "People leader who has mentored engineers and run weekly design reviews."
    first.insert_paragraph_before("Machine learning engineer focused on recommendation systems.")
    doc.save(str(original))
    source = align_to_document(original, master_cv.model_copy(deep=True))
    assert source.basics.summary == (
        f"Machine learning engineer focused on recommendation systems.\n\n"
        f"{master_cv.basics.summary}"
    )

    tailored = source.model_copy(deep=True)
    tailored.basics.summary = f"ML engineer building production recommenders.\n\n{second}"
    write_like_original(original, source, tailored, tmp_path / "two.docx")
    texts = [p.text for p in Document(str(tmp_path / "two.docx")).paragraphs]
    profile = texts.index("Profile")
    assert texts[profile + 1 : profile + 3] == [
        "ML engineer building production recommenders.",
        second,
    ]

    tailored.basics.summary = "ML engineer building production recommenders for 4M users."
    write_like_original(original, source, tailored, tmp_path / "one.docx")
    texts = [p.text for p in Document(str(tmp_path / "one.docx")).paragraphs]
    assert texts[profile + 1 : profile + 3] == [tailored.basics.summary, "Experience"]


def test_suggested_headlines_are_written_under_the_cvs_own(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    from src.cv.docx_original import write_like_original

    options = ["Machine Learning Engineer | Recommender Systems", "ML Engineer | MLOps"]
    original = _original_cv(master_cv, tmp_path / "original.docx")
    write_like_original(original, master_cv, master_cv, tmp_path / "cv.docx", options)
    after = Document(str(tmp_path / "cv.docx")).paragraphs
    assert [p.text for p in after[1:4]] == [master_cv.basics.headline, *options]
    assert after[2].style.name == after[1].style.name  # in the headline's own formatting

    tailored = TailoredCV(
        cv=master_cv, target_title="ML Engineer", keyword_coverage=1, headline_options=options
    )
    export_docx(tailored, tmp_path / "classic.docx")
    texts = [p.text for p in Document(str(tmp_path / "classic.docx")).paragraphs]
    assert texts[1:4] == [master_cv.basics.headline, *options]


def test_the_review_reads_the_cv_as_printed_and_knows_the_rules(master_cv: MasterCV) -> None:
    from src.cv.tailor import CRITIC_SYSTEM, cv_to_text

    tag = master_cv.experience[0].bullets[0].skills[0]
    printed = cv_to_text(master_cv, tags=False).splitlines().count(tag)
    assert printed < cv_to_text(master_cv).splitlines().count(tag)  # no per-role tag lines
    assert "never removing" in CRITIC_SYSTEM and "own headline" in CRITIC_SYSTEM


def test_footer_carries_the_cvs_own_name_and_todays_date(
    master_cv: MasterCV, tmp_path: Path
) -> None:
    from datetime import date

    from src.cv.docx_original import write_like_original

    cv = master_cv.model_copy(deep=True)
    cv.basics.name = "Alex Example Morgan, PhD"
    original = _original_cv(cv, tmp_path / "original.docx")
    doc = Document(str(original))
    footer = doc.sections[0].footer.paragraphs[0]
    footer.add_run("Alex B. Example, PhD | Master CV | 4 October 2026 | Page ")
    footer.add_run("1")
    doc.save(str(original))
    write_like_original(original, cv, cv, tmp_path / "tailored.docx")
    out = Document(str(tmp_path / "tailored.docx")).sections[0].footer.paragraphs[0]
    today = f"{date.today().day} {date.today():%B %Y}"
    assert out.text == f"Alex Example Morgan, PhD | CV | {today} | Page 1"
    assert [r.text for r in out.runs][1] == "1"  # the page number run is untouched

    # A footer that already has the full name keeps it, even when the CV's name is shorter.
    cv.basics.name = "Alex Example"
    write_like_original(tmp_path / "tailored.docx", cv, cv, tmp_path / "again.docx")
    again = Document(str(tmp_path / "again.docx")).sections[0].footer.paragraphs[0]
    assert again.text.startswith("Alex Example Morgan, PhD | CV")


def test_a_rewrite_may_not_drop_the_lines_numbers_or_most_of_its_detail(
    master_cv: MasterCV,
) -> None:
    def rewrite(text: str) -> Any:
        plan = TailoringPlan(rewritten_bullets=[RewrittenBullet(source_id="nimbus-1", text=text)])
        return apply_plan(master_cv, plan, JD).changes[0]

    assert "drops numbers" in (rewrite("Built a real-time recommender in Python.").reason or "")
    short = rewrite("Built a recommendation service serving 4M users, CTR up 18%.")
    assert "too much" in (short.reason or "")
    assert rewrite(
        "Built a real-time recommendation service in Python and PyTorch for 4M users, "
        "lifting click-through rate by 18%."
    ).accepted
