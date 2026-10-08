"""Low-level python-docx helpers. No CV knowledge here, only document primitives."""

from __future__ import annotations

import copy
from collections.abc import Collection
from difflib import SequenceMatcher
from typing import Any

from docx.document import Document as DocxDocument
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


def hex_to_rgb(color: str) -> RGBColor:
    """'#1F3864' -> RGBColor."""
    return RGBColor.from_string(color.lstrip("#").upper())


def set_margins(doc: DocxDocument, inches: float) -> None:
    """Apply equal page margins to every section."""
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(inches)
        section.left_margin = section.right_margin = Inches(inches)


def set_base_font(doc: DocxDocument, name: str, size_pt: float) -> None:
    """Set the Normal style font (including East-Asian fallback so Word honours it)."""
    style = doc.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(size_pt)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), name)
    pf = style.paragraph_format
    pf.space_after = Pt(2)
    pf.space_before = Pt(0)


def add_bottom_border(paragraph: Paragraph, color: str = "808080", size: int = 6) -> None:
    """Draw a horizontal rule under a paragraph."""
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color.lstrip("#"))
    borders.append(bottom)
    ppr.append(borders)


def add_text(
    doc: DocxDocument,
    text: str,
    *,
    size_pt: float | None = None,
    bold: bool = False,
    italic: bool = False,
    color: str | None = None,
    align: int | None = None,
    space_after_pt: float | None = None,
) -> Paragraph:
    """Add a single-run paragraph."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold, run.italic = bold, italic
    if size_pt:
        run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = hex_to_rgb(color)
    if align is not None:
        p.alignment = align  # type: ignore[assignment]
    if space_after_pt is not None:
        p.paragraph_format.space_after = Pt(space_after_pt)
    return p


def add_split_line(
    doc: DocxDocument, left: str, right: str, *, bold_left: bool = True, right_tab_in: float = 7.0
) -> Paragraph:
    """One line with `left` text and `right` text flush to a right tab stop."""
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(right_tab_in), WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.space_before = Pt(4)
    p.add_run(left).bold = bold_left
    if right:
        p.add_run(f"\t{right}")
    return p


def add_bullet(doc: DocxDocument, text: str) -> Paragraph:
    """Add a 'List Bullet' paragraph (ATS-friendly native Word list)."""
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    return p


def _plain(text: str) -> str:
    return " ".join(text.casefold().replace("\u2019", "'").split())


def all_paragraphs(doc: DocxDocument) -> list[Paragraph]:
    """Body paragraphs, including those inside table cells (CV layouts often use tables)."""
    out = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                out.extend(p for p in cell.paragraphs if p not in out)
    return out


def find_paragraph(
    doc: DocxDocument, text: str, min_ratio: float = 0.8, exclude: Collection[Any] = ()
) -> Paragraph | None:
    """The paragraph whose text best matches `text` (exact, then contained, then similar).
    `exclude`: paragraph elements (`p._p`) already claimed by other text."""
    target = _plain(text)
    if not target:
        return None
    best: tuple[float, Paragraph | None] = (0.0, None)
    for p in all_paragraphs(doc):
        here = _plain(p.text)
        if not here or any(p._p is e for e in exclude):
            continue
        if here == target:
            return p
        ratio = 0.95 if target in here or (here in target and len(here) > 20) else 0.0
        ratio = max(ratio, SequenceMatcher(None, here, target).ratio())
        if ratio > best[0]:
            best = (ratio, p)
    return best[1] if best[0] >= min_ratio else None


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    """Replace a paragraph's text, keeping its style and its first run's formatting."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


def insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    """A copy of `paragraph` (same style, bullet and formatting) holding `text`, placed after it."""
    new = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(new)
    out = Paragraph(new, paragraph._parent)
    set_paragraph_text(out, text)
    return out


def remove_paragraph(paragraph: Paragraph) -> None:
    paragraph._p.getparent().remove(paragraph._p)


def reorder_paragraphs(paragraphs: list[Paragraph], order: list[Paragraph]) -> None:
    """Put `order` into the places `paragraphs` occupy (document order), each keeping its own
    formatting. Paragraphs not in `order` are removed; `order` may not add new ones."""
    slots = sorted(paragraphs, key=_position)
    marks = []
    for p in slots:
        mark = OxmlElement("w:p")
        p._p.addprevious(mark)
        marks.append(mark)
        p._p.getparent().remove(p._p)
    for mark, p in zip(marks, order, strict=False):
        mark.addprevious(p._p)
    for mark in marks:
        mark.getparent().remove(mark)


def _position(paragraph: Paragraph) -> int:
    body = paragraph._p.getroottree().getroot()
    return next(i for i, e in enumerate(body.iter()) if e is paragraph._p)


def replace_text_after(paragraph: Paragraph, prefix_len: int, text: str) -> None:
    """Replace the paragraph's text after its first `prefix_len` characters (e.g. a bold
    "Programming:" label), keeping the label's runs and formatting."""
    seen = 0
    runs = paragraph.runs
    for index, run in enumerate(runs):
        end = seen + len(run.text)
        if end > prefix_len or (end == prefix_len and index == len(runs) - 1):
            run.text = run.text[: prefix_len - seen] + text
            for later in runs[index + 1 :]:
                later._r.getparent().remove(later._r)
            return
        seen = end
    paragraph.add_run(text)
