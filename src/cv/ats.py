"""Check an exported .docx the way an applicant-tracking system reads it.

ATS parsers read the file's text, not its look: contact details must be plain text in the
body (not in headers, text boxes or icons), keywords must be present as words, and tables
can scramble the reading order. The page count is estimated from the word count and
reported, not warned about: the user prefers a full-length CV (user's decision, 2026-10-07).
Layout warnings follow what parsers are known to miss (user's decision 2026-10-14): two
columns read straight across, contact details in a header or footer, text in images or text
boxes, role dates the parser cannot find (it counts years of experience from them), no
standard section headings, and unusual fonts. They are warnings only: a tailored CV keeps the
design of the user's own Word file.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn

from src.cv.ats_score import term_found
from src.cv.models import ATSReport, MasterCV

WORDS_PER_PAGE = 500  # a dense single-column CV page
UNVERIFIED_PREFIX = "Skills to confirm before use (not yet claimed):"
REVIEW_PREFIX = "[CONFIRM "
INLINE_REVIEW = re.compile(r"\s*(?:·\s*)?\[CONFIRM[^\]]*\]")
COMMON_FONTS = {
    "arial", "calibri", "cambria", "garamond", "georgia", "helvetica", "verdana", "tahoma",
    "times new roman", "aptos", "book antiqua", "palatino linotype", "segoe ui", "trebuchet ms",
    "century gothic", "lato", "open sans", "roboto", "source sans pro", "carlito",
    "liberation sans", "liberation serif", "arial narrow", "franklin gothic book", "gill sans mt",
    "corbel", "candara", "constantia", "symbol", "wingdings",
}  # fmt: skip
STANDARD_HEADINGS = {
    "experience": ("experience", "employment", "work history", "positions", "career"),
    "education": ("education", "qualifications", "degrees"),
}
YEAR = re.compile(r"\b(?:19|20)\d{2}\b")


def docx_text(path: Path) -> tuple[str, int]:
    """Body text as a parser sees it (paragraphs, then table cells) and the table count."""
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts += [cell.text for row in table.rows for cell in row.cells]
    return "\n".join(parts), len(doc.tables)


def _plain(text: str) -> str:
    """Case and spacing do not matter to a parser (a name set in capitals is still read)."""
    return " ".join(text.casefold().split())


def _header_text(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for s in doc.sections for p in s.header.paragraphs)


def _footer_text(doc: DocxDocument) -> str:
    return "\n".join(p.text for s in doc.sections for p in s.footer.paragraphs)


def _columns(doc: DocxDocument) -> bool:
    """A section set in more than one column (parsers read straight across them)."""
    for sect in doc.element.body.iter(qn("w:sectPr")):
        cols = sect.find(qn("w:cols"))
        if cols is not None and int(cols.get(qn("w:num"), "1") or "1") > 1:
            return True
    return False


def _boxes(doc: DocxDocument) -> tuple[int, int]:
    """(images, text boxes) in the page body: a parser reads no text inside either."""
    body = doc.element.body
    boxes = len(list(body.iter(qn("w:txbxContent"))))
    images = len(list(body.iter(qn("w:drawing")))) + len(list(body.iter(qn("w:pict"))))
    return max(0, images - boxes), boxes


def _body_lines(doc: DocxDocument) -> list[str]:
    """The page body in reading order: each paragraph, and each table row as one line (a
    role's dates often sit in the next cell)."""
    lines: list[str] = []
    for child in doc.element.body:
        if child.tag == qn("w:p"):
            lines.append("".join(child.itertext()))
        elif child.tag == qn("w:tbl"):
            for row in child.iter(qn("w:tr")):
                cells = ["".join(c.itertext()) for c in row.iter(qn("w:tc"))]
                lines.append(" | ".join(cells))
    return lines


def _undated(lines: list[str], cv: MasterCV) -> list[str]:
    """Roles whose heading line (or the two lines after it) has no year a parser can read."""
    out = []
    for role in cv.experience:
        index = next(
            (i for i, line in enumerate(lines) if role.company and role.company in line), None
        )
        if index is not None and not YEAR.search(" ".join(lines[index : index + 3])):
            out.append(role.company)
    return out


def _missing_headings(lines: list[str]) -> list[str]:
    short = [_plain(line) for line in lines if 0 < len(line.strip()) < 40]
    return [
        kind
        for kind, words in STANDARD_HEADINGS.items()
        if not any(word in line for line in short for word in words)
    ]


def _fonts(doc: DocxDocument) -> list[str]:
    names = {f.get(qn("w:ascii")) for f in doc.element.iter(qn("w:rFonts")) if f.get(qn("w:ascii"))}
    normal = doc.styles["Normal"].font.name if "Normal" in [s.name for s in doc.styles] else None
    names |= {normal} if normal else set()
    return sorted(n for n in names if n and n.casefold() not in COMMON_FONTS)


def layout_warnings(path: Path, cv: MasterCV) -> list[str]:
    """What in the file's layout a parser is likely to miss or misread (warnings only)."""
    doc = Document(str(path))
    lines = _body_lines(doc)
    b = cv.basics
    warnings: list[str] = []
    if _columns(doc):
        warnings.append("Two or more columns: parsers read straight across and mix them")
    footer = _plain(_footer_text(doc))
    if any(v and _plain(v) in footer for v in (b.email, b.phone)):
        warnings.append("Contact details in the page footer: many ATS parsers skip footers")
    images, boxes = _boxes(doc)
    if boxes:
        warnings.append(f"{boxes} text box(es): parsers do not read the text inside them")
    if images:
        warnings.append(f"{images} image(s): any text in an image is not read")
    if undated := _undated(lines, cv):
        warnings.append(
            "No year next to the role: " + ", ".join(undated) + " (systems count your years "
            "of experience from the dates; use e.g. 'Mar 2021 – Present')"
        )
    if missing := _missing_headings(lines):
        warnings.append(
            "No standard section heading for: " + ", ".join(missing) + " (e.g. 'Experience', "
            "'Education'); parsers file content by these headings"
        )
    if fonts := _fonts(doc):
        warnings.append("Unusual font(s): " + ", ".join(fonts) + " (Arial or Calibri read safest)")
    return warnings


def check_docx(path: Path, cv: MasterCV, keywords: list[str] | None = None) -> ATSReport:
    """Read `path` back and report what an ATS would miss (`keywords`: the job's)."""
    text, tables = docx_text(path)
    # red suggestions are not claims: whole red lines, and red words added inside a line
    confirmed_text = "\n".join(
        INLINE_REVIEW.sub("", line)
        for line in text.splitlines()
        if not line.startswith((UNVERIFIED_PREFIX, REVIEW_PREFIX))
    )
    b = cv.basics
    wanted = {"name": b.name, "email": b.email, "phone": b.phone}
    missing = [k for k, v in wanted.items() if v and _plain(v) not in _plain(text)]
    keywords = list(dict.fromkeys(keywords or []))
    absent = [k for k in keywords if not term_found(confirmed_text, k)]  # as ATS search
    words = len(text.split())
    pages = round(words / WORDS_PER_PAGE, 1)
    warnings: list[str] = []
    if missing:
        warnings.append(f"Contact details not readable as text: {', '.join(missing)}")
    if tables:
        warnings.append(f"{tables} table(s): some ATS parsers scramble their reading order")
    if any(v and _plain(v) in _plain(_header_text(path)) for v in wanted.values()):
        warnings.append("Contact details in the page header: many ATS parsers skip headers")
    warnings += layout_warnings(path, cv)
    if confirmed_text != text:
        warnings.append("Red suggestions are unverified; confirm or remove them before applying")
    return ATSReport(
        words=words,
        est_pages=pages,
        contact_missing=missing,
        keyword_coverage=(len(keywords) - len(absent)) / len(keywords) if keywords else 1.0,
        missing_keywords=absent,
        warnings=warnings,
    )
