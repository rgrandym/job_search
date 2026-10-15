"""Render a `MasterCV` / `TailoredCV` to a styled, ATS-friendly Word document.

ATS constraints honoured by every template: single column, no tables or text boxes,
no headers/footers holding content, native bullet lists, standard section names.
"""

from __future__ import annotations

from pathlib import Path
from typing import Literal

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph
from pydantic import BaseModel, Field

from src.cv.models import CoverLetter, MasterCV, ReviewAddition, TailoredCV
from src.tools import docx_tools as dt

Section = Literal[
    "summary", "experience", "skills", "projects", "education", "certifications", "languages"
]

DEFAULT_SECTIONS: list[Section] = [
    "summary",
    "experience",
    "skills",
    "projects",
    "education",
    "certifications",
    "languages",
]
REVIEW_RED = RGBColor(0xC0, 0x00, 0x00)  # a changed line in a draft; a questioned letter line
_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


class DocxTemplate(BaseModel):
    """Visual style of an exported CV."""

    name: str
    font: str = "Calibri"
    body_pt: float = Field(10.5, ge=8, le=14)
    name_pt: float = 20
    heading_pt: float = 12
    accent: str = "#1F3864"
    margins_in: float = 0.7
    centered_header: bool = True
    uppercase_headings: bool = True
    sections: list[Section] = Field(default_factory=lambda: list(DEFAULT_SECTIONS))


TEMPLATES: dict[str, DocxTemplate] = {
    "classic": DocxTemplate(name="classic"),
    "modern": DocxTemplate(
        name="modern",
        font="Arial",
        body_pt=10,
        name_pt=22,
        accent="#0F766E",
        centered_header=False,
        uppercase_headings=False,
    ),
    "compact": DocxTemplate(
        name="compact",
        font="Calibri",
        body_pt=9.5,
        name_pt=16,
        heading_pt=11,
        margins_in=0.5,
        sections=["summary", "skills", "experience", "education", "certifications"],
    ),
}


def export_docx(
    cv: MasterCV | TailoredCV, path: Path, template: str = "classic", final: bool = False
) -> Path:
    """Write `cv` to `path` (.docx) using a named template from `TEMPLATES`. A tailored
    draft shows every rewritten or added line in red; `final` writes the clean version: no
    red, no suggested headlines and nothing pending review."""
    if template not in TEMPLATES:
        raise ValueError(f"Unknown template {template!r}. Choose from {sorted(TEMPLATES)}")
    tpl = TEMPLATES[template]
    tailored = clean_copy(cv) if isinstance(cv, TailoredCV) and final else cv
    draft = tailored if isinstance(tailored, TailoredCV) else None
    master = draft.cv if draft else cv if isinstance(cv, MasterCV) else cv.cv
    options = draft.headline_options if draft else []
    changed = {c.source_id for c in draft.changes if c.accepted} if draft else set()
    sections = list(tpl.sections)
    if draft and draft.academic and "education" in sections:
        sections.remove("education")  # academic readers look for education first
        sections.insert(sections.index("summary") + 1 if "summary" in sections else 0, "education")

    doc = Document()
    dt.set_margins(doc, tpl.margins_in)
    dt.set_base_font(doc, tpl.font, tpl.body_pt)
    _header(doc, master, tpl, options)
    marked = list(doc.paragraphs[1 + bool(master.basics.headline) :][: len(options)])
    for section in sections:
        before = len(doc.paragraphs)
        if section == "experience" and draft:
            _experience(doc, master, tpl, draft.review_additions)
            texts = {b.text for b in master.bullet_index().values() if b.id in changed}
            marked += [p for p in doc.paragraphs[before:] if p.text in texts]
        else:
            RENDERERS[section](doc, master, tpl)
        if section == "summary" and "summary" in changed:
            marked += doc.paragraphs[before + 1 :]
        if section == "summary" and draft:
            marked += _summary_extras(doc, draft)
    for paragraph in [] if final else marked:
        for run in paragraph.runs:
            run.font.color.rgb = REVIEW_RED
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def clean_copy(cv: TailoredCV) -> TailoredCV:
    """The final version's content: no suggested headlines, nothing pending review."""
    return cv.model_copy(
        update={
            "headline_options": [],
            "skills_to_confirm": [],
            "review_additions": [a for a in cv.review_additions if a.status != "pending"],
            "changes": [],
        }
    )


def _summary_extras(doc: DocxDocument, cv: TailoredCV) -> list[Paragraph]:
    """The lines under the summary: the relevant-skills line (returned, to be marked as a
    change) and, in red, what is still to confirm."""
    marked = [doc.add_paragraph(f"• {text}") for text in cv.key_strengths]
    if cv.role_skills:
        line = f"Relevant skills for this role: {' · '.join(cv.role_skills)}"
        marked.append(doc.add_paragraph(line))
    pending = [f"[CONFIRM SKILL BEFORE USE] {a.text}" for a in cv.review_additions
               if a.status == "pending" and a.kind == "skill"]  # fmt: skip
    if cv.skills_to_confirm:
        skills = " · ".join(s.skill for s in cv.skills_to_confirm)
        pending.insert(0, f"Skills to confirm before use (not yet claimed): {skills}")
    for text in pending:
        doc.add_paragraph().add_run(text).font.color.rgb = REVIEW_RED
    return marked


def fmt_date(ym: str | None) -> str:
    """'2021-03' -> 'Mar 2021'; None -> 'Present'."""
    if ym is None:
        return "Present"
    year, month = ym.split("-")[:2] if "-" in ym else (ym, "")
    return f"{_MONTHS[int(month) - 1]} {year}" if month else year


def _heading(doc: DocxDocument, text: str, tpl: DocxTemplate) -> None:
    label = text.upper() if tpl.uppercase_headings else text
    p = dt.add_text(doc, label, size_pt=tpl.heading_pt, bold=True, color=tpl.accent)
    p.paragraph_format.space_before = Pt(10)
    dt.add_bottom_border(p, color=tpl.accent)


def _header(
    doc: DocxDocument, cv: MasterCV, tpl: DocxTemplate, headline_options: list[str]
) -> None:
    """Name, headline (then any suggested headlines, for the user to keep or delete), contact."""
    align = WD_ALIGN_PARAGRAPH.CENTER if tpl.centered_header else WD_ALIGN_PARAGRAPH.LEFT
    b = cv.basics
    dt.add_text(
        doc, b.name, size_pt=tpl.name_pt, bold=True, color=tpl.accent, align=align, space_after_pt=0
    )
    headlines = [b.headline] if b.headline else []
    headlines += [h for h in headline_options if h.strip() and h not in headlines]
    for headline in headlines:
        dt.add_text(doc, headline, size_pt=tpl.body_pt + 1, align=align, space_after_pt=0)
    contact = [x for x in (b.location, b.email, b.phone) if x]
    contact += [link.url for link in b.links]
    if contact:
        dt.add_text(doc, "  |  ".join(contact), size_pt=tpl.body_pt - 1, align=align)


def _summary(doc: DocxDocument, cv: MasterCV, tpl: DocxTemplate) -> None:
    if cv.basics.summary:
        _heading(doc, "Summary", tpl)
        doc.add_paragraph(cv.basics.summary)


def _experience(
    doc: DocxDocument,
    cv: MasterCV,
    tpl: DocxTemplate,
    additions: list[ReviewAddition] | None = None,
) -> None:
    if not cv.experience:
        return
    _heading(doc, "Experience", tpl)
    for e in cv.experience:
        dt.add_split_line(
            doc,
            f"{e.title}, {e.company}",
            f"{fmt_date(e.start)} – {fmt_date(e.end)}",
            right_tab_in=8.5 - 2 * tpl.margins_in,
        )
        if e.location:
            dt.add_text(doc, e.location, italic=True, size_pt=tpl.body_pt - 0.5, space_after_pt=1)
        for bullet in e.bullets:
            dt.add_bullet(doc, bullet.text)
        for item in additions or []:
            if item.status == "pending" and item.kind == "experience" and item.role_id == e.id:
                run = dt.add_bullet(doc, f"[CONFIRM EXPERIENCE BEFORE USE] {item.text}").runs[0]
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)


def _skills(doc: DocxDocument, cv: MasterCV, tpl: DocxTemplate) -> None:
    if not cv.skills:
        return
    _heading(doc, "Skills", tpl)
    for g in cv.skills:
        p = doc.add_paragraph()
        p.add_run(f"{g.category}: ").bold = True
        p.add_run(", ".join(g.items))


def _projects(doc: DocxDocument, cv: MasterCV, tpl: DocxTemplate) -> None:
    if not cv.projects:
        return
    _heading(doc, "Projects", tpl)
    for pr in cv.projects:
        p = doc.add_paragraph()
        p.add_run(pr.name).bold = True
        p.add_run(f": {pr.description}")


def _education(doc: DocxDocument, cv: MasterCV, tpl: DocxTemplate) -> None:
    if not cv.education:
        return
    _heading(doc, "Education", tpl)
    for ed in cv.education:
        degree = f"{ed.degree}, {ed.field}" if ed.field else ed.degree
        dates = (
            f"{fmt_date(ed.start)} – {fmt_date(ed.end)}"
            if ed.start
            else fmt_date(ed.end)
            if ed.end
            else ""
        )
        dt.add_split_line(
            doc, f"{degree}, {ed.institution}", dates, right_tab_in=8.5 - 2 * tpl.margins_in
        )
        for d in ed.details:
            dt.add_bullet(doc, d)


def _certifications(doc: DocxDocument, cv: MasterCV, tpl: DocxTemplate) -> None:
    if not cv.certifications:
        return
    _heading(doc, "Certifications", tpl)
    for c in cv.certifications:
        bits = [c.name] + [x for x in (c.issuer, str(c.year) if c.year else None) if x]
        dt.add_bullet(doc, ", ".join(bits))


def _languages(doc: DocxDocument, cv: MasterCV, tpl: DocxTemplate) -> None:
    if cv.languages:
        _heading(doc, "Languages", tpl)
        doc.add_paragraph(", ".join(cv.languages))


def export_cover_letter(
    letter: CoverLetter, cv: MasterCV, path: Path, template: str = "classic"
) -> Path:
    """Write a checked cover letter to `path` (.docx), styled like the CV template ("original",
    a CV in its own Word design, gives the letter the classic style). Flagged paragraphs are
    red, for the user to review."""
    template = "classic" if template == "original" else template
    if template not in TEMPLATES:
        raise ValueError(f"Unknown template {template!r}. Choose from {sorted(TEMPLATES)}")
    tpl = TEMPLATES[template]
    doc = Document()
    dt.set_margins(doc, max(tpl.margins_in, 0.9))
    dt.set_base_font(doc, tpl.font, tpl.body_pt + 0.5)
    greeting = doc.add_paragraph(letter.greeting)
    greeting.paragraph_format.space_after = Pt(12)
    flagged = {flag.paragraph for flag in letter.flags}
    for index, para in enumerate(letter.paragraphs):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(para)
        if index in flagged:
            run.font.color.rgb = REVIEW_RED
        paragraph.paragraph_format.space_after = Pt(9)
    closing = doc.add_paragraph(letter.closing)
    closing.paragraph_format.space_before = Pt(3)
    doc.add_paragraph(cv.basics.name)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


RENDERERS = {
    "summary": _summary,
    "experience": _experience,
    "skills": _skills,
    "projects": _projects,
    "education": _education,
    "certifications": _certifications,
    "languages": _languages,
}
